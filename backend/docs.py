"""Documents API: upload PDF/DOCX/ODF, write notes, choose what the LLM sees.

    GET    /api/documents
    POST   /api/documents                 multipart: file, kind
    POST   /api/documents/notes           {title, text, kind?, use_for_scoring?}: a note written in the app
    POST   /api/documents/summary         consolidate the ticked documents into one new summary document (LLM)
    PATCH  /api/documents/{id}            {kind?, use_for_scoring?, title?, text?, concise_text?, use_concise?}
    POST   /api/documents/{id}/concise    write (or rewrite) the document's concise version (LLM)
    GET    /api/documents/{id}/text       the original and concise text
    DELETE /api/documents/{id}

Notes and summaries have no file (``filetype`` "note" / "summary"): the title is ``filename``, the body
``text``, and both can be edited. An uploaded file keeps the text extracted from it. Any document can also have a
concise version (``concise_text``, editable); ``use_concise`` picks which one the LLM gets (Document.text_for_llm).
Changing what the LLM gets changes the profile hash, so earlier scores show as stale.

Kinds (DOCUMENT_KINDS): resume, cover_letter, preferences, other, letter_template. Preferences say what the candidate
wants: the scorer, the chat and title suggestions read them, cover letters never do (backend/cover_letters.py).

A letter template (kind letter_template) is a .docx or .odt letterhead with a {letter} marker, checked at upload
(backend/letter_files.py). Its file is what the cover-letter downloads are poured into; its text is shown but never
sent to the AI: ``use_for_scoring`` stays off, the kind can't be changed into or out of it, and every prompt builder
skips the kind. The workspace's current template (``workspaces.letter_template_id``):

    GET    /api/documents/letter-template   the template in use, the formats it gives and its fonts (export_options)
    PUT    /api/documents/letter-template   {id} chooses a template of this workspace; {id: null} the plain layout

Both signed in, like the rest.

The concise version and the summary are written by the configured LLM at the cover-letter task's reasoning
level (both write from the documents), and must keep every fact: the scorer quotes them as evidence.
"""
from __future__ import annotations

import io
import re
import uuid
from datetime import datetime
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from backend import crypto, letter_files, llm_providers
from backend.db import (
    DOCUMENT_KINDS,
    LETTER_TEMPLATE_KIND,
    NOTE_FILETYPE,
    SUMMARY_FILETYPE,
    UPLOAD_DIR,
    Document,
    Workspace,
    get_db,
)
from backend.workspaces import current_workspace, owned

router = APIRouter(prefix="/api/documents", tags=["documents"])

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".odt", ".ods", ".odp", ".txt", ".md"}
TITLE_MAX = 200
TEXT_MAX = 300_000  # characters; the largest bundles seen are about 230k for all documents together
EDITABLE_FILETYPES = (NOTE_FILETYPE, SUMMARY_FILETYPE)


def _extract_pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    return "\n".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(data)).pages).strip()


def _extract_docx_text(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts).strip()


def _extract_odf_text(data: bytes) -> str:
    from odf import teletype
    from odf import text as odf_text
    from odf.opendocument import load

    doc = load(io.BytesIO(data))
    return "\n".join(teletype.extractText(p) for p in doc.getElementsByType(odf_text.P)).strip()


def extract_text(data: bytes, suffix: str) -> str:
    """The text of an uploaded file, read from memory: the plaintext file never touches the disk."""
    suffix = suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf_text(data)
    if suffix == ".docx":
        return _extract_docx_text(data)
    if suffix in {".odt", ".ods", ".odp"}:
        return _extract_odf_text(data)
    if suffix in {".txt", ".md"}:
        return data.decode("utf-8", errors="replace").strip()
    raise ValueError(f"Unsupported file type: {suffix}")


# --------------------------------------------------------------------------
# Uploads: sealed under the owner's key in data/uploads as "<random hex>.sealed" (accounts design §2.2)
# --------------------------------------------------------------------------

SEALED_SUFFIX = ".sealed"
MAX_UPLOAD_BYTES = 20 * 1024 * 1024  # held in memory (never spooled to disk in plaintext): app.py sizes the spool


def file_label(stored_name: str) -> str:
    """The AAD label of an upload: the file opens only under its own name."""
    return f"file|{stored_name}"


def new_stored_name() -> str:
    return f"{uuid.uuid4().hex}{SEALED_SUFFIX}"


def save_upload(data: bytes) -> str:
    """Seal ``data`` under the current user's key into a new upload file. Returns its stored name."""
    stored = new_stored_name()
    crypto.seal_file(data, UPLOAD_DIR / stored, file_label(stored))
    return stored


def read_upload(stored_name: str) -> bytes:
    """An upload's plaintext, in memory (KeyMismatch / Corrupt if it isn't this user's, or was moved)."""
    return crypto.open_file(UPLOAD_DIR / stored_name, file_label(stored_name))


def delete_upload(stored_name: Optional[str]) -> None:
    if stored_name:
        (UPLOAD_DIR / stored_name).unlink(missing_ok=True)


def _iso(value: Optional[datetime]) -> Optional[str]:
    return value.isoformat() if value else None


def source_of(d: Document) -> str:
    """note | summary | file"""
    return d.filetype if d.filetype in EDITABLE_FILETYPES else "file"


def _current_template_id(db: Session, ws: int) -> Optional[int]:
    return db.query(Workspace.letter_template_id).filter(Workspace.id == ws).scalar()


def doc_to_dict(d: Document) -> dict:
    from sqlalchemy.orm import object_session

    concise = (d.concise_text or "").strip()
    is_template = d.kind == LETTER_TEMPLATE_KIND
    db = object_session(d) if is_template else None
    return {
        # a letter template: whether cover-letter downloads use this one (backend/cover_letters.py)
        "current_template": bool(is_template and db is not None and _current_template_id(db, d.workspace_id) == d.id),
        "id": d.id,
        "filename": d.filename,
        "filetype": d.filetype,
        "source": source_of(d),
        "editable": d.filetype in EDITABLE_FILETYPES,  # the text itself can be edited (notes, summaries)
        "kind": d.kind or "other",
        "use_for_scoring": bool(d.use_for_scoring),
        "created_at": _iso(d.created_at),
        "updated_at": _iso(d.updated_at),
        "text_len": len(d.text or ""),
        "concise_len": len(concise),
        "use_concise": bool(d.use_concise and concise),  # the LLM gets the concise version
        "concise_model": d.concise_model,
        "concise_at": _iso(d.concise_at),
        "sent_len": len(d.text_for_llm()),
        "summary_of": [i for i in (d.summary_of or []) if isinstance(i, int)],
    }


def _kind(kind: str) -> str:
    if kind not in DOCUMENT_KINDS:
        raise HTTPException(status_code=422, detail=f"kind must be one of {DOCUMENT_KINDS}")
    return kind


class NoteCreate(BaseModel):
    title: str = Field(min_length=1, max_length=TITLE_MAX)
    text: str = Field(min_length=1, max_length=TEXT_MAX)
    kind: str = "other"
    use_for_scoring: bool = True


class DocumentUpdate(BaseModel):
    kind: Optional[str] = None
    use_for_scoring: Optional[bool] = None
    title: Optional[str] = Field(default=None, max_length=TITLE_MAX)        # rename
    text: Optional[str] = Field(default=None, max_length=TEXT_MAX)          # notes and summaries only
    concise_text: Optional[str] = Field(default=None, max_length=TEXT_MAX)  # "" removes the concise version
    use_concise: Optional[bool] = None                                      # send the concise version


@router.get("")
def list_documents(db: Session = Depends(get_db), ws: int = Depends(current_workspace)):
    rows = db.query(Document).filter(Document.workspace_id == ws).order_by(Document.created_at.desc())
    return [doc_to_dict(d) for d in rows]


@router.post("")
def upload_document(file: UploadFile = File(...), kind: str = Form("resume"), db: Session = Depends(get_db),
                    ws: int = Depends(current_workspace)):
    _kind(kind)
    original = file.filename or "upload"
    suffix = Path(original).suffix.lower()
    template = kind == LETTER_TEMPLATE_KIND
    if template and suffix not in letter_files.TEMPLATE_SUFFIXES:
        raise HTTPException(status_code=400, detail="A letter template is a .docx or .odt file with a {letter} marker "
                                                    "where the letter should go.")
    if suffix not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"Unsupported file type {suffix!r}. Allowed: {sorted(ALLOWED_EXTENSIONS)}")
    data = file.file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"That file is over {MAX_UPLOAD_BYTES // (1024 * 1024)} MB. "
                                                    "Upload a smaller copy (for example, without images).")
    if template:
        try:
            letter_files.check_template(data, suffix)  # refused unless it has the marker and can be filled
        except letter_files.TemplateError as exc:
            raise HTTPException(status_code=400, detail=str(exc))
    try:
        text = extract_text(data, suffix)
    except Exception as exc:
        if template:
            text = ""
        else:
            raise HTTPException(status_code=400, detail=f"Could not extract text: {exc}")
    if not text and not template:  # a template may be a letterhead of pictures and a marker: its text doesn't matter
        raise HTTPException(status_code=400, detail="No text found in the file (scanned PDF?)")
    stored = save_upload(data)
    doc = Document(workspace_id=ws, filename=original, filetype=suffix, stored_name=stored, text=text, kind=kind,
                   use_for_scoring=not template)
    db.add(doc)
    try:
        db.flush()
        if template and not _current_template_id(db, ws):  # the first template becomes the one in use
            db.get(Workspace, ws).letter_template_id = doc.id
        db.commit()
    except Exception:
        db.rollback()
        delete_upload(stored)
        raise
    return doc_to_dict(doc)


@router.post("/notes", status_code=201)
def create_note(payload: NoteCreate, db: Session = Depends(get_db), ws: int = Depends(current_workspace)):
    """A note written in the app: a titled piece of plain text (project details, preferences...) that takes part
    like an uploaded document. It has no file."""
    title, text = payload.title.strip(), payload.text.strip()
    if not title or not text:
        raise HTTPException(status_code=422, detail="A note needs a title and some text.")
    now = datetime.utcnow()
    doc = Document(workspace_id=ws, filename=title, filetype=NOTE_FILETYPE, stored_name=None, text=text,
                   kind=_kind(payload.kind), use_for_scoring=payload.use_for_scoring, created_at=now, updated_at=now)
    db.add(doc)
    db.commit()
    return doc_to_dict(doc)


@router.patch("/{doc_id}")
def update_document(doc_id: int, payload: DocumentUpdate, db: Session = Depends(get_db),
                    ws: int = Depends(current_workspace)):
    doc = owned(db, Document, doc_id, ws, "Document")
    edited = False
    if payload.kind is not None:
        if (payload.kind == LETTER_TEMPLATE_KIND) != (doc.kind == LETTER_TEMPLATE_KIND) and payload.kind != doc.kind:
            raise HTTPException(status_code=400, detail="A letter template is uploaded as one (its file is checked for "
                                                        "the {letter} marker); a document's type can't be changed "
                                                        "into or out of it.")
        doc.kind = _kind(payload.kind)
    if payload.use_for_scoring is not None:
        if payload.use_for_scoring and doc.kind == LETTER_TEMPLATE_KIND:
            raise HTTPException(status_code=400, detail="A letter template is never sent to the AI.")
        doc.use_for_scoring = payload.use_for_scoring
    if payload.title is not None:
        if not payload.title.strip():
            raise HTTPException(status_code=422, detail="The title can't be empty.")
        doc.filename, edited = payload.title.strip(), True
    if payload.text is not None:
        if doc.filetype not in EDITABLE_FILETYPES:
            raise HTTPException(status_code=400, detail="The text of an uploaded file can't be edited. "
                                                        "Edit its concise version, or write a note instead.")
        if not payload.text.strip():
            raise HTTPException(status_code=422, detail="The text can't be empty; delete the document instead.")
        doc.text, edited = payload.text.strip(), True
    if payload.concise_text is not None:
        doc.concise_text = payload.concise_text.strip() or None
        if doc.concise_text is None:
            doc.use_concise = False
        edited = True
    if payload.use_concise is not None:
        if payload.use_concise and not (doc.concise_text or "").strip():
            raise HTTPException(status_code=409, detail="This document has no concise version yet.")
        doc.use_concise = payload.use_concise
    if edited:
        doc.updated_at = datetime.utcnow()
    db.commit()
    return doc_to_dict(doc)


@router.get("/{doc_id}/text")
def document_text(doc_id: int, db: Session = Depends(get_db), ws: int = Depends(current_workspace)):
    doc = owned(db, Document, doc_id, ws, "Document")
    return {**doc_to_dict(doc), "text": doc.text or "", "concise_text": doc.concise_text or ""}


@router.delete("/{doc_id}")
def delete_document(doc_id: int, db: Session = Depends(get_db), ws: int = Depends(current_workspace)):
    doc = owned(db, Document, doc_id, ws, "Document")
    stored = doc.stored_name  # notes and summaries have no file
    if _current_template_id(db, ws) == doc_id:  # the template in use: cover letters go back to the plain layout
        db.get(Workspace, ws).letter_template_id = None
    db.delete(doc)
    db.commit()
    delete_upload(stored)
    return {"ok": True, "id": doc_id}


# --------------------------------------------------------------------------
# The letter template in use (per workspace)
# --------------------------------------------------------------------------

class TemplateChoice(BaseModel):
    id: Optional[int] = None  # None: the plain layout


@router.get("/letter-template")
def letter_template(db: Session = Depends(get_db), ws: int = Depends(current_workspace)):
    from backend.cover_letters import export_options

    return export_options(db, ws)


@router.put("/letter-template")
def choose_letter_template(payload: TemplateChoice, db: Session = Depends(get_db), ws: int = Depends(current_workspace)):
    from backend.cover_letters import export_options

    if payload.id is not None:
        doc = owned(db, Document, payload.id, ws, "Document")
        if doc.kind != LETTER_TEMPLATE_KIND or not doc.stored_name:
            raise HTTPException(status_code=400, detail="That document isn't a letter template.")
    db.get(Workspace, ws).letter_template_id = payload.id
    db.commit()
    return export_options(db, ws)


# --------------------------------------------------------------------------
# Concise versions and the summary document (written by the LLM)
# --------------------------------------------------------------------------

_FACTS = """Keep every concrete fact, worded as in the original:
- employers, organisations, clients, products and projects; job titles; dates and durations
- numbers and outcomes: team sizes, budgets, user counts, percentages, results
- skills, methods, tools and technologies
- qualifications, degrees, certifications, licences and clearances, each with its status and caveats exactly as written (e.g. "not yet conferred", "in progress", "expected 2027")
- distinctive phrases worth quoting: keep them verbatim"""

_NO_INVENTION = """Never invent, infer, round or generalise anything, never combine separate numbers or achievements into one, and never make a claim stronger than the original ("contributed to" stays "contributed to"). Drop only repetition, filler, greetings, sign-offs and claims with no specifics ("passionate team player"). When unsure whether something matters, keep it."""

CONCISE_PROMPT = f"""You shorten one document from a job seeker's files. An automated job-fit assessor will read your version instead of the original and quote it as evidence, so accuracy matters more than brevity.

{_FACTS}{{kind_rule}}

{_NO_INVENTION}

Write plain text: short headings and "- " bullet points are fine; no tables, no remarks about what you changed. Return only the shortened document."""

_KIND_RULES = {
    "preferences": "This document says what the candidate wants from a job: keep every preference, condition and "
                   "deal-breaker, and how strongly it is held (\"must\", \"prefer\", \"open to\").",
    "cover_letter": "This is a cover letter: keep what it says about the candidate; leave out the address, the "
                    "salutation and what it says only about that employer.",
}

SUMMARY_PROMPT = f"""You merge a job seeker's documents (resumes, cover letters, notes and other supporting documents) into ONE factual profile. An automated job-fit assessor will read it instead of the documents and quote it as evidence, so accuracy matters more than brevity.

How to merge:
- State each fact once, in its most complete and specific form. Keep every detail that appears in only one document.
- Where documents disagree (e.g. different dates for the same role), keep both versions and say which document says which; never pick one.
- Cover letters were written to particular employers: keep the facts they give about the candidate, not what they say about that employer or job.
- Leave out what the candidate wants from a job (pay, location, work mode, kind of role): that belongs in their preferences, not in this profile.

{_FACTS.replace("as in the original", "as in the documents")}

{_NO_INVENTION.replace("than the original", "than the documents do")}

Layout: plain text with these headings and "- " bullet points; leave out a section the documents say nothing for.
SUMMARY: 2-4 sentences, only facts from the documents
EXPERIENCE: one entry per role (title, employer, dates), most recent first, with its responsibilities, achievements and numbers
PROJECTS AND ACHIEVEMENTS outside those roles
SKILLS, METHODS AND TOOLS
EDUCATION, QUALIFICATIONS AND CERTIFICATIONS, each with its status
OTHER FACTS (e.g. languages, work rights, clearances, availability)

Return only the profile, with no remarks about what you did."""


def concise_prompt(kind: Optional[str]) -> str:
    rule = _KIND_RULES.get(kind or "")
    return CONCISE_PROMPT.replace("{kind_rule}", f"\n\n{rule}" if rule else "")


def _llm_config(ws: int, purpose: str) -> dict:
    from backend.scorer import config_problem, get_config

    cfg = {**get_config("cover_letter"), "purpose": purpose, "workspace_id": ws, "long_output": True}
    cfg["timeout_s"] = max(cfg["timeout_s"], 600)  # a long document takes minutes to write
    if config_problem(cfg):
        raise HTTPException(status_code=400, detail=config_problem(cfg))
    return cfg


def _clean(text: str) -> str:
    text = (text or "").strip()
    fence = re.match(r"^```\w*\n(.*?)\n```$", text, re.S)
    return fence.group(1).strip() if fence else text


def _rewrite(cfg: dict, system: str, user: str, task: str, document: Optional[int] = None) -> str:
    """The model's rewrite; in the tasks tray as ``task`` ("concise" for ``document``, or "summary")."""
    from backend import tasks
    from backend.scorer import ScorerError, call_model

    try:
        with tasks.track(task, cfg["workspace_id"], document=document):
            text = _clean(call_model([{"role": "system", "content": system}, {"role": "user", "content": user}],
                                     cfg, json_mode=False))
    except ScorerError as exc:
        raise HTTPException(status_code=502, detail=llm_providers.user_message(exc))
    if not text:
        raise HTTPException(status_code=502, detail="The model returned no text; try again.")
    return text


@router.post("/{doc_id}/concise")
def make_concise(doc_id: int, db: Session = Depends(get_db), ws: int = Depends(current_workspace)):
    """Write the document's concise version from its original text (replacing an earlier one) and send that to
    the LLM from now on. The user can edit it, or switch back to the original."""
    from backend.scorer import document_heading

    doc = owned(db, Document, doc_id, ws, "Document")
    if doc.kind == LETTER_TEMPLATE_KIND:
        raise HTTPException(status_code=400, detail="A letter template is never sent to the AI, so it has no concise version.")
    original = (doc.text or "").strip()
    if not original:
        raise HTTPException(status_code=400, detail="This document has no text to shorten.")
    cfg = _llm_config(ws, "concise")
    user = f"THE DOCUMENT:\n{document_heading(doc)}\n<<<\n{original}\n>>>\n\nReturn only the shortened document."
    doc.concise_text = _rewrite(cfg, concise_prompt(doc.kind), user, "concise", doc_id)
    now = datetime.utcnow()
    doc.use_concise, doc.concise_model, doc.concise_at, doc.updated_at = True, cfg["model"], now, now
    db.commit()
    return {**doc_to_dict(doc), "text": doc.text or "", "concise_text": doc.concise_text}


@router.post("/summary", status_code=201)
def make_summary(db: Session = Depends(get_db), ws: int = Depends(current_workspace)):
    """Consolidate every ticked document (not preferences) into one new summary document, from their full text.
    It is created ticked, as a resume (the scorer's primary source). The originals are left as they are: the UI
    suggests unticking them, so their content isn't sent twice."""
    from backend.scorer import _KIND_ORDER, document_heading

    docs = [d for d in db.query(Document).filter(Document.workspace_id == ws)
            if d.use_for_scoring and d.kind not in ("preferences", LETTER_TEMPLATE_KIND) and (d.text or "").strip()]
    if not docs:
        raise HTTPException(status_code=400, detail="Tick \"use for scoring\" on the documents to summarise "
                                                    "first (preferences are kept separate).")
    docs.sort(key=lambda d: (_KIND_ORDER.get(d.kind or "other", 2), d.id))
    cfg = _llm_config(ws, "summary")
    body = "\n\n".join(f"{document_heading(d)}\n{d.text.strip()}" for d in docs)
    user = (f"THE CANDIDATE'S DOCUMENTS ({len(docs)}):\n<<<\n{body}\n>>>\n\n"
            "Return only the consolidated profile.")
    text = _rewrite(cfg, SUMMARY_PROMPT, user, "summary")
    now = datetime.utcnow()
    summary = Document(workspace_id=ws, filename=f"Summary of {len(docs)} document{'s' if len(docs) != 1 else ''}",
                       filetype=SUMMARY_FILETYPE, stored_name=None, text=text, kind="resume", use_for_scoring=True,
                       summary_of=[d.id for d in docs], created_at=now, updated_at=now)
    db.add(summary)
    db.commit()
    return {**doc_to_dict(summary), "text": text, "concise_text": "",
            "sources": [{"id": d.id, "filename": d.filename} for d in docs]}
