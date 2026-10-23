"""Cover letter template, and a saved letter as .docx, .odt or .pdf.

    GET    /api/letter-template           {template: {filename, filetype, created_at} | null, pdf, pdf_uses_template}
    POST   /api/letter-template           multipart: file (.docx or .odt): replaces the workspace's template
    DELETE /api/letter-template

All of them need a signed-in user. A workspace has at most one template: a Word (.docx) or OpenDocument (.odt) file
with a paragraph holding ``{letter}`` where the letter goes (an upload without one is refused). The file is an upload
like a document's: sealed under the owner's key (``docs.save_upload``), its name sealed in the row.

**Filling** (``fill_docx``, ``fill_odt``): ``{company}``, ``{role}``, ``{location}`` and ``{date}``
(``letter_blocks.PLACEHOLDERS``) are filled in every paragraph of the template (body, tables, text boxes, headers and
footers; a placeholder with no value stays as written). The first body paragraph with ``{letter}`` takes the letter:
the letter is split into paragraphs at blank lines, the marker paragraph gets the first one (in place of ``{letter}``,
with the paragraph's own text before it), and each further one is a new paragraph after it, a copy of the marker
paragraph (its style, and its first run's character formatting), the last one ending with the marker paragraph's text
after ``{letter}``. A line break inside a letter paragraph stays a line break. Any other ``{letter}`` is removed. A
paragraph whose placeholders each sit in one run (span) keeps its formatting; one with a placeholder split across runs
is written again as one run with its first run's formatting (Word and LibreOffice split text where its formatting or
spell check changed).

**Downloads** (``render``, for ``GET /api/jobs/{id}/cover-letter.{docx|odt|pdf}``, backend/cover_letters.py):

- ``.docx``: the filled template when it is a .docx, else a plain document (one paragraph per blank-line block);
- ``.odt``: the filled template when it is an .odt, else a plain document;
- ``.pdf``: with LibreOffice installed (``pdf_backend()`` "libreoffice"), the filled template (of either type), or the
  plain .docx, converted by LibreOffice; without it ("plain"), a plain-text PDF (fpdf2, Helvetica, A4) that never
  uses the template. ``GET`` says which (``pdf``, ``pdf_uses_template``).

The LibreOffice conversion is the only place a letter is written to disk: into a new folder of its own inside the data
folder (``.letter-pdf-*``, readable by this user only), removed again when the conversion ends however it ends; a
folder a crash left is removed by the next conversion (and by the reset tool). The log says only that a conversion
failed and how (exit status, time-out), never a name, a path or any text.
"""
from __future__ import annotations

import copy
import io
import logging
import os
import re
import shutil
import subprocess
import tempfile
import time
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile
from sqlalchemy.orm import Session

from backend import crypto, docs
from backend.db import LetterTemplate, get_db
from backend.export_formats import clean
from backend.letter_blocks import _PLACEHOLDER, PLACEHOLDERS, placeholder_values
from backend.workspaces import current_workspace

log = logging.getLogger(__name__)

router = APIRouter(prefix="/api/letter-template", tags=["letter templates"])

TEMPLATE_TYPES = (".docx", ".odt")
FORMATS = ("docx", "odt", "pdf")
MEDIA_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "odt": "application/vnd.oasis.opendocument.text",
    "pdf": "application/pdf",
}
MARKER = "{letter}"
FILENAME_MAX = 255
PDF_TIMEOUT_S = 120
TEMP_PREFIX = ".letter-pdf-"
STALE_TEMP_S = 15 * 60  # older than any conversion can run: left by a crash


class TemplateError(Exception):
    """A template that can't be used. ``kind``: "unreadable" or "no_marker" (never the file's text)."""

    def __init__(self, kind: str) -> None:
        super().__init__(kind)
        self.kind = kind


class PdfError(Exception):
    """LibreOffice didn't make the PDF. The message says how (an exit status, a time-out), never a name or text."""


UNREADABLE = "That file could not be read as a Word/OpenDocument document."
NO_MARKER = "The template needs a paragraph containing {letter} where the letter goes."
PDF_FAILED = "LibreOffice could not turn the letter into a PDF."
TEMPLATE_GONE = "The cover letter template could not be opened. Upload it again on the Documents tab."


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------

def letter_paragraphs(letter: str) -> list[str]:
    """The letter's paragraphs (split at blank lines), each trimmed, with no empty ones; at least one."""
    text = clean((letter or "").replace("\r\n", "\n").replace("\r", "\n"))
    parts = [p.strip() for p in re.split(r"\n[ \t]*\n", text.strip())]
    return [p for p in parts if p] or [""]


def _values(values: dict) -> dict[str, str]:
    """The placeholders' values as single lines without characters XML can't hold."""
    out = {}
    for name in PLACEHOLDERS:
        out[name] = re.sub(r"\s+", " ", clean(str(values.get(name) or ""))).strip()
    return out


def _fill(text: str, values: dict[str, str]) -> str:
    """``text`` with each placeholder that has a value filled in; any other ``{name}`` is left as written."""
    return _PLACEHOLDER.sub(lambda m: values.get(m.group(1)) or m.group(0), text)


def _wants(text: str, values: dict[str, str]) -> bool:
    """Whether ``text`` still holds a placeholder that has a value (or a ``{letter}``)."""
    return any(values.get(name) or name == "letter" for name in _PLACEHOLDER.findall(text))


# ---------------------------------------------------------------------------
# DOCX (python-docx, working on the XML)
# ---------------------------------------------------------------------------

_W_TEXT = ("t", "tab", "br", "cr", "noBreakHyphen", "softHyphen", "ptab")


def _docx_open(data: bytes):
    import docx

    try:
        return docx.Document(io.BytesIO(data))
    except Exception:  # noqa: BLE001  (a zip, XML or package problem: the file isn't a Word document)
        raise TemplateError("unreadable") from None


def _docx_parts(document) -> list:
    """The root elements to fill placeholders in: the body, then every header and footer part (each once)."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    roots, seen = [document.element.body], set()
    for rel in list(document.part.rels.values()):
        if rel.is_external or rel.reltype not in (RT.HEADER, RT.FOOTER):
            continue
        part = rel.target_part
        if id(part) not in seen and hasattr(part, "element"):
            seen.add(id(part))
            roots.append(part.element)
    return roots


def _docx_paragraphs(root) -> list:
    from docx.oxml.ns import qn

    return list(root.iter(qn("w:p")))


def _docx_runs(p) -> list:
    """The paragraph's own runs (in hyperlinks, fields, content controls too), not those of a text box inside it."""
    from docx.oxml.ns import qn

    w_p = qn("w:p")
    return [r for r in p.iter(qn("w:r")) if next(r.iterancestors(w_p), None) is p]


def _docx_text(p) -> str:
    return "".join(r.text for r in _docx_runs(p))


def _is_text_child(el) -> bool:
    from docx.oxml.ns import qn

    if el.tag == qn("w:br"):
        return el.get(qn("w:type")) in (None, "textWrapping")  # a page or column break isn't text: it stays
    return el.tag in {qn(f"w:{name}") for name in _W_TEXT}


def _docx_clear_run(r) -> None:
    for child in list(r):
        if _is_text_child(child):
            r.remove(child)


def _docx_add_text(r, text: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for n, line in enumerate(text.split("\n")):
        if n:
            r.append(OxmlElement("w:br"))
        for m, piece in enumerate(line.split("\t")):
            if m:
                r.append(OxmlElement("w:tab"))
            if piece:
                t = OxmlElement("w:t")
                t.text = piece
                t.set(qn("xml:space"), "preserve")
                r.append(t)


def _docx_set_text(p, text: str) -> None:
    """The paragraph's text becomes ``text``, in its first run with text (keeping that run's formatting); the other
    runs lose their text (pictures, fields' codes and page breaks in them stay)."""
    from docx.oxml import OxmlElement

    runs = _docx_runs(p)
    first = next((r for r in runs if r.text), runs[0] if runs else None)
    if first is None:
        first = OxmlElement("w:r")
        p.append(first)
    for r in runs:
        _docx_clear_run(r)
    _docx_add_text(first, clean(text))


def _docx_fill_paragraph(p, values: dict[str, str]) -> None:
    """Placeholders filled (and a stray ``{letter}`` removed) in place where each sits in one piece of text; the
    paragraph is written again only when one is split across runs."""
    from docx.oxml.ns import qn

    for r in _docx_runs(p):
        for t in r.iter(qn("w:t")):
            if t.text and "{" in t.text:
                t.text = _fill(t.text, values).replace(MARKER, "")
    text = _docx_text(p)
    if _wants(text, values):
        _docx_set_text(p, _fill(text, values).replace(MARKER, ""))


def _docx_marker(document):
    for p in _docx_paragraphs(document.element.body):
        if MARKER in _docx_text(p):
            return p
    return None


def _docx_paragraph_copy(marker):
    """A new paragraph like ``marker``: its paragraph properties (without a section break) and its first text run's
    character formatting, no content."""
    from docx.oxml.ns import qn

    p = copy.deepcopy(marker)
    runs = _docx_runs(p)
    keep = next((r for r in runs if r.text), runs[0] if runs else None)
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    ppr = p.find(qn("w:pPr"))
    if ppr is not None:
        for sect in ppr.findall(qn("w:sectPr")):
            ppr.remove(sect)
    if keep is not None:
        for child in list(keep):
            if child.tag != qn("w:rPr"):
                keep.remove(child)
        p.append(keep)
    return p


def fill_docx(data: bytes, letter: str, values: dict) -> bytes:
    """The .docx template ``data`` with the placeholders and the letter filled in (see the module's docstring)."""
    from docx.oxml.ns import qn

    document = _docx_open(data)
    marker = _docx_marker(document)
    if marker is None:
        raise TemplateError("no_marker")
    values = _values(values)
    for root in _docx_parts(document):
        for p in _docx_paragraphs(root):
            if p is not marker:
                _docx_fill_paragraph(p, values)
    blank = _docx_paragraph_copy(marker)
    before, _, after = _fill(_docx_text(marker), values).partition(MARKER)
    after = after.replace(MARKER, "")
    paragraphs = letter_paragraphs(letter)
    texts = [before + paragraphs[0]] + paragraphs[1:]
    texts[-1] += after
    _docx_set_text(marker, texts[0])
    last = marker
    for text in texts[1:]:
        p = copy.deepcopy(blank)
        _docx_set_text(p, text)
        last.addnext(p)
        last = p
    ppr = marker.find(qn("w:pPr"))
    sect = ppr.find(qn("w:sectPr")) if ppr is not None else None
    if sect is not None and last is not marker:  # the section ends after the letter, not after its first paragraph
        last.get_or_add_pPr().append(sect)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def plain_docx(letter: str) -> bytes:
    """The letter as a plain Word document: one Normal paragraph per blank-line block."""
    import docx

    document = docx.Document()
    for para in letter_paragraphs(letter):
        document.add_paragraph(para)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# ODT (odfpy, working on its element tree)
# ---------------------------------------------------------------------------

def _odt_open(data: bytes):
    from odf.opendocument import load

    try:
        document = load(io.BytesIO(data))
    except Exception:  # noqa: BLE001  (a zip or XML problem: the file isn't an OpenDocument)
        raise TemplateError("unreadable") from None
    if getattr(document, "text", None) is None:  # a spreadsheet or a drawing: no text body
        raise TemplateError("unreadable")
    return document


def _odt_ns():
    from odf.namespaces import DRAWNS, OFFICENS, TEXTNS

    return DRAWNS, OFFICENS, TEXTNS


def _odt_skipped(el) -> bool:
    """Children whose text isn't the paragraph's own: frames (text boxes, pictures), notes and comments."""
    drawns, officens, textns = _odt_ns()
    ns, name = el.qname
    return ns == drawns or (ns == textns and name in ("note", "tracked-changes")) or \
        (ns == officens and name in ("annotation", "annotation-end"))


def _odt_paragraphs(root) -> list:
    """Every text:p and text:h under ``root``, in document order."""
    _drawns, _officens, textns = _odt_ns()
    out, stack = [], [root]
    while stack:
        el = stack.pop()
        if el.nodeType != el.ELEMENT_NODE:
            continue
        if el.qname in ((textns, "p"), (textns, "h")):
            out.append(el)
        stack.extend(reversed(el.childNodes))
    return out


def _odt_text_nodes(el) -> list:
    """The paragraph's own text nodes (in spans and links too), in order."""
    out = []
    for child in el.childNodes:
        if child.nodeType == child.TEXT_NODE:
            out.append(child)
        elif child.nodeType == child.ELEMENT_NODE and not _odt_skipped(child):
            out.extend(_odt_text_nodes(child))
    return out


def _odt_text(el) -> str:
    _drawns, _officens, textns = _odt_ns()
    out = []
    for child in el.childNodes:
        if child.nodeType == child.TEXT_NODE:
            out.append(child.data)
        elif child.nodeType == child.ELEMENT_NODE:
            if child.qname == (textns, "s"):
                try:
                    out.append(" " * max(1, int(child.getAttrNS(textns, "c") or 1)))
                except ValueError:
                    out.append(" ")
            elif child.qname == (textns, "tab"):
                out.append("\t")
            elif child.qname == (textns, "line-break"):
                out.append("\n")
            elif not _odt_skipped(child):
                out.append(_odt_text(child))
    return "".join(out)


def _odt_first_span_style(p) -> Optional[str]:
    """The style of the span holding the paragraph's first text, if it is in one (its "first run's formatting")."""
    _drawns, _officens, textns = _odt_ns()
    for node in _odt_text_nodes(p):
        if not node.data.strip():
            continue
        parent = node.parentNode
        while parent is not None and parent is not p:
            if parent.qname == (textns, "span") and parent.getAttrNS(textns, "style-name"):
                return parent.getAttrNS(textns, "style-name")
            parent = parent.parentNode
        return None
    return None


def _odt_detach_children(el) -> list:
    """Takes every child out of ``el`` (directly: odfpy's removeChild trips over text nodes in its caches)."""
    children = list(el.childNodes)
    for child in children:
        child.parentNode = child.previousSibling = child.nextSibling = None
    el.childNodes = []
    return children


def _odt_set_text(p, text: str, style: Optional[str] = None) -> None:
    """The paragraph's text becomes ``text`` (in a span of ``style`` when given); frames anchored in it stay."""
    from odf import teletype
    from odf.text import Span

    kept = [c for c in _odt_detach_children(p) if c.nodeType == c.ELEMENT_NODE and _odt_skipped(c)]
    for child in kept:
        p.appendChild(child)
    target = p
    if style:
        target = Span(stylename=style)
        p.appendChild(target)
    teletype.addTextToElement(target, clean(text))


def _odt_fill_paragraph(p, values: dict[str, str]) -> None:
    for node in _odt_text_nodes(p):
        if "{" in node.data:
            node.data = _fill(node.data, values).replace(MARKER, "")
    text = _odt_text(p)
    if _wants(text, values):
        _odt_set_text(p, _fill(text, values).replace(MARKER, ""), _odt_first_span_style(p))


def _odt_marker(document):
    for p in _odt_paragraphs(document.text):
        if MARKER in _odt_text(p):
            return p
    return None


def fill_odt(data: bytes, letter: str, values: dict) -> bytes:
    """The .odt template ``data`` with the placeholders and the letter filled in (see the module's docstring). Headers
    and footers are in the master pages, which odfpy reads into the same tree: they are filled too."""
    from odf.element import Element

    document = _odt_open(data)
    marker = _odt_marker(document)
    if marker is None:
        raise TemplateError("no_marker")
    values = _values(values)
    for p in _odt_paragraphs(document.topnode):
        if p is not marker:
            _odt_fill_paragraph(p, values)
    style = _odt_first_span_style(marker)
    attributes = {k: v for k, v in marker.attributes.items() if k[1] != "id"}  # (xml:id / text:id stay unique)
    before, _, after = _fill(_odt_text(marker), values).partition(MARKER)
    after = after.replace(MARKER, "")
    paragraphs = letter_paragraphs(letter)
    texts = [before + paragraphs[0]] + paragraphs[1:]
    texts[-1] += after
    _odt_set_text(marker, texts[0], style)
    parent, last = marker.parentNode, marker
    for text in texts[1:]:
        p = Element(qname=marker.qname, qattributes=dict(attributes), check_grammar=False)
        parent.insertBefore(p, last.nextSibling)
        _odt_set_text(p, text, style)
        last = p
    out = io.BytesIO()
    document.write(out)
    return out.getvalue()


def plain_odt(letter: str) -> bytes:
    """The letter as a plain OpenDocument text: one Standard paragraph per blank-line block."""
    from odf import teletype
    from odf.opendocument import OpenDocumentText
    from odf.style import Style
    from odf.text import P

    document = OpenDocumentText()
    standard = Style(name="Standard", family="paragraph")
    document.styles.addElement(standard)
    for para in letter_paragraphs(letter):
        p = P(stylename=standard)
        teletype.addTextToElement(p, para)
        document.text.addElement(p)
    out = io.BytesIO()
    document.write(out)
    return out.getvalue()


def check_template(data: bytes, suffix: str) -> None:
    """Raises TemplateError unless ``data`` opens as a ``suffix`` document with a ``{letter}`` body paragraph."""
    if suffix == ".docx":
        found = _docx_marker(_docx_open(data)) is not None
    elif suffix == ".odt":
        found = _odt_marker(_odt_open(data)) is not None
    else:
        raise TemplateError("unreadable")
    if not found:
        raise TemplateError("no_marker")


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_backend: Optional[tuple[str, Optional[str]]] = None  # ("libreoffice", path) | ("plain", None), found once


def _candidates() -> list[str]:
    found = [shutil.which(name) for name in ("soffice", "libreoffice")]
    found.append("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    for env, default in (("ProgramFiles", r"C:\Program Files"), ("ProgramFiles(x86)", r"C:\Program Files (x86)")):
        found.append(os.path.join(os.environ.get(env) or default, "LibreOffice", "program", "soffice.exe"))
    return [f for f in found if f]


def _find() -> tuple[str, Optional[str]]:
    global _backend
    if _backend is None:
        exe = next((c for c in _candidates() if os.path.isfile(c) and os.access(c, os.X_OK)), None)
        _backend = ("libreoffice", exe) if exe else ("plain", None)
    return _backend


def reset_pdf_backend() -> None:
    """Look for LibreOffice again on the next PDF (tests; an install made while the app runs needs a restart)."""
    global _backend
    _backend = None


def pdf_backend() -> str:
    """"libreoffice" when LibreOffice is installed (PDFs from the filled template), else "plain"."""
    return _find()[0]


def _temp_parent() -> Path:
    from backend.db import DATA_DIR

    return Path(DATA_DIR)


def remove_temp_folders(older_than_s: float = 0) -> int:
    """Removes the conversion folders (``.letter-pdf-*`` in the data folder) older than ``older_than_s`` seconds: the
    ones a crash left. Returns how many."""
    parent, removed = _temp_parent(), 0
    if not parent.is_dir():
        return 0
    now = time.time()
    for path in parent.glob(TEMP_PREFIX + "*"):
        try:
            if path.is_dir() and not path.is_symlink() and now - path.stat().st_mtime >= older_than_s:
                shutil.rmtree(path, ignore_errors=True)
                removed += 1
        except OSError:
            continue
    return removed


def to_pdf_libreoffice(data: bytes, suffix: str) -> bytes:
    """``data`` (a .docx or .odt) as a PDF, converted by LibreOffice in a folder of its own that is removed after."""
    exe = _find()[1]
    if not exe:
        raise PdfError("LibreOffice not found")
    remove_temp_folders(STALE_TEMP_S)
    parent = _temp_parent()
    parent.mkdir(mode=0o700, parents=True, exist_ok=True)
    tmp = tempfile.mkdtemp(prefix=TEMP_PREFIX, dir=parent)  # (0700)
    try:
        source = Path(tmp) / f"letter{suffix}"
        source.write_bytes(data)
        command = [exe, "--headless", "--norestore", "--nologo",
                   f"-env:UserInstallation={(Path(tmp) / 'profile').as_uri()}",
                   "--convert-to", "pdf", "--outdir", tmp, str(source)]
        try:
            done = subprocess.run(command, capture_output=True, timeout=PDF_TIMEOUT_S, stdin=subprocess.DEVNULL)
        except subprocess.TimeoutExpired:
            raise PdfError(f"timed out after {PDF_TIMEOUT_S} s") from None
        except OSError as exc:
            raise PdfError(f"could not start ({exc.__class__.__name__})") from None
        result = Path(tmp) / "letter.pdf"
        if done.returncode != 0:
            raise PdfError(f"exit status {done.returncode}")
        if not result.is_file():
            raise PdfError("no PDF written")
        return result.read_bytes()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


_NO_DECOMPOSITION = {"Ł": "L", "ł": "l", "Đ": "D", "đ": "d", "Ħ": "H", "ħ": "h", "ı": "i", "Ŧ": "T", "ŧ": "t"}


def cp1252_text(text: str) -> str:
    """``text`` in the characters the PDF core fonts have (Windows-1252): invisible characters dropped, accents
    taken off what has no accented form there (NFKD), anything else left a "?". Punctuation and braces stay."""
    out = []
    for ch in unicodedata.normalize("NFC", text):
        if ch in "\n\t":
            out.append(" " if ch == "\t" else ch)
            continue
        if unicodedata.category(ch) in ("Cc", "Cf", "Cs"):
            continue
        try:
            ch.encode("cp1252")
            out.append(ch)
            continue
        except UnicodeEncodeError:
            pass
        base = _NO_DECOMPOSITION.get(ch) or \
            "".join(c for c in unicodedata.normalize("NFKD", ch) if not unicodedata.combining(c))
        try:
            base.encode("cp1252")
            out.append(base)
        except UnicodeEncodeError:
            out.append("?")
    return "".join(out)


def to_pdf_plain(letter: str) -> bytes:
    """The letter as a plain-text PDF: A4, 25 mm margins, Helvetica 11 pt, a blank line between paragraphs."""
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF(orientation="portrait", unit="mm", format="A4")
    pdf.set_margins(25, 25, 25)
    pdf.set_auto_page_break(True, margin=25)
    pdf.core_fonts_encoding = "windows-1252"  # (the core fonts' WinAnsiEncoding; fpdf2 assumes Latin-1 otherwise)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    line = 5.5
    for n, para in enumerate(letter_paragraphs(letter)):
        if n:
            pdf.ln(line)
        pdf.multi_cell(0, line, cp1252_text(para), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    return bytes(pdf.output())


# ---------------------------------------------------------------------------
# Rendering a saved letter
# ---------------------------------------------------------------------------

def workspace_template(db: Session, ws: int) -> Optional[LetterTemplate]:
    return db.query(LetterTemplate).filter(LetterTemplate.workspace_id == ws).order_by(LetterTemplate.id.desc()).first()


def _filled(template: LetterTemplate, letter: str, values: dict) -> bytes:
    try:
        data = docs.read_upload(template.stored_name)
    except (OSError, crypto.CryptoError) as exc:
        log.warning("Cover letter template %s could not be opened: %s", template.id, exc.__class__.__name__)
        raise HTTPException(status_code=409, detail=TEMPLATE_GONE) from None
    try:
        return (fill_docx if template.filetype == ".docx" else fill_odt)(data, letter, values)
    except TemplateError as exc:
        log.warning("Cover letter template %s could not be filled: %s", template.id, exc.kind)
        raise HTTPException(status_code=409, detail=TEMPLATE_GONE) from None


def render(db: Session, ws: int, job, letter_text: str, fmt: str) -> tuple[bytes, str, str]:
    """The letter for ``job`` as ``fmt`` ("docx", "odt" or "pdf"): (the file, its media type, its extension)."""
    if fmt not in FORMATS:
        raise ValueError("unknown format")
    template = workspace_template(db, ws)
    values = placeholder_values(job)
    if fmt in ("docx", "odt"):
        if template is not None and template.filetype == f".{fmt}":
            data = _filled(template, letter_text, values)
        else:
            data = plain_docx(letter_text) if fmt == "docx" else plain_odt(letter_text)
        return data, MEDIA_TYPES[fmt], fmt
    if pdf_backend() == "libreoffice":
        if template is not None:
            source, suffix = _filled(template, letter_text, values), template.filetype
        else:
            source, suffix = plain_docx(letter_text), ".docx"
        try:
            data = to_pdf_libreoffice(source, suffix)
        except PdfError as exc:
            log.warning("LibreOffice could not turn a cover letter into a PDF: %s", exc)
            raise HTTPException(status_code=502, detail=PDF_FAILED) from None
    else:
        data = to_pdf_plain(letter_text)
    return data, MEDIA_TYPES["pdf"], "pdf"


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------

def _state(template: Optional[LetterTemplate]) -> dict:
    backend = pdf_backend()
    return {
        "template": None if template is None else {
            "filename": template.filename or "",
            "filetype": template.filetype,
            "created_at": template.created_at.isoformat() if template.created_at else None,
        },
        "pdf": backend,
        "pdf_uses_template": template is not None and backend == "libreoffice",
    }


@router.get("")
def get_template(db: Session = Depends(get_db), ws: int = Depends(current_workspace)):
    return _state(workspace_template(db, ws))


@router.post("")
def upload_template(file: UploadFile = File(...), db: Session = Depends(get_db), ws: int = Depends(current_workspace)):
    original = (file.filename or "template").strip()[:FILENAME_MAX] or "template"
    suffix = Path(original).suffix.lower()
    if suffix not in TEMPLATE_TYPES:
        raise HTTPException(status_code=400, detail="A cover letter template must be a .docx or .odt file.")
    limit = docs.MAX_UPLOAD_BYTES
    data = file.file.read(limit + 1)
    if len(data) > limit:
        raise HTTPException(status_code=413, detail=f"That file is over {limit // (1024 * 1024)} MB. "
                                                    "Upload a smaller copy (for example, without images).")
    try:
        check_template(data, suffix)
    except TemplateError as exc:
        raise HTTPException(status_code=400, detail=NO_MARKER if exc.kind == "no_marker" else UNREADABLE) from None
    old = db.query(LetterTemplate).filter(LetterTemplate.workspace_id == ws).all()
    old_files = [t.stored_name for t in old]
    stored = docs.save_upload(data)
    for t in old:
        db.delete(t)
    template = LetterTemplate(workspace_id=ws, filename=original, filetype=suffix, stored_name=stored,
                              created_at=datetime.utcnow())
    db.add(template)
    try:
        db.commit()
    except Exception:
        db.rollback()
        docs.delete_upload(stored)
        raise
    for name in old_files:  # (after the commit: a failed upload leaves the old template working)
        docs.delete_upload(name)
    return _state(template)


@router.delete("", status_code=204)
def delete_template(db: Session = Depends(get_db), ws: int = Depends(current_workspace)):
    rows = db.query(LetterTemplate).filter(LetterTemplate.workspace_id == ws).all()
    if not rows:
        raise HTTPException(status_code=404, detail="There is no cover letter template.")
    files = [t.stored_name for t in rows]
    for t in rows:
        db.delete(t)
    db.commit()
    for name in files:
        docs.delete_upload(name)
    return Response(status_code=204)
