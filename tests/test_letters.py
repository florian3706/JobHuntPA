"""Cover letters as files: letter templates (.docx and .odt), the plain layout, the PDF and the downloads.

The filler (backend/letter_files.py) is tested on documents built here with python-docx and odfpy; the PDF
(backend/letter_pdf.py) on this computer's fonts, with a made-up font name for the substitution. The endpoints:
signed in and the job's workspace (a signed-out call is refused, another workspace answers 404), 409 for a format
the template can't give. No network: the model is mocked at backend.scorer.call_model.

Run:  python3 -m unittest tests.test_letters -v
"""
from __future__ import annotations

import io
import json
import os
import tempfile
import unittest
from unittest import mock

_TMP = tempfile.mkdtemp()
os.environ.setdefault("JOBHUNT_DB_PATH", os.path.join(_TMP, "test.db"))  # never the real database
os.environ.setdefault("JOBHUNT_SCHEDULER", "off")
os.environ.setdefault("JOBHUNT_ENV_PATH", os.path.join(_TMP, "test.env"))

from tests import helpers  # noqa: E402  (signed-in clients; sets the test environment)

from fastapi.testclient import TestClient  # noqa: E402

import backend.scorer as sc  # noqa: E402
from backend import letter_files as lf  # noqa: E402
from backend import letter_pdf as lp  # noqa: E402
from backend.app import app  # noqa: E402
from backend.db import Document, Job, SessionLocal, Workspace, init_db  # noqa: E402

init_db()

LETTER = "Dear Hiring Manager,\n\nI led the ZQX launch.\nIt shipped on time.\n\nKind regards,\nJane Doe"
LLM_ENV = {"LLM_API_KEY": "k", "LLM_BASE_URL": "https://llm.example/v1", "LLM_MODEL": "m"}
MISSING_FONT = "No Such Font Zq"


def setUpModule():
    helpers.enter_admin_scope()


def tearDownModule():
    helpers.exit_admin_scope()


def docx_template(font: str = "Garamond", marker_runs=("{{letter}}",), header: str = "Jane Doe · {{date}}",
                  before: str = "{{company}}: {{job_title}}", after: str = "Signed, {{name}} {{unknown}}",
                  align_right: bool = False) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = header
    d.sections[0].footer.paragraphs[0].text = "Footer {{company}}"
    d.add_paragraph(before)
    p = d.add_paragraph()
    if align_right:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(14)
    for i, piece in enumerate(marker_runs):
        run = p.add_run(piece)
        run.font.name, run.font.size, run.italic = font, Pt(12.5), True
    d.add_paragraph(after)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def odt_template(marker: str = "{{letter}}", span: bool = True) -> bytes:
    from odf import teletype
    from odf.opendocument import OpenDocumentText
    from odf.style import ParagraphProperties, Style, TextProperties
    from odf.text import P, Span

    doc = OpenDocumentText()
    para = Style(name="Body", family="paragraph")
    para.addElement(ParagraphProperties(textalign="justify", marginbottom="0.4cm"))
    para.addElement(TextProperties(fontfamily="Caladea", fontsize="11.5pt"))
    doc.styles.addElement(para)
    bold = Style(name="Strong", family="text")
    bold.addElement(TextProperties(fontweight="bold"))
    doc.styles.addElement(bold)
    p = P(stylename=para)
    teletype.addTextToElement(p, "To {{company}} for the {{job_title}} role")
    doc.text.addElement(p)
    p = P(stylename=para)
    if span:
        holder = Span(stylename=bold)
        p.addElement(holder)
    else:
        holder = p
    teletype.addTextToElement(holder, marker)
    doc.text.addElement(p)
    p = P(stylename=para)
    teletype.addTextToElement(p, "{{name}}")
    doc.text.addElement(p)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


class FillerTest(unittest.TestCase):
    values = lf.marker_values(LETTER, "Acme Pty Ltd", "Program Manager")

    def test_docx_marker_paragraph_becomes_the_letter_in_its_formatting(self):
        model = lf.docx_letter(lf.fill_docx(docx_template(align_right=True), LETTER, self.values))
        texts = [p.text for p in model.body]
        self.assertEqual(texts, ["Acme Pty Ltd: Program Manager", "Dear Hiring Manager,",
                                 "I led the ZQX launch.\nIt shipped on time.", "Kind regards,\nJane Doe",
                                 "Signed, Jane Doe {{unknown}}"])
        for p in model.body[1:4]:  # each letter paragraph: the marker run's font, size and italic, its alignment
            self.assertEqual([(r.font, r.size, r.italic) for r in p.runs], [("Garamond", 12.5, True)])
            self.assertEqual((p.align, p.space_after), ("right", 14.0))
        self.assertEqual(model.body[0].align, "left")
        self.assertEqual([p.text for p in model.header], [f"Jane Doe · {lf.long_date()}"])
        self.assertEqual([p.text for p in model.footer], ["Footer Acme Pty Ltd"])
        self.assertIn("Garamond", model.families())

    def test_docx_marker_split_over_runs_and_a_second_marker(self):
        data = docx_template(marker_runs=("{{", "let", "ter}} {{letter}}"))
        model = lf.docx_letter(lf.fill_docx(data, "One paragraph only.", self.values))
        self.assertEqual([p.text for p in model.body], ["Acme Pty Ltd: Program Manager", "One paragraph only.",
                                                        "Signed, Jane Doe {{unknown}}"])

    def test_docx_without_the_marker_is_refused(self):
        with self.assertRaises(lf.TemplateError) as caught:
            lf.fill_docx(docx_template(marker_runs=("no marker here",)), LETTER, self.values)
        self.assertIn("{letter}", str(caught.exception))
        with self.assertRaises(lf.TemplateError):
            lf.check_template(b"not a zip", ".docx")
        with self.assertRaises(lf.TemplateError):
            lf.check_template(docx_template(), ".pdf")

    def test_docx_pictures_header_and_page_size_are_read(self):
        import docx
        from docx.shared import Inches

        d = docx.Document()
        d.sections[0].page_width, d.sections[0].page_height = Inches(8.5), Inches(11)
        d.add_paragraph("{{letter}}")
        from PIL import Image  # (reportlab brings Pillow)

        png_file = io.BytesIO()
        Image.new("RGB", (8, 8), "white").save(png_file, "PNG")
        png = png_file.getvalue()
        d.add_paragraph().add_run().add_picture(io.BytesIO(png), width=Inches(1))
        out = io.BytesIO()
        d.save(out)
        model = lf.docx_letter(lf.fill_docx(out.getvalue(), "Hello.", self.values))
        self.assertEqual((model.page_width, model.page_height), lf.US_LETTER)
        self.assertEqual(len(model.body[-1].pictures), 1)
        self.assertAlmostEqual(model.body[-1].pictures[0].width, 72.0, places=1)

    def test_odt_marker_paragraph_becomes_the_letter_in_its_formatting(self):
        model = lf.odt_letter(lf.fill_odt(odt_template(), LETTER, self.values))
        self.assertEqual([p.text for p in model.body], ["To Acme Pty Ltd for the Program Manager role",
                                                        "Dear Hiring Manager,", "I led the ZQX launch.\nIt shipped on time.",
                                                        "Kind regards,\nJane Doe", "Jane Doe"])
        for p in model.body[1:4]:  # the paragraph style (justified, Caladea 11.5) and the marker's span (bold)
            self.assertEqual((p.align, p.font, p.size), ("justify", "Caladea", 11.5))
            self.assertTrue(all(r.bold and r.font == "Caladea" for r in p.runs), p.runs)
            self.assertAlmostEqual(p.space_after, 0.4 * 72 / 2.54, places=2)
        self.assertFalse(model.body[0].runs[0].bold)
        self.assertEqual(model.families(), ["Caladea"])

    def test_odt_without_the_marker_is_refused_and_a_bare_marker_works(self):
        with self.assertRaises(lf.TemplateError):
            lf.fill_odt(odt_template(marker="nothing"), LETTER, self.values)
        model = lf.odt_letter(lf.fill_odt(odt_template(span=False), "Only one.", self.values))
        self.assertEqual([p.text for p in model.body][1], "Only one.")
        self.assertFalse(model.body[1].runs[0].bold)

    def test_plain_layouts_agree(self):
        for suffix in (".docx", ".odt"):
            model = lf.letter_model(lf.plain(suffix, LETTER, "Liberation Sans"), suffix)
            self.assertEqual([p.text for p in model.body], lf.split_paragraphs(LETTER), suffix)
            self.assertEqual((model.page_width, model.page_height), lf.A4, suffix)
            self.assertAlmostEqual(model.margin_left, 72.0, places=0)
            self.assertEqual(model.families(), ["Liberation Sans"], suffix)
            self.assertEqual({r.size for p in model.body for r in p.runs}, {11.0}, suffix)

    def test_helpers(self):
        self.assertEqual(lf.sign_off_name(LETTER), "Jane Doe")
        self.assertEqual(lf.sign_off_name("Dear team,\n\nA long last sentence that ends with a full stop."), "")
        self.assertEqual(lf.split_paragraphs("a\r\n\r\n\n b \n\nc"), ["a", "b", "c"])
        self.assertAlmostEqual(lf.parse_length("2.54cm"), 72.0, places=5)
        self.assertEqual(lf.parse_length("12pt"), 12.0)
        self.assertEqual(lf.parse_length("junk", 7.0), 7.0)
        self.assertEqual(lf.marker_font(lf.check_template(docx_template(), ".docx")), ("Garamond", 12.5))


class PdfTest(unittest.TestCase):
    def test_fonts_on_this_computer_and_stand_ins(self):
        cat = lp.catalogue()
        for family, files in list(cat.items())[:3]:
            self.assertTrue(all(p.is_file() for p in files.values()), family)
        used, files = lp.pick(MISSING_FONT)
        self.assertNotEqual(used, MISSING_FONT)
        self.assertEqual(lp.classify("Garamond"), "serif")
        self.assertEqual(lp.classify("Consolas"), "mono")
        self.assertEqual(lp.classify("Calibri"), "sans")
        self.assertIn(lp.pick("Garamond")[0], ("Garamond",) + lp.FALLBACKS["serif"][0] + ("Times-Roman",))
        self.assertEqual(lp.read_names(__file__), None)  # not a font
        report = lp.font_report([MISSING_FONT, lp.plain_font()])
        self.assertEqual(report[0]["found"], False)
        self.assertEqual(report[1]["found"] or report[1]["used"] == "Helvetica", True)

    def test_render_a_template_reports_substitutions(self):
        from pypdf import PdfReader

        data = lf.fill_docx(docx_template(font=MISSING_FONT), LETTER, lf.marker_values(LETTER, "Acme", "PM"))
        pdf, substitutions = lp.render(lf.docx_letter(data))
        self.assertTrue(pdf.startswith(b"%PDF-"))
        self.assertEqual([s["asked"] for s in substitutions if s["asked"] == MISSING_FONT], [MISSING_FONT])
        text = PdfReader(io.BytesIO(pdf)).pages[0].extract_text()
        for piece in ("Dear Hiring Manager,", "ZQX launch", "Jane Doe", "Footer Acme", lf.long_date()):
            self.assertIn(piece, text)

    def test_render_the_plain_odt(self):
        pdf, substitutions = lp.render(lf.odt_letter(lf.plain_odt(LETTER, lp.plain_font())))
        self.assertTrue(pdf.startswith(b"%PDF-"))
        self.assertEqual(substitutions, [])


class EndpointTest(unittest.TestCase):
    """Each test: its own workspace with one job and a saved letter."""

    def setUp(self):
        self.client = helpers.client()
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.ws = self.client.post("/api/workspaces", json={"name": self._testMethodName[:80]}).json()["id"]
        self.h = {"X-Workspace": str(self.ws)}
        self.addCleanup(lambda: self.client.delete(f"/api/workspaces/{self.ws}"))
        db = SessionLocal()
        job = Job(workspace_id=self.ws, url=f"https://letters.example/{self.ws}", title="Program Manager", company="Acme",
                  description="Lead programs.", status="to_review", detail_status="full")
        db.add(job)
        db.commit()
        self.job_id = job.id
        db.close()
        r = self.client.put(f"/api/jobs/{self.job_id}/cover-letter", headers=self.h, json={"text": LETTER})
        self.assertEqual(r.status_code, 200, r.text)

    def upload(self, name: str, data: bytes, kind: str = "letter_template"):
        return self.client.post("/api/documents", headers=self.h, files={"file": (name, data, "application/octet-stream")},
                                data={"kind": kind})

    def download(self, fmt: str):
        return self.client.get(f"/api/jobs/{self.job_id}/cover-letter.{fmt}", headers=self.h)

    def test_plain_layout_gives_every_format(self):
        info = self.client.get("/api/documents/letter-template", headers=self.h).json()
        self.assertIsNone(info["template"])
        self.assertEqual(info["pdf"], "reportlab")  # (tests/helpers.py: never this computer's LibreOffice)
        self.assertTrue(all(f["available"] for f in info["formats"].values()))
        self.assertEqual(list(info["formats"]), ["pdf", "odt", "docx"])
        for fmt, magic in (("pdf", b"%PDF-"), ("odt", b"PK"), ("docx", b"PK")):
            r = self.download(fmt)
            self.assertEqual(r.status_code, 200, fmt)
            self.assertTrue(r.content.startswith(magic), fmt)
            self.assertIn(f'filename="Cover letter - Acme - Program Manager.{fmt}"', r.headers["content-disposition"])
        self.assertEqual(json.loads(self.download("pdf").headers["x-font-substitutions"]), [])

    def test_the_plain_layouts_font_can_be_chosen(self):
        info = self.client.get("/api/documents/letter-template", headers=self.h).json()
        self.assertIsNone(info["font_setting"])
        self.assertIsInstance(info["font_choices"], list)
        default = info["letter_font"]
        r = self.client.put("/api/documents/letter-font", headers=self.h, json={"family": " Liberation Serif ", "size": 12})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["letter_font"], {"family": "Liberation Serif", "size": 12.0})
        self.assertEqual(r.json()["font_setting"], {"family": "Liberation Serif", "size": 12.0})
        body = lf.docx_letter(self.download("docx").content).body
        self.assertEqual({(p.font, p.size) for p in body if p.text.strip()}, {("Liberation Serif", 12.0)})
        body = lf.odt_letter(self.download("odt").content).body
        self.assertEqual({(p.font, p.size) for p in body if p.text.strip()}, {("Liberation Serif", 12.0)})
        self.assertEqual(self.download("pdf").status_code, 200)
        # the size alone, the family alone; bad values are refused; both null is the default again
        r = self.client.put("/api/documents/letter-font", headers=self.h, json={"size": 10.5})
        self.assertEqual(r.json()["letter_font"], {"family": default["family"], "size": 10.5})
        r = self.client.put("/api/documents/letter-font", headers=self.h, json={"family": "Georgia"})
        self.assertEqual(r.json()["letter_font"], {"family": "Georgia", "size": default["size"]})
        for bad in ({"size": 7}, {"size": 19}, {"family": "x" * 81}):
            self.assertEqual(self.client.put("/api/documents/letter-font", headers=self.h, json=bad).status_code, 422, bad)
        self.assertEqual(self.client.put("/api/documents/letter-font", headers=self.h, json={"family": "<b>Arial</b>"}).status_code, 400)
        r = self.client.put("/api/documents/letter-font", headers=self.h, json={"family": None, "size": None})
        self.assertEqual((r.json()["letter_font"], r.json()["font_setting"]), (default, None))
        # a template's letter keeps the template's font
        self.client.put("/api/documents/letter-font", headers=self.h, json={"family": "Georgia", "size": 9})
        self.upload("letterhead.docx", docx_template(font=MISSING_FONT))
        info = self.client.get("/api/documents/letter-template", headers=self.h).json()
        self.assertEqual(info["letter_font"], {"family": MISSING_FONT, "size": 12.5})
        self.assertEqual(info["font_setting"], {"family": "Georgia", "size": 9.0})
        # other workspaces: the default; anonymous: refused
        other = self.client.post("/api/workspaces", json={"name": "other"}).json()["id"]
        self.addCleanup(lambda: self.client.delete(f"/api/workspaces/{other}"))
        self.assertIsNone(self.client.get("/api/documents/letter-template", headers={"X-Workspace": str(other)}).json()["font_setting"])
        with TestClient(app) as anonymous:
            self.assertEqual(anonymous.put("/api/documents/letter-font", json={"size": 12}).status_code, 401)

    def test_docx_template_gives_docx_and_pdf_with_the_substitution_header(self):
        r = self.upload("letterhead.docx", docx_template(font=MISSING_FONT))
        self.assertEqual(r.status_code, 200, r.text)
        doc = r.json()
        self.assertEqual((doc["kind"], doc["use_for_scoring"], doc["current_template"]), ("letter_template", False, True))
        info = self.client.get("/api/documents/letter-template", headers=self.h).json()
        self.assertEqual(info["template"]["id"], doc["id"])
        self.assertEqual((info["formats"]["pdf"]["available"], info["formats"]["docx"]["available"],
                          info["formats"]["odt"]["available"]), (True, True, False))
        self.assertIn(".odt", info["formats"]["odt"]["reason"])
        self.assertEqual(info["letter_font"], {"family": MISSING_FONT, "size": 12.5})
        self.assertIn({"asked": MISSING_FONT, "used": lp.pick(MISSING_FONT)[0], "found": False}, info["fonts"])
        r = self.download("pdf")
        self.assertEqual(r.status_code, 200)
        subs = json.loads(r.headers["x-font-substitutions"])
        self.assertIn(MISSING_FONT, [s["asked"] for s in subs])
        r = self.download("docx")
        self.assertEqual(r.status_code, 200)
        self.assertIn("Dear Hiring Manager,", [p.text for p in lf.docx_letter(r.content).body])
        r = self.download("odt")
        self.assertEqual(r.status_code, 409)
        self.assertIn("Word (.docx) file", r.json()["detail"])

    def test_odt_template_gives_odt_and_pdf(self):
        self.assertEqual(self.upload("letterhead.odt", odt_template()).status_code, 200)
        self.assertEqual(self.download("odt").status_code, 200)
        self.assertEqual(self.download("pdf").status_code, 200)
        r = self.download("docx")
        self.assertEqual(r.status_code, 409)
        self.assertIn("OpenDocument (.odt) file", r.json()["detail"])

    def test_template_uploads_are_checked(self):
        r = self.upload("no-marker.docx", docx_template(marker_runs=("plain text",)))
        self.assertEqual(r.status_code, 400)
        self.assertIn("{letter}", r.json()["detail"])
        self.assertEqual(self.upload("cv.pdf", b"%PDF-1.4 junk").status_code, 400)  # a template is .docx or .odt
        self.assertEqual(self.upload("broken.docx", b"PK junk").status_code, 400)
        self.assertEqual(self.client.get("/api/documents/letter-template", headers=self.h).json()["template"], None)

    def test_choosing_among_templates_and_deleting_the_current_one(self):
        first = self.upload("first.docx", docx_template()).json()
        second = self.upload("second.odt", odt_template()).json()
        self.assertEqual((first["current_template"], second["current_template"]), (True, False))
        r = self.client.put("/api/documents/letter-template", headers=self.h, json={"id": second["id"]})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["template"]["filename"], "second.odt")
        listed = {d["id"]: d["current_template"] for d in self.client.get("/api/documents", headers=self.h).json()}
        self.assertEqual(listed, {first["id"]: False, second["id"]: True})
        cv = self.client.post("/api/documents/notes", headers=self.h, json={"title": "cv", "text": "Jane.", "kind": "resume"}).json()
        self.assertEqual(self.client.put("/api/documents/letter-template", headers=self.h, json={"id": cv["id"]}).status_code, 400)
        r = self.client.put("/api/documents/letter-template", headers=self.h, json={"id": None})
        self.assertIsNone(r.json()["template"])
        self.client.put("/api/documents/letter-template", headers=self.h, json={"id": second["id"]})
        self.assertEqual(self.client.delete(f"/api/documents/{second['id']}", headers=self.h).status_code, 200)
        self.assertIsNone(self.client.get("/api/documents/letter-template", headers=self.h).json()["template"])
        self.assertEqual(self.download("docx").status_code, 200)  # the plain layout again

    def test_a_template_is_never_sent_to_the_ai(self):
        from backend.cover_letters import build_prompt

        self.upload("letterhead.docx", docx_template(header="SECRET LETTERHEAD {{date}}"))
        tid = self.client.get("/api/documents", headers=self.h).json()[0]["id"]
        self.assertEqual(self.client.patch(f"/api/documents/{tid}", headers=self.h, json={"use_for_scoring": True}).status_code, 400)
        self.assertEqual(self.client.patch(f"/api/documents/{tid}", headers=self.h, json={"kind": "resume"}).status_code, 400)
        self.assertEqual(self.client.post(f"/api/documents/{tid}/concise", headers=self.h).status_code, 400)
        cv = self.client.post("/api/documents/notes", headers=self.h, json={"title": "cv", "text": "Jane led things.", "kind": "resume"}).json()
        self.assertEqual(self.client.patch(f"/api/documents/{cv['id']}", headers=self.h, json={"kind": "letter_template"}).status_code, 400)
        db = SessionLocal()
        try:
            db.get(Document, tid).use_for_scoring = True  # even forced on in the database
            db.commit()
            profile, _ = sc.build_profile(db, self.ws)
            prompt = build_prompt(db, db.get(Job, self.job_id))
        finally:
            db.close()
        self.assertNotIn("letterhead", profile.lower())
        self.assertNotIn("letterhead", prompt.lower())
        self.assertIn("Jane led things.", prompt)

    def test_draft_edits_and_discard(self):
        self.client.post("/api/documents/notes", headers=self.h, json={"title": "cv", "text": "Jane led things.", "kind": "resume"})
        with mock.patch.object(sc, "call_model", return_value="Dear Hiring Manager,\n\nDrafted.\n\nJane"), \
                mock.patch.dict(os.environ, LLM_ENV):
            r = self.client.post(f"/api/jobs/{self.job_id}/cover-letter", headers=self.h, json={})
        self.assertEqual(r.status_code, 200, r.text)
        letter = r.json()
        self.assertEqual((letter["draft_text"], letter["edited"]), (letter["text"], False))
        r = self.client.put(f"/api/jobs/{self.job_id}/cover-letter", headers=self.h, json={"text": "Edited by me."})
        self.assertEqual((r.json()["edited"], r.json()["draft_text"]), (True, letter["draft_text"]))
        r = self.client.put(f"/api/jobs/{self.job_id}/cover-letter", headers=self.h, json={"text": letter["draft_text"]})
        self.assertEqual((r.json()["edited"], r.json()["text"]), (False, letter["draft_text"]))  # back to the draft

    def test_access_is_the_job_workspace_signed_in(self):
        other = self.client.post("/api/workspaces", json={"name": "other"}).json()["id"]
        self.addCleanup(lambda: self.client.delete(f"/api/workspaces/{other}"))
        for fmt in ("pdf", "odt", "docx"):
            self.assertEqual(self.client.get(f"/api/jobs/{self.job_id}/cover-letter.{fmt}",
                                             headers={"X-Workspace": str(other)}).status_code, 404, fmt)
        with TestClient(app) as anonymous:
            for fmt in ("pdf", "odt", "docx"):
                self.assertEqual(anonymous.get(f"/api/jobs/{self.job_id}/cover-letter.{fmt}").status_code, 401, fmt)
            self.assertEqual(anonymous.get("/api/documents/letter-template").status_code, 401)
            self.assertEqual(anonymous.put("/api/documents/letter-template", json={"id": None}).status_code, 401)
        template = self.upload("letterhead.docx", docx_template()).json()
        r = self.client.put("/api/documents/letter-template", headers={"X-Workspace": str(other)}, json={"id": template["id"]})
        self.assertEqual(r.status_code, 404)  # another workspace's template can't be chosen
        self.assertIsNone(self.client.get("/api/documents/letter-template", headers={"X-Workspace": str(other)}).json()["template"])

    def test_a_copied_workspace_keeps_its_template_and_the_export_names_it(self):
        from backend import import_file

        template = self.upload("letterhead.docx", docx_template()).json()
        copy = self.client.post("/api/workspaces", json={"name": "copy", "copy_from": self.ws, "copy_documents": True}).json()["id"]
        self.addCleanup(lambda: self.client.delete(f"/api/workspaces/{copy}"))
        info = self.client.get("/api/documents/letter-template", headers={"X-Workspace": str(copy)}).json()
        self.assertEqual(info["template"]["filename"], "letterhead.docx")
        self.assertNotEqual(info["template"]["id"], template["id"])
        db = SessionLocal()
        try:
            data, _ = import_file.workspace_data(db, db.get(Workspace, self.ws), {})
        finally:
            db.close()
        docs = {d["ref"]: d for d in data["documents"]}
        self.assertEqual(docs[data["letter_template"]]["kind"], "letter_template")
        job = next(j for j in data["jobs"] if j["cover_letter"])
        self.assertEqual(job["cover_letter"]["text"], LETTER)
        import_file.WorkspaceIn.model_validate(data)  # what an import reads back


# --------------------------------------------------------------------------
# Markers in either brace style, the same values as cover letter blocks
# --------------------------------------------------------------------------

def single_brace_docx(marker=("{let", "ter}"), section_break: bool = False) -> bytes:
    """{company} in the header, {date} in the footer, {role} in a table cell, the {letter} paragraph (style "Body
    Text") split across two runs, and a paragraph with {location} and {{job_title}} in double braces."""
    import docx

    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Application to {company}"
    d.sections[0].footer.paragraphs[0].text = "Sent on {date}"
    d.add_paragraph("Dear Hiring Manager,")
    p = d.add_paragraph(style="Body Text")
    for piece in marker:
        p.add_run(piece)
    d.add_paragraph("In {location} as {{ job_title }}, {unknown} and {{company} stay")
    d.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run("Role: {role}")
    if section_break:  # the marker paragraph ends a section (a new one, landscape, follows)
        from docx.enum.section import WD_SECTION

        d.add_section(WD_SECTION.NEW_PAGE)
        body = d.element.body
        breaker = [el for el in body.iter() if el.tag.endswith("}sectPr") and el.getparent().tag.endswith("}pPr")][0]
        breaker.getparent().remove(breaker)
        p._p.get_or_add_pPr().append(breaker)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def single_brace_odt() -> bytes:
    """The same as an .odt: a master page with {company} in the header and {date} in the footer, the {letter}
    paragraph, and {role} in a table cell."""
    from odf.opendocument import OpenDocumentText
    from odf.style import Footer, Header, MasterPage, PageLayout, Style
    from odf.table import Table, TableCell, TableColumn, TableRow
    from odf.text import P

    doc = OpenDocumentText()
    body = Style(name="Letter Body", family="paragraph")
    doc.styles.addElement(body)
    layout = PageLayout(name="Page")
    doc.automaticstyles.addElement(layout)
    master = MasterPage(name="Standard", pagelayoutname=layout)
    header, footer = Header(), Footer()
    header.addElement(P(text="Application to {company}"))
    footer.addElement(P(text="Sent on {date}"))
    master.addElement(header)
    master.addElement(footer)
    doc.masterstyles.addElement(master)
    doc.text.addElement(P(stylename=body, text="{letter}"))
    table = Table(name="T")
    table.addElement(TableColumn())
    row, cell = TableRow(), TableCell()
    cell.addElement(P(text="Role: {role} in {{location}}"))
    row.addElement(cell)
    table.addElement(row)
    doc.text.addElement(table)
    out = io.BytesIO()
    doc.write(out)
    return out.getvalue()


def table_marker_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run("{letter}")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


class MarkerTest(unittest.TestCase):
    from datetime import date as _date

    values = lf.marker_values(LETTER, "Acme", "Program Manager", _date(2026, 10, 3), location="Sydney")

    def test_values_for_templates_and_blocks_agree(self):
        from backend.letter_blocks import PLACEHOLDERS, placeholder_values

        self.assertEqual(self.values, {"date": "3 October 2026", "company": "Acme", "job_title": "Program Manager",
                                       "role": "Program Manager", "location": "Sydney", "name": "Jane Doe"})
        job = Job(company=" Acme ", title="Program Manager", location_text="Sydney ")
        self.assertEqual(placeholder_values(job, self._date(2026, 10, 3)),
                         {name: self.values[name] for name in PLACEHOLDERS})
        self.assertEqual(set(PLACEHOLDERS) - set(lf.MARKER_NAMES), set(), "a block placeholder a template lacks")

    def test_blocks_take_either_brace_style(self):
        from backend.letter_blocks import assemble_letter, placeholders_in

        class B:
            text = "Dear {{company}} team, the {role} in {{ location }} ({{x}})"

        job = Job(company="Acme", title="PM", location_text="Sydney")
        text, unfilled = assemble_letter([B()], job)
        self.assertEqual(text, "Dear Acme team, the PM in {{ location }} ({{x}})")  # (no spaces inside, in blocks)
        self.assertEqual(unfilled, ["x"])
        self.assertEqual(placeholders_in("{{a}} {b} {{c}"), ["a", "b", "c"])

    def test_docx_single_braces(self):
        import docx

        d = docx.Document(io.BytesIO(lf.fill_docx(single_brace_docx(), LETTER, self.values)))
        body = [(p.style.name, p.text) for p in d.paragraphs]
        self.assertEqual(body[0], ("Normal", "Dear Hiring Manager,"))
        self.assertEqual(body[1:4], [("Body Text", t) for t in lf.split_paragraphs(LETTER)])
        self.assertEqual(body[4][1], "In Sydney as Program Manager, {unknown} and {Acme stay")
        self.assertEqual(d.sections[0].header.paragraphs[0].text, "Application to Acme")
        self.assertEqual(d.sections[0].footer.paragraphs[0].text, "Sent on 3 October 2026")
        self.assertEqual(d.tables[0].cell(0, 0).text, "Role: Program Manager")
        self.assertFalse(any("{letter}" in t for _, t in body))

    def test_odt_single_braces(self):
        from odf import teletype
        from odf.opendocument import load
        from odf.text import P

        doc = load(io.BytesIO(lf.fill_odt(single_brace_odt(), LETTER, self.values)))
        texts = [teletype.extractText(p) for root in (doc.masterstyles, doc.text) for p in root.getElementsByType(P)]
        self.assertEqual(texts[:2], ["Application to Acme", "Sent on 3 October 2026"])
        self.assertEqual(texts[2:5], lf.split_paragraphs(LETTER))
        self.assertEqual(texts[5], "Role: Program Manager in Sydney")

    def test_an_empty_letter_keeps_a_paragraph_in_the_cell(self):
        import docx
        from docx.oxml.ns import qn

        d = docx.Document(io.BytesIO(lf.fill_docx(table_marker_docx(), "", self.values)))
        cells = list(d.element.body.iter(qn("w:tc")))
        self.assertTrue(cells and all(tc.findall(qn("w:p")) for tc in cells), "Word refuses a cell without a paragraph")
        self.assertEqual(d.tables[0].cell(0, 0).text, "")
        from odf.opendocument import load
        from odf.table import TableCell
        from odf.text import P

        doc = load(io.BytesIO(lf.fill_odt(single_brace_odt(), "", self.values)))
        self.assertTrue(all(c.getElementsByType(P) for c in doc.text.getElementsByType(TableCell)))

    def test_a_section_break_moves_to_the_last_letter_paragraph(self):
        import docx
        from docx.oxml.ns import qn

        template = single_brace_docx(section_break=True)
        before = len(list(docx.Document(io.BytesIO(template)).element.body.iter(qn("w:sectPr"))))
        d = docx.Document(io.BytesIO(lf.fill_docx(template, LETTER, self.values)))
        self.assertEqual(len(list(d.element.body.iter(qn("w:sectPr")))), before, "no section break per paragraph")
        letter = d.paragraphs[1:4]
        self.assertEqual([p.text for p in letter], lf.split_paragraphs(LETTER))
        self.assertEqual([p._p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None for p in letter], [False, False, True])


class ReviewFixTest(unittest.TestCase):
    def setUp(self):
        self.client = helpers.client()
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.ws = self.client.post("/api/workspaces", json={"name": self._testMethodName[:80]}).json()["id"]
        self.h = {"X-Workspace": str(self.ws)}
        self.addCleanup(lambda: self.client.delete(f"/api/workspaces/{self.ws}"))

    def plant_template(self, data: bytes, filename="imported.docx") -> int:
        """A letter template document that never went through the upload check (an import brings them in so)."""
        from backend import docs

        db = SessionLocal()
        try:
            doc = Document(workspace_id=self.ws, filename=filename, filetype=".docx", kind="letter_template",
                           stored_name=docs.save_upload(data), text="", use_for_scoring=False)
            db.add(doc)
            db.flush()
            db.get(Workspace, self.ws).letter_template_id = doc.id
            db.commit()
            return doc.id
        finally:
            db.close()

    def test_a_template_that_no_longer_reads_says_why(self):
        import docx

        d = docx.Document()
        d.add_paragraph("no marker")
        out = io.BytesIO()
        d.save(out)
        self.plant_template(out.getvalue())
        r = self.client.get("/api/documents/letter-template", headers=self.h)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn("{letter}", r.json()["problem"])
        self.assertFalse(any(f["available"] for f in r.json()["formats"].values()))
        self.plant_template(b"PK not a zip", "broken.docx")
        r = self.client.get("/api/documents/letter-template", headers=self.h)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()["problem"])

    def test_the_document_list_reads_the_current_template_once(self):
        from sqlalchemy import event

        from backend.db import engine

        for n in range(3):
            self.assertEqual(self.client.post("/api/documents", headers=self.h, data={"kind": "letter_template"},
                                              files={"file": (f"t{n}.docx", docx_template(), "application/octet-stream")}
                                              ).status_code, 200)
        seen = []

        def count(_conn, _cursor, statement, *_args):
            if statement.startswith("SELECT workspaces.letter_template_id"):  # (_current_template_id)
                seen.append(statement)

        event.listen(engine, "before_cursor_execute", count)
        try:
            listed = self.client.get("/api/documents", headers=self.h).json()
        finally:
            event.remove(engine, "before_cursor_execute", count)
        self.assertEqual(len(seen), 1, seen)
        self.assertEqual(sum(d["current_template"] for d in listed), 1)

    def test_a_request_never_waits_for_the_font_scan(self):
        empty = tempfile.mkdtemp()
        with mock.patch.object(lp, "_catalogue", None), mock.patch.object(lp, "_warming", None), \
                mock.patch.object(lp, "font_dirs", return_value=[__import__("pathlib").Path(empty)]):
            with mock.patch.object(lp, "_scan", side_effect=lambda: __import__("time").sleep(0.3) or {}):
                self.assertEqual(lp.pick("Garamond"), ("Times-Roman", None), "the core font while the scan runs")
                self.assertEqual(lp.plain_font(), "Helvetica")
                warming = lp._warming
                self.assertIsNotNone(warming, "the scan was started in the background")
                lp.warm_up()
                self.assertIs(lp._warming, warming, "once")
                warming.result(timeout=10)
            self.assertEqual(lp._catalogue, {})


class RetiredTemplateTableTest(unittest.TestCase):
    def test_a_development_builds_template_table_is_dropped_with_its_files(self):
        """A build before this merge kept templates in a table of their own (``letter_templates``): start-up drops it
        (backend/db.py RETIRED_TABLES) and deletes the sealed files its rows named."""
        from sqlalchemy import inspect, text

        from backend import db as db_module

        stray = db_module.UPLOAD_DIR / "0123456789abcdef-retired.sealed"
        db_module.UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        stray.write_bytes(b"sealed template bytes")
        with db_module.engine.begin() as conn:
            conn.execute(text("CREATE TABLE letter_templates (id INTEGER PRIMARY KEY, workspace_id INTEGER, filename BLOB, "
                              "filetype VARCHAR NOT NULL, stored_name VARCHAR NOT NULL, created_at DATETIME)"))
            conn.execute(text("INSERT INTO letter_templates (workspace_id, filetype, stored_name) VALUES (1, '.docx', :n)"),
                         {"n": stray.name})
        self.assertIn("letter_templates", db_module.RETIRED_TABLES)
        self.assertTrue(db_module._retire_removed_features())
        self.assertNotIn("letter_templates", inspect(db_module.engine).get_table_names())
        self.assertFalse(stray.exists())
        self.assertFalse(db_module._retire_removed_features(), "nothing left to do the next time")


# --------------------------------------------------------------------------
# PDFs through LibreOffice (a fake soffice stands in)
# --------------------------------------------------------------------------

FAKE_SOFFICE = """#!{python}
import json, pathlib, sys
pathlib.Path({calls!r}).write_text(json.dumps(sys.argv[1:]))
args = sys.argv[1:]
if {fail!r}:
    sys.exit(1)
outdir = pathlib.Path(args[args.index("--outdir") + 1])
source = pathlib.Path(args[-1])
(outdir / (source.stem + ".pdf")).write_bytes(b"%PDF-1.4 fake from " + source.suffix.encode())
"""


@unittest.skipIf(os.name == "nt", "the fake soffice is a script with a #! line")
class LibreOfficeTest(unittest.TestCase):
    def setUp(self):
        import sys
        from pathlib import Path

        self.python = sys.executable
        self.dir = Path(tempfile.mkdtemp())
        self.data = self.dir / "data"
        self.data.mkdir()
        self.calls = self.dir / "calls.json"
        for patcher in (mock.patch.object(lp, "_temp_parent", return_value=self.data),
                        mock.patch.object(lp, "_backend", lp._backend)):  # (put back as tests/helpers.py set it)
            patcher.start()
            self.addCleanup(patcher.stop)

    def fake(self, fail=False) -> str:
        path = self.dir / ("fail-soffice" if fail else "soffice")
        path.write_text(FAKE_SOFFICE.format(python=self.python, calls=str(self.calls), fail=fail))
        path.chmod(0o755)
        lp._backend = ("libreoffice", str(path))
        return str(path)

    def test_the_conversion_and_its_folder(self):
        from pathlib import Path

        self.fake()
        self.assertEqual(lp.pdf_backend(), "libreoffice")
        self.assertEqual(lp.to_pdf_libreoffice(b"the filled template", ".odt"), b"%PDF-1.4 fake from .odt")
        self.assertEqual(list(self.data.iterdir()), [], "the conversion's folder is gone")
        args = json.loads(self.calls.read_text())
        self.assertEqual(args[:3], ["--headless", "--norestore", "--nologo"])
        self.assertTrue(args[3].startswith("-env:UserInstallation=file://"), args[3])
        self.assertEqual(args[4:6], ["--convert-to", "pdf"])
        self.assertTrue(args[-1].endswith("letter.odt"))
        self.assertTrue(Path(args[-1]).parent.name.startswith(lp.TEMP_PREFIX))

    def test_a_failed_conversion(self):
        self.fake(fail=True)
        with self.assertRaises(lp.PdfError) as caught:
            lp.to_pdf_libreoffice(b"x", ".docx")
        self.assertEqual(str(caught.exception), "exit status 1")
        self.assertEqual(list(self.data.iterdir()), [])

    def test_stale_folders_are_removed(self):
        self.fake()
        stale = self.data / f"{lp.TEMP_PREFIX}old"
        stale.mkdir()
        (stale / "letter.docx").write_bytes(b"left by a crash")
        os.utime(stale, (1, 1))
        fresh = self.data / f"{lp.TEMP_PREFIX}running"
        fresh.mkdir()
        lp.to_pdf_libreoffice(b"x", ".docx")
        self.assertFalse(stale.exists())
        self.assertTrue(fresh.exists(), "a conversion running now is left alone")
        self.assertEqual(lp.remove_temp_folders(), 1)
        self.assertFalse(fresh.exists())

    def test_over_http(self):
        with helpers.client() as c:
            ws = c.post("/api/workspaces", json={"name": "LibreOffice"}).json()["id"]
            h = {"X-Workspace": str(ws)}
            try:
                db = SessionLocal()
                job = Job(workspace_id=ws, url=f"https://letters.example/lo/{ws}", title="PM", company="Acme",
                          status="to_review", detail_status="full")
                db.add(job)
                db.commit()
                job_id = job.id
                db.close()
                self.assertEqual(c.put(f"/api/jobs/{job_id}/cover-letter", headers=h, json={"text": LETTER}).status_code, 200)
                self.fake()
                self.assertEqual(c.get("/api/documents/letter-template", headers=h).json()["pdf"], "libreoffice")
                r = c.get(f"/api/jobs/{job_id}/cover-letter.pdf", headers=h)  # the plain .docx goes to LibreOffice
                self.assertEqual((r.status_code, r.content), (200, b"%PDF-1.4 fake from .docx"))
                self.assertEqual(json.loads(r.headers["x-font-substitutions"]), [])
                r = c.post("/api/documents", headers=h, data={"kind": "letter_template"},
                           files={"file": ("t.odt", odt_template(), "application/octet-stream")})
                self.assertEqual(r.status_code, 200, r.text)
                r = c.get(f"/api/jobs/{job_id}/cover-letter.pdf", headers=h)  # the filled .odt
                self.assertEqual((r.status_code, r.content), (200, b"%PDF-1.4 fake from .odt"))
                self.fake(fail=True)
                with self.assertLogs("backend.cover_letters", "WARNING") as logged:
                    r = c.get(f"/api/jobs/{job_id}/cover-letter.pdf", headers=h)
                self.assertEqual(r.status_code, 200, "reportlab draws it instead")
                self.assertTrue(r.content.startswith(b"%PDF-") and b"fake" not in r.content)
                self.assertEqual(logged.output, ["WARNING:backend.cover_letters:LibreOffice could not turn a cover "
                                                 "letter into a PDF: exit status 1"])
                self.assertEqual(list(self.data.iterdir()), [])
            finally:
                c.delete(f"/api/workspaces/{ws}")

    def test_finding_libreoffice(self):
        lp.reset_pdf_backend()
        with mock.patch.object(lp, "_candidates", return_value=[str(self.dir / "missing")]):
            self.assertEqual(lp.pdf_backend(), "reportlab")
        lp.reset_pdf_backend()
        exe = self.dir / "soffice-found"
        exe.write_text("#!/bin/sh\n")
        exe.chmod(0o755)
        with mock.patch.object(lp, "_candidates", return_value=[str(self.dir / "missing"), str(exe)]):
            self.assertEqual(lp._find(), ("libreoffice", str(exe)))
            with mock.patch.object(lp, "_candidates", return_value=[]):
                self.assertEqual(lp.pdf_backend(), "libreoffice", "found once, for the process")


if __name__ == "__main__":
    unittest.main()
