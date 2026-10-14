"""Export file types and the import (backend/export_formats.py, backend/import_file.py, backend/data_import.py;
docs/design/open-source.md §1.7 and Part 7)."""
from __future__ import annotations

import io
import json
import os
import unittest
import zipfile
from datetime import date, datetime, timedelta
from unittest import mock

from tests import helpers
from tests import test_export

from backend import data_import, docs, import_file
from backend.db import (
    CompanyProfile,
    CompanySource,
    Document,
    Job,
    Posting,
    SessionLocal,
    Workspace,
    engine,
    utcnow,
)


def zip_of(files: dict[str, bytes | str], *, method=zipfile.ZIP_DEFLATED) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", method) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return out.getvalue()


def entries(blob: bytes) -> dict[str, bytes]:
    z = zipfile.ZipFile(io.BytesIO(blob))
    return {n: z.read(n) for n in z.namelist()}


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        helpers.started()

    def setUp(self):
        self.bob = helpers.make_user(f"MARK-name-bob-{os.urandom(4).hex()}")
        self.upload = b"%PDF-1.4 MARK-upload-bytes \x00\xff\x10"
        with helpers.as_account(self.bob):
            test_export.ExportTest.plant(self, self.bob, "bob")
        self.addCleanup(self.restore_research, self.research_now())

    def add_posting(self, **fields) -> None:
        """A posting in the shared pool for this test (taken out again after it)."""
        db = SessionLocal()
        try:
            db.add(Posting(detail_status="full", **fields))
            db.commit()
        finally:
            db.close()
        self.addCleanup(self.drop_posting, fields["url"])

    @staticmethod
    def drop_posting(url: str) -> None:
        db = SessionLocal()
        try:
            db.query(Posting).filter(Posting.url == url).delete()
            db.commit()
        finally:
            db.close()

    @staticmethod
    def research_now() -> dict:
        """The shared company research as it is (an import may change it; other tests read it)."""
        db = SessionLocal()
        try:
            columns = [c.name for c in CompanyProfile.__table__.columns]
            return {p.key: {c: getattr(p, c) for c in columns} for p in db.query(CompanyProfile)}
        finally:
            db.close()

    @staticmethod
    def restore_research(saved: dict) -> None:
        db = SessionLocal()
        try:
            for p in db.query(CompanyProfile):
                if p.key not in saved:
                    db.delete(p)
                else:
                    for column, value in saved[p.key].items():
                        setattr(p, column, value)
            db.commit()
        finally:
            db.close()

    def export(self, who=None, formats=None) -> bytes:
        with helpers.client((who or self.bob).email) as c:
            r = c.post("/api/account/export", json={"formats": formats} if formats else None)
        self.assertEqual(r.status_code, 200, r.text[:300])
        return r.content

    def post(self, path, blob, who=None, **params):
        with helpers.client((who or self.bob).email) as c:
            return c.post(path, content=blob, params=params, headers={"Content-Type": "application/zip"})

    def check(self, blob, who=None):
        return self.post("/api/account/import/check", blob, who)

    def do_import(self, blob, pick="0", who=None):
        return self.post("/api/account/import", blob, who, pick=pick)

    def workspaces(self, who=None):
        who = who or self.bob
        db = SessionLocal()
        try:
            with helpers.as_account(who):
                return [(w.id, w.name) for w in db.query(Workspace).filter(Workspace.owner_id == who.id).order_by(Workspace.id)]
        finally:
            db.close()

    def snapshot(self, ws_id, who=None):
        """The workspace as the import file describes it (without its name and the original files' paths)."""
        db = SessionLocal()
        try:
            with helpers.as_account(who or self.bob):
                data, _ = import_file.workspace_data(db, db.get(Workspace, ws_id), {})
        finally:
            db.close()
        data.pop("name")
        data.pop("created_at")
        return data

    def expected(self):
        """Bob's workspace as an import should bring it back: the same, except that a page a chat answer used comes
        back only when its address is a web address (the planted "javascript:" one doesn't)."""
        data = self.snapshot(self.bob.workspace_id)
        for chat in data["chats"]:
            for m in chat["messages"]:
                if m["sources"] is not None:
                    m["sources"] = [p for p in m["sources"] if str(p.get("url", "")).startswith("https://")]
        return data

    def import_data(self, blob) -> dict:
        name = next(n for n in zipfile.ZipFile(io.BytesIO(blob)).namelist() if n.endswith(import_file.FILE_NAME))
        return json.loads(zipfile.ZipFile(io.BytesIO(blob)).read(name))

    def replace_import_file(self, blob, data) -> bytes:
        files = entries(blob)
        name = next(n for n in files if n.endswith(import_file.FILE_NAME))
        files[name] = json.dumps(data).encode()
        return zip_of(files)


class ExportFileTypesTest(Base):
    def test_markdown_docx_and_odt_each_hold_the_jobs_text_and_open_with_their_libraries(self):
        import docx
        from odf import teletype
        from odf import text as odf_text
        from odf.opendocument import load

        files = entries(self.export(formats=["md", "docx", "odt"]))
        job = "Workspaces/My job search/Jobs/Shortlisted/Acme Corp - Senior Analyst MARK-jobtitle-bob"
        root = next(iter(files)).split("/")[0]
        for ext in (".md", ".docx", ".odt"):
            self.assertIn(f"{root}/{job}{ext}", files)
            self.assertIn(f"{root}/README{ext}", files)
        document = docx.Document(io.BytesIO(files[f"{root}/{job}.docx"]))
        words = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("MARK-description-bob", words)
        self.assertIn("<script>alert('x')</script>", words, "the ad's text is text in a Word file, unescaped")
        self.assertIn("Senior Analyst MARK-jobtitle-bob", words)
        odt = load(io.BytesIO(files[f"{root}/{job}.odt"]))
        words = "\n".join(teletype.extractText(p) for p in odt.getElementsByType(odf_text.P) + odt.getElementsByType(odf_text.H))
        self.assertIn("MARK-description-bob", words)
        self.assertIn("MARK-fitsummary-bob", words)
        links = [a.getAttribute("href") for a in odt.getElementsByType(odf_text.A)]
        self.assertTrue(any(h.endswith("Cover%20letter.odt") for h in links), links)
        self.assertFalse(any(h.startswith("javascript") for h in links))
        # the originals are there once, and so is the import file
        self.assertEqual(sum(1 for n in files if "/Original files/" in n), 1)
        self.assertIn(f"{root}/{import_file.FILE_NAME}", files)

    def test_only_the_chosen_types_are_written(self):
        names = list(entries(self.export(formats=["docx"])))
        self.assertTrue(any(n.endswith(".docx") for n in names))
        self.assertFalse(any(n.endswith((".md", ".odt")) for n in names if "/Original files/" not in n))
        default = list(entries(self.export()))
        self.assertFalse(any(n.endswith((".docx", ".odt")) for n in default if "/Original files/" not in n))

    def test_word_and_opendocument_files_are_made_in_memory(self):
        import tempfile
        data_dir = os.environ["JOBHUNT_DATA_DIR"]
        before = {os.path.join(d, f) for d, _s, files in os.walk(data_dir) for f in files}
        with mock.patch.object(tempfile, "SpooledTemporaryFile", side_effect=AssertionError("no spooled files")), \
                mock.patch.object(tempfile, "NamedTemporaryFile", side_effect=AssertionError("no temporary files")), \
                mock.patch.object(tempfile, "TemporaryFile", side_effect=AssertionError("no temporary files")), \
                mock.patch.object(tempfile, "mkstemp", side_effect=AssertionError("no temporary files")):
            self.export(formats=["docx", "odt"])
        after = {os.path.join(d, f) for d, _s, files in os.walk(data_dir) for f in files}
        self.assertEqual({p for p in after - before if not p.endswith(("-wal", "-shm", ".db-journal"))}, set())
        from backend import export_formats
        with open(export_formats.__file__, encoding="utf-8") as f:
            src = f.read()
        for forbidden in ("tempfile", "open(", "write_bytes", "write_text", "mkdtemp"):
            self.assertNotIn(forbidden, src)

    def test_the_converter_keeps_other_peoples_text_as_text(self):
        from backend import export, export_formats
        hostile = "Apply <b>now</b> & [here](javascript:alert(1)) \\[x] a | b \x00\x07 end"
        md = f"# {export.line('Title')}\n\n{export.text(hostile)}\n\n- {export.link('Ad', 'https://x.example/a')}\n"
        blocks = export_formats.parse(md, ".docx")
        para = next(b for b in blocks if b.kind == "para")
        self.assertEqual("".join(r.text for r in para.runs), "Apply <b>now</b> & [here](javascript:alert(1)) \\[x] a | b  end")
        self.assertTrue(all(r.href is None for r in para.runs), "no link forms from the text")
        bullet = next(b for b in blocks if b.kind == "bullet")
        self.assertEqual([r.href for r in bullet.runs if r.href], ["https://x.example/a"])
        self.assertEqual(export_formats._target("../Jobs/A%20B.md", ".odt"), "../Jobs/A%20B.odt")
        for bad in ("javascript:alert(1)", "file:///etc/passwd", "/etc/passwd", "data:text/html,x"):
            self.assertIsNone(export_formats._target(bad, ".docx"), bad)
        for fmt in ("docx", "odt"):
            self.assertTrue(export_formats.convert(md, fmt).startswith(b"PK"), fmt)

    def test_at_least_one_known_type(self):
        with helpers.client(self.bob.email) as c:
            for bad in ([], ["pdf"], ["md", "md", "md", "md"]):
                self.assertEqual(c.post("/api/account/export", json={"formats": bad}).status_code, 422, bad)

    def test_long_names_keep_room_for_the_longest_ending(self):
        files = entries(self.export(formats=["md", "docx", "odt"]))
        for name in files:
            for part in name.split("/"):
                self.assertLessEqual(len(part), 80, part)
                self.assertLessEqual(len(part.encode()), 255, part)

    def test_the_import_file_holds_no_secrets(self):
        from backend.db import UserKeys
        data = self.import_data(self.export())
        self.assertEqual((data["format"], data["format_version"]), ("jobhuntpa-import", import_file.FORMAT_VERSION))
        text = json.dumps(data)
        db = SessionLocal()
        try:
            keys = db.get(UserKeys, self.bob.id)
            for secret in (keys.pw_wrapped_dek, keys.x25519_pub):
                self.assertNotIn(secret.hex(), text)
        finally:
            db.close()
        for word in ("password", "csrf", "api_key", "LLM_API_KEY", "wrapped", "session"):
            self.assertNotIn(word.lower(), text.lower(), word)
        self.assertNotIn(helpers.USER_PASSWORD, text)


class RoundTripTest(Base):
    def test_export_then_import_gives_the_same_workspace(self):
        before = self.workspaces()
        blob = self.export(formats=["md", "docx"])
        r = self.check(blob)
        self.assertEqual(r.status_code, 200, r.text)
        listed = r.json()["workspaces"]
        self.assertEqual([w["name"] for w in listed], ["My job search"])
        self.assertEqual(listed[0]["jobs"], 7)  # (the merged copy isn't counted)
        self.assertEqual((listed[0]["documents"], listed[0]["chats"], listed[0]["cover_letters"]), (3, 1, 1))
        r = self.do_import(blob)
        self.assertEqual(r.status_code, 200, r.text)
        made = r.json()["workspaces"][0]
        today = date.today().isoformat()
        self.assertEqual(made["name"], f"My job search (imported {today})")
        self.assertEqual(self.workspaces()[:-1], before, "nothing was merged into the old workspace")
        self.assertEqual(self.snapshot(made["id"]), self.expected())
        # the original file came back byte for byte, sealed under a new name
        db = SessionLocal()
        try:
            with helpers.as_account(self.bob):
                old = [d for d in db.query(Document).filter(Document.workspace_id == self.bob.workspace_id) if d.stored_name]
                new = [d for d in db.query(Document).filter(Document.workspace_id == made["id"]) if d.stored_name]
                self.assertEqual(len(new), 1)
                self.assertNotEqual(new[0].stored_name, old[0].stored_name)
                self.assertEqual(docs.read_upload(new[0].stored_name), self.upload)
        finally:
            db.close()

    def test_starting_points_and_the_applied_date_come_back(self):
        """Phase C's starting points (a pin start names the copy of its pin; a gone pin's start is left out) and
        Phase D's applied date travel in the import file."""
        from backend import commute
        from backend.db import Pin, UserProfile

        when = datetime(2026, 10, 3, 9, 30)
        db = SessionLocal()
        try:
            with helpers.as_account(self.bob):
                ws = self.bob.workspace_id
                pin = db.query(Pin).filter(Pin.workspace_id == ws).one()
                db.get(UserProfile, ws).starting_points = [
                    {"label": "MARK-start-pin", "pin_id": pin.id, "place": None, "use": "driving"},
                    {"label": "MARK-start-place", "pin_id": None, "place": "MARK-place Hornsby Station", "use": "transit"},
                    {"label": "MARK-start-gone", "pin_id": pin.id + 100000, "place": None, "use": "both"},
                ]
                weird = next(j for j in db.query(Job).filter(Job.workspace_id == ws) if j.url == "https://ads.example/bob/weird")
                weird.applied_at = when
                db.commit()
        finally:
            db.close()
        blob = self.export()
        setup = self.import_data(blob)["workspaces"][0]["search_setup"]
        self.assertEqual(setup["starting_points"], [
            {"label": "MARK-start-pin", "pin": 1, "place": None, "use": "driving"},
            {"label": "MARK-start-place", "pin": None, "place": "MARK-place Hornsby Station", "use": "transit"},
        ])
        r = self.do_import(blob)
        self.assertEqual(r.status_code, 200, r.text)
        made = r.json()["workspaces"][0]["id"]
        db = SessionLocal()
        try:
            with helpers.as_account(self.bob):
                new_pin = db.query(Pin).filter(Pin.workspace_id == made).one()
                profile = db.get(UserProfile, made)
                self.assertEqual(commute.saved_starting_points(profile), [
                    {"label": "MARK-start-pin", "pin_id": new_pin.id, "place": None, "use": "driving"},
                    {"label": "MARK-start-place", "pin_id": None, "place": "MARK-place Hornsby Station", "use": "transit"},
                ])
                self.assertIsNone(profile.commute_from)
                job = next(j for j in db.query(Job).filter(Job.workspace_id == made) if j.url == "https://ads.example/bob/weird")
                self.assertEqual(job.applied_at, when)
        finally:
            db.close()
        self.assertEqual(self.snapshot(made), self.expected())

    def test_an_older_file_s_stations_and_a_missing_schedule(self):
        """The first import files wrote the "Commute from" stations as text: each comes in as a public-transport place.
        A file with no schedule gives a new workspace's schedule: off."""
        from backend import commute, schedule
        from backend.db import UserProfile

        blob = self.export()
        data = self.import_data(blob)
        data["workspaces"][0]["search_setup"]["starting_points"] = ["MARK-old-station", {"label": "x", "pin": 99}]
        data["workspaces"][0]["schedule"] = None
        r = self.do_import(self.replace_import_file(blob, data))
        self.assertEqual(r.status_code, 200, r.text)
        made = r.json()["workspaces"][0]["id"]
        db = SessionLocal()
        try:
            with helpers.as_account(self.bob):
                self.assertEqual(commute.saved_starting_points(db.get(UserProfile, made)), [
                    {"label": "MARK-old-station", "pin_id": None, "place": "MARK-old-station", "use": "transit"}])
                self.assertFalse(schedule.get_schedule(db, made).enabled)
        finally:
            db.close()

    def test_into_another_account(self):
        alice = helpers.make_user(f"MARK-name-alice-{os.urandom(4).hex()}")
        with helpers.as_account(alice):
            db = SessionLocal()
            ws = db.get(Workspace, alice.workspace_id)
            ws.name = "Alice's search"
            db.commit()
            db.close()
        r = self.do_import(self.export(), who=alice)
        self.assertEqual(r.status_code, 200, r.text)
        made = r.json()["workspaces"][0]
        self.assertEqual(made["name"], "My job search", "a free name is kept as it is")
        self.assertEqual(self.snapshot(made["id"], alice), self.expected())

    def test_the_name_counter(self):
        blob = self.export()
        today = date.today().isoformat()
        names = [self.do_import(blob).json()["workspaces"][0]["name"] for _ in range(3)]
        self.assertEqual(names, [f"My job search (imported {today})", f"My job search (imported {today}, 2)",
                                 f"My job search (imported {today}, 3)"])
        taken = {"a", "a (imported 2026-10-13)"}
        self.assertEqual(data_import.new_name("A", taken, date(2026, 10, 13)), "A (imported 2026-10-13, 2)")
        self.assertEqual(data_import.new_name("  B  ", taken, date(2026, 10, 13)), "B")
        long = data_import.new_name("x" * 200, {"x" * 80}, date(2026, 10, 13))
        self.assertLessEqual(len(long), 80)
        self.assertTrue(long.endswith(" (imported 2026-10-13)"))

    def test_picking_some_workspaces(self):
        with helpers.as_account(self.bob):
            helpers.new_workspace("Second search", owner=self.bob.id)
        blob = self.export()
        self.assertEqual([w["name"] for w in self.check(blob).json()["workspaces"]], ["My job search", "Second search"])
        r = self.do_import(blob, pick="1")
        self.assertEqual([w["name"] for w in r.json()["workspaces"]],
                         [f"Second search (imported {date.today().isoformat()})"])
        for bad in ("", "2", "0,0", "x", "-1"):
            r = self.do_import(blob, pick=bad)
            self.assertEqual((r.status_code, r.json().get("code")), (422, "import_pick"), bad)

    def test_pool_linking_and_never_writing_to_the_pool(self):
        self.add_posting(url="https://ads.example/bob/main", title="Pool title", company="Acme Corp",
                         description="the pool's ad")
        db = SessionLocal()
        try:
            postings = db.query(Posting).count()
        finally:
            db.close()
        r = self.do_import(self.export())
        ws = r.json()["workspaces"][0]["id"]
        db = SessionLocal()
        try:
            self.assertEqual(db.query(Posting).count(), postings, "import never adds to the pool")
            with helpers.as_account(self.bob):
                jobs = {j.url: j for j in db.query(Job).filter(Job.workspace_id == ws)}
                main = jobs["https://ads.example/bob/main"]
                self.assertIsNotNone(main.posting_id)
                self.assertEqual(main.origin, "pool")
                # the exported job's values (its own, then: a private job's) sit on top of the pool's
                self.assertIn("MARK-description-bob", main.description)
                self.assertEqual(main.status, "shortlisted")
                other = jobs["https://ads.example/bob/hidden"]
                self.assertIsNone(other.posting_id)
                self.assertEqual(other.origin, "private")
                pool = db.query(Posting).filter(Posting.url == "https://ads.example/bob/main").one()
                self.assertEqual(pool.description, "the pool's ad", "the pool's own ad is untouched")
        finally:
            db.close()

    def test_only_the_workspace_s_own_values_go_on_top_of_a_pool_posting(self):
        self.add_posting(url="https://ads.example/bob/same0", title="Same", company="Same",
                         description="pool text for same0")
        blob = self.export()
        data = self.import_data(blob)
        job = next(j for j in data["workspaces"][0]["jobs"] if j["ad"]["url"] == "https://ads.example/bob/same0")
        job["origin"], job["own"] = "pool", ["title"]
        job["ad"]["title"], job["ad"]["description"] = "My own title", "an old copy of the ad"
        ws = self.do_import(self.replace_import_file(blob, data)).json()["workspaces"][0]["id"]
        db = SessionLocal()
        try:
            with helpers.as_account(self.bob):
                linked = next(j for j in db.query(Job).filter(Job.workspace_id == ws) if j.url == "https://ads.example/bob/same0")
                self.assertEqual((linked.title, linked.description), ("My own title", "pool text for same0"))
        finally:
            db.close()

    def test_company_research_only_where_the_install_has_none_or_older(self):
        blob = self.export()
        data = self.import_data(blob)
        new_time = (utcnow() - timedelta(days=1)).replace(microsecond=0)
        data["companies"] = [
            {"key": "acme corp", "name": "Acme Corp", "business_model": "MARK-newer-research", "researched_at": new_time.isoformat() + "Z"},
            {"key": "hidco", "name": "Hidco", "business_model": "MARK-hidco-research", "researched_at": new_time.isoformat() + "Z"},
            {"key": "filtco", "name": "Filtco", "business_model": "MARK-older-filtco", "researched_at": "2020-01-01T00:00:00Z"},
            {"key": "elsewhere inc", "name": "Elsewhere Inc", "business_model": "not a company of these jobs"},
        ]
        db = SessionLocal()
        try:
            db.merge(CompanyProfile(key="filtco", name="Filtco", status="done", business_model="MARK-current-filtco",
                                    researched_at=datetime(2025, 1, 1)))
            acme = db.get(CompanyProfile, "acme corp")
            acme.researched_at = datetime(2024, 1, 1)
            db.commit()
        finally:
            db.close()
        r = self.do_import(self.replace_import_file(blob, data))
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["companies"], 2)
        db = SessionLocal()
        try:
            self.assertEqual(db.get(CompanyProfile, "acme corp").business_model, "MARK-newer-research")
            self.assertEqual(db.get(CompanyProfile, "hidco").business_model, "MARK-hidco-research")
            self.assertEqual(db.get(CompanyProfile, "filtco").business_model, "MARK-current-filtco", "the install's is newer")
            self.assertIsNone(db.get(CompanyProfile, "elsewhere inc"))
        finally:
            db.close()

    def test_hostile_values_are_stored_as_text_or_refused(self):
        blob = self.export()
        data = self.import_data(blob)
        ws = data["workspaces"][0]
        ws["name"] = "<img src=x onerror=alert(1)>"
        ws["sources"].append({"url": "javascript:alert(1)", "label": "bad"})
        ws["documents"][0]["original"] = "../../server-keys.json"
        r = self.do_import(self.replace_import_file(blob, data))
        self.assertEqual(r.status_code, 200, r.text)
        made = r.json()["workspaces"][0]
        self.assertEqual(made["name"], "<img src=x onerror=alert(1)>", "kept as text; the page escapes it")
        self.assertEqual(made["sources"], 1, "what was imported (the javascript: source wasn't)")
        db = SessionLocal()
        try:
            with helpers.as_account(self.bob):
                from backend.db import CompanySource
                urls = [s.careers_url for s in db.query(CompanySource).filter(CompanySource.workspace_id == made["id"])]
                self.assertEqual(urls, ["https://careers.example/bob"])
                self.assertFalse(any(d.stored_name for d in db.query(Document).filter(Document.workspace_id == made["id"])),
                                 "a path that climbs out names no original")
        finally:
            db.close()
        for change in ({"status": "owner"}, {"hidden": "yes please"}, {"ad": {"work_mode": "moon"}}):
            data = self.import_data(blob)
            data["workspaces"][0]["jobs"][0].update(change)
            r = self.do_import(self.replace_import_file(blob, data))
            self.assertEqual((r.status_code, r.json()["code"]), (422, "import_damaged"), change)
            self.assertIn("workspaces.0.jobs.0", r.json()["detail"])
            self.assertNotIn("moon", r.json()["detail"], "the refusal never quotes the file's text")


class ShapeTest(Base):
    """Research and fit results come in only in the shapes the app reads by fixed keys: the research rows are shared
    by every account on the install, and the file is someone else's."""

    def malformed(self, blob) -> bytes:
        data = self.import_data(blob)
        ws = data["workspaces"][0]
        main = next(j for j in ws["jobs"] if j["ad"]["url"] == "https://ads.example/bob/main")
        main["ad"]["location_text"], main["ad"]["work_mode"], main["office"] = "Sydney NSW", "hybrid", None
        main["fit"]["evidence"] = {"summary": ["not text"], "gaps": {"a": 1}, "score": "high",
                                   "requirements": [{"matched": True}, 3, {"point": "MARK-real-point", "evidence": "x"},
                                                    {"point": "MARK-matched", "matched": True,
                                                     "evidence": [{"bullet": "MARK-bullet", "sub_bullets": ["s", 7, None]},
                                                                  {"sub_bullets": []}]}]}
        data["companies"] = [{
            "key": "acme corp", "name": "Acme Corp", "status": "done", "researched_at": "2099-01-01T00:00:00Z",
            "news_researched_at": "2099-01-01T00:00:00Z", "sentiment_researched_at": "2099-01-01T00:00:00Z",
            "offices_checked_at": "2099-01-01T00:00:00Z", "research_version": 999,
            "controversies": [{"title": "MARK-controversy", "year": 2020, "sources": "x"}, {"year": 2020}, "str",
                              {"title": {"a": 1}}],
            "sources": {"a": 1},
            "sentiment": {"rating": {"x": 1}, "positives": "abc", "negatives": ["MARK-neg", "", 5],
                          "sources": [{"url": "javascript:alert(1)"}, {"title": "ok", "url": "https://reviews.example/a"}]},
            "offices": [{"name": "No address"}, {"address": "1 Real St, Sydney NSW", "url": "javascript:x"}, 7,
                        {"address": ["x"]}, {"address": "1 REAL ST, SYDNEY NSW"}],
            "offices_cities": "sydney",
        }]
        return self.replace_import_file(blob, data)

    def test_malformed_research_and_evidence_come_in_as_the_app_reads_them(self):
        from backend import job_chat, research
        import backend.offices as of
        import backend.pipeline as pl

        db = SessionLocal()
        try:
            db.query(CompanyProfile).filter(CompanyProfile.key == "acme corp").delete()
            db.commit()
        finally:
            db.close()
        r = self.do_import(self.malformed(self.export()))
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["companies"], 1)
        ws = r.json()["workspaces"][0]["id"]
        db = SessionLocal()
        try:
            row = db.get(CompanyProfile, "acme corp")
            for name in ("researched_at", "news_researched_at", "sentiment_researched_at", "offices_checked_at"):
                self.assertLessEqual(getattr(row, name), utcnow(), f"{name}: no time later than now")
            self.assertEqual(row.research_version, research.RESEARCH_VERSION)
            p = row.to_dict()
            self.assertEqual(p["controversies"], [{"title": "MARK-controversy", "year": "2020", "summary": "", "sources": []}])
            self.assertEqual(p["sources"], [])
            self.assertEqual(p["sentiment"], {"summary": "", "rating": "", "positives": [], "negatives": ["MARK-neg", "5"],
                                              "sources": [{"title": "ok", "url": "https://reviews.example/a"}]})
            self.assertEqual(p["offices"], [{"name": "", "address": "1 Real St, Sydney NSW", "url": ""}])
            self.assertIsNone(row.offices_cities, "a list or nothing")
            with helpers.as_account(self.bob):
                job = next(j for j in db.query(Job).filter(Job.workspace_id == ws) if j.url == "https://ads.example/bob/main")
                job_id = job.id
                block = job_chat.job_block(db, 1, job, 2000)  # the job chat's context
                fit = json.loads(db.query(import_file.FitResult).filter(import_file.FitResult.job_id == job_id).one().evidence_json)
            self.assertIn("MARK-real-point", block)
            self.assertIn("Employee sentiment", block)
            self.assertEqual(fit, {"summary": "", "gaps": [], "score": None, "requirements": [
                {"point": "MARK-real-point", "must_have": True, "matched": False, "evidence": []},
                {"point": "MARK-matched", "must_have": True, "matched": True,
                 "evidence": [{"bullet": "MARK-bullet", "sub_bullets": ["s", "7"]}]}]})
        finally:
            db.close()
        # the job page (with the employer's offices near the ad's city) and its fit section
        point = {"lat": -33.87, "lng": 151.21}
        with mock.patch.object(of, "geocode", return_value=point), mock.patch.object(pl, "geocode", return_value=point), \
                helpers.client(self.bob.email) as c:
            r = c.get(f"/api/jobs/{job_id}", headers={"X-Workspace": str(ws)})
            self.assertEqual(r.status_code, 200, r.text[:300])
            self.assertEqual([o["address"] for o in r.json()["office_candidates"]], ["1 Real St, Sydney NSW"])
            self.assertEqual(len(r.json()["fit"]["requirements"]), 2)
            r = c.get("/api/companies", headers={"X-Workspace": str(ws)})  # the company panel's data
            self.assertEqual(r.status_code, 200)
            self.assertEqual(r.json()["acme corp"]["offices"][0]["address"], "1 Real St, Sydney NSW")
        # "Research again" for that company keeps the offices the install knows
        research._save({"key": "acme corp", "name": "Acme Corp"}, model="test", profile={
            "is_company": True, "official_name": "Acme", "is_recruiter": False, "business_model": "b", "ownership": "o",
            "headquarters": "h", "controversies": [], "controversy_note": "", "sources": [], "employee_count": "",
            "glassdoor_url": "", "sentiment": {"summary": "", "rating": "", "positives": [], "negatives": [], "sources": []},
            "offices": [{"name": "Head office", "address": "2 Other St, Sydney NSW", "url": "https://acme.example/contact"}]})
        db = SessionLocal()
        try:
            offices = db.get(CompanyProfile, "acme corp").to_dict()["offices"]
            self.assertEqual([o["address"] for o in offices], ["2 Other St, Sydney NSW", "1 Real St, Sydney NSW"])
        finally:
            db.close()

    def test_offices_the_install_found_are_kept(self):
        """A row the office search made (no research yet) takes the file's research, and keeps its offices."""
        blob = self.export()
        data = self.import_data(blob)
        data["companies"] = [{"key": "hidco", "name": "Hidco", "researched_at": "2026-01-01T00:00:00Z",
                              "offices": [{"name": "HQ", "address": "1 Known St, Sydney NSW", "url": "https://h.example/a"},
                                          {"name": "", "address": "5 New Rd, Sydney NSW", "url": "https://h.example/b"}],
                              "offices_cities": ["Melbourne"]}]
        checked = datetime(2026, 9, 1)
        db = SessionLocal()
        try:
            db.merge(CompanyProfile(key="hidco", name="Hidco", status="missing", offices_checked_at=checked,
                                    offices_json=json.dumps([{"name": "Mine", "address": "1 known st, sydney nsw", "url": ""}]),
                                    offices_cities=json.dumps(["sydney"])))
            db.commit()
        finally:
            db.close()
        r = self.do_import(self.replace_import_file(blob, data))
        self.assertEqual(r.status_code, 200, r.text)
        db = SessionLocal()
        try:
            row = db.get(CompanyProfile, "hidco")
            self.assertEqual([o["address"] for o in json.loads(row.offices_json)],
                             ["1 known st, sydney nsw", "5 New Rd, Sydney NSW"])
            self.assertEqual(row.offices_checked_at, checked, "the file had no time: the install's stays")
            self.assertEqual(json.loads(row.offices_cities), ["melbourne", "sydney"])
            self.assertIsNotNone(row.researched_at)
        finally:
            db.close()

    def test_no_time_is_later_than_now(self):
        blob = self.export()
        data = self.import_data(blob)
        for job in data["workspaces"][0]["jobs"]:
            job["first_seen"] = job["last_seen"] = "2099-05-05T00:00:00Z"
        ws = self.do_import(self.replace_import_file(blob, data)).json()["workspaces"][0]["id"]
        db = SessionLocal()
        try:
            with helpers.as_account(self.bob):
                for job in db.query(Job).filter(Job.workspace_id == ws):
                    self.assertLessEqual(job.first_seen, utcnow())
                    self.assertLessEqual(job.last_seen, utcnow())
        finally:
            db.close()


class LimitTest(Base):
    def test_the_counts_are_what_was_imported(self):
        blob = self.export()
        data = self.import_data(blob)
        jobs = data["workspaces"][0]["jobs"]
        twice = dict(jobs[0])  # the same job ref again: skipped
        jobs.append(twice)
        data["workspaces"][0]["sources"].append(dict(data["workspaces"][0]["sources"][0]))  # a repeated address
        blob = self.replace_import_file(blob, data)
        listed = self.check(blob).json()["workspaces"][0]
        self.assertEqual((listed["jobs"], listed["sources"]), (8, 2), "the check lists what the file holds")
        made = self.do_import(blob).json()["workspaces"][0]
        self.assertEqual({k: made[k] for k in ("jobs", "documents", "chats", "cover_letters", "sources", "pins")},
                         {"jobs": 7, "documents": 3, "chats": 1, "cover_letters": 1, "sources": 1, "pins": 1})

    def test_a_file_of_too_many_things_is_refused_before_they_are_made(self):
        blob = self.export()
        with mock.patch.object(data_import, "MAX_ITEMS", 5), \
                mock.patch.object(import_file.ImportFile, "model_validate", side_effect=AssertionError("validated")):
            r = self.check(blob)
        self.assertEqual((r.status_code, r.json()["code"]), (413, "import_too_many"))

    def test_the_limits_take_in_whatever_an_export_holds(self):
        from backend import export
        self.assertGreater(data_import.MAX_ZIP_BYTES, export.MAX_EXPORT_BYTES)
        self.assertGreaterEqual(data_import.MAX_FILE_BYTES, export.MAX_EXPORT_BYTES)
        self.assertGreaterEqual(data_import.MAX_UNPACKED_BYTES, export.MAX_EXPORT_BYTES)
        self.assertGreaterEqual(import_file.DOC_TEXT, docs.MAX_UPLOAD_BYTES, "an uploaded .txt keeps all of its text")
        self.assertGreaterEqual(import_file.DOC_TEXT, export.MAX_EXPORT_BYTES, "any text an export can hold")
        long = "x" * (import_file.LONG + 10)
        doc = import_file.DocumentIn.model_validate({"ref": 1, "filename": "big.txt", "text": long, "concise_text": long})
        self.assertEqual(len(doc.text), len(long))
        self.assertNotIn("at most", data_import.too_large().detail)


class AddressTest(Base):
    def test_only_web_addresses_come_in_where_pages_make_links(self):
        blob = self.export()
        data = self.import_data(blob)
        job = next(j for j in data["workspaces"][0]["jobs"] if j["ad"]["url"] == "https://ads.example/bob/hidden")
        job["ad"]["url"] = "javascript:alert(1)"
        data["companies"] = [{"key": "hidco", "name": "Hidco", "researched_at": "2026-01-01T00:00:00Z",
                              "glassdoor_url": "javascript:alert(2)",
                              "sources": [{"title": "ok", "url": "https://hidco.example/about"},
                                          {"title": "bad", "url": "javascript:alert(3)"}],
                              "controversies": [{"headline": "x", "url": "data:text/html,y"}]}]
        r = self.do_import(self.replace_import_file(blob, data))
        self.assertEqual(r.status_code, 200, r.text)
        ws = r.json()["workspaces"][0]["id"]
        db = SessionLocal()
        try:
            with helpers.as_account(self.bob):
                urls = [j.url for j in db.query(Job).filter(Job.workspace_id == ws)]
            self.assertNotIn("javascript:alert(1)", urls)
            self.assertIn(None, urls, "the job came in, without an address")
            hidco = db.get(CompanyProfile, "hidco").to_dict()
            self.assertIsNone(hidco["glassdoor_url"])
            self.assertEqual(hidco["sources"], [{"title": "ok", "url": "https://hidco.example/about"}],
                             "a page without a web address is left out")
            self.assertEqual(hidco["controversies"], [], "a controversy without a title is left out")
        finally:
            db.close()


class RefusalTest(Base):
    def test_a_newer_format_version_is_refused_with_a_plain_message(self):
        blob = self.export()
        data = self.import_data(blob)
        data["format_version"] = import_file.FORMAT_VERSION + 1
        for path in ("/api/account/import/check", "/api/account/import"):
            r = self.post(path, self.replace_import_file(blob, data), pick="0")
            self.assertEqual((r.status_code, r.json()["code"]), (422, "import_newer"))
            self.assertIn("newer JobHuntPA", r.json()["detail"])
            self.assertIn("Update JobHuntPA", r.json()["detail"])

    def test_an_export_with_no_import_file_says_which_version_writes_one(self):
        files = {n: b for n, b in entries(self.export()).items() if not n.endswith(import_file.FILE_NAME)}
        r = self.check(zip_of(files))
        self.assertEqual((r.status_code, r.json()["code"]), (422, "import_no_file"))
        self.assertIn(import_file.FIRST_APP_VERSION, r.json()["detail"])
        r = self.check(zip_of({"holiday.jpg": b"\xff\xd8"}))
        self.assertEqual((r.status_code, r.json()["code"]), (422, "import_not_export"))
        r = self.check(b"not a zip at all")
        self.assertEqual((r.status_code, r.json()["code"]), (422, "import_not_export"))
        r = self.check(b"")
        self.assertEqual((r.status_code, r.json()["code"]), (422, "import_empty"))

    def test_a_zip_with_a_path_that_climbs_out_is_refused(self):
        files = entries(self.export())
        for bad in ("../evil.txt", "JobHuntPA export/../../evil.txt", "/etc/evil", "C:/evil.txt", "a\\..\\evil.txt"):
            r = self.check(zip_of({**files, bad: b"x"}))
            self.assertEqual((r.status_code, r.json()["code"]), (422, "import_unsafe_path"), bad)
        self.assertTrue(data_import.unsafe_path("a/./b"))
        self.assertFalse(data_import.unsafe_path("JobHuntPA export 2026-10-13/Workspaces/A/Documents/Original files/x.pdf"))

    def test_a_zip_too_large_when_unpacked_is_refused(self):
        blob = self.export()
        data = self.import_data(blob)
        data["padding"] = " " * (3 * 1024 * 1024)  # compresses about a thousand times
        r = self.check(self.replace_import_file(blob, data))
        self.assertEqual((r.status_code, r.json()["code"]), (413, "import_too_large"))
        with mock.patch.object(data_import, "MAX_UNPACKED_BYTES", 1000):
            r = self.check(blob)
            self.assertEqual((r.status_code, r.json()["code"]), (413, "import_too_large"))
        with mock.patch.object(data_import, "MAX_FILE_BYTES", 1000):
            self.assertEqual(self.check(blob).status_code, 413)
        with mock.patch.object(data_import, "MAX_ZIP_BYTES", 1000):
            self.assertEqual(self.check(blob).status_code, 413)

    def test_a_lying_entry_size_is_caught_while_unpacking(self):
        blob = self.export()
        archive = data_import.ExportZip(blob)
        name = archive.base + import_file.FILE_NAME
        archive.infos[name].file_size = 10  # the zip says less than the entry holds
        with self.assertRaises(data_import.AuthProblem) as caught:
            archive.read(name, data_import.MAX_FILE_BYTES)
        self.assertIn(caught.exception.status, (413, 422))

    def test_signed_out_no_csrf_and_no_recent_password_are_refused(self):
        from fastapi.testclient import TestClient
        from backend.app import app
        blob = self.export()
        with TestClient(app) as c:
            for path in ("/api/account/import/check", "/api/account/import?pick=0"):
                self.assertEqual(c.post(path, content=blob).status_code, 401)
        with helpers.client(self.bob.email) as c:
            c.headers.pop("X-CSRF-Token", None)
            for path in ("/api/account/import/check", "/api/account/import?pick=0"):
                r = c.post(path, content=blob)
                self.assertEqual((r.status_code, r.json()["code"]), (403, "csrf"))
        with helpers.client(self.bob.email) as c:
            with engine.begin() as conn:
                from sqlalchemy import text as sql
                conn.execute(sql("UPDATE sessions SET reauth_at = :t WHERE user_id = :u"),
                             {"t": utcnow() - timedelta(minutes=30), "u": self.bob.id})
            for path in ("/api/account/import/check", "/api/account/import?pick=0"):
                r = c.post(path, content=blob)
                self.assertEqual((r.status_code, r.json()["code"]), (403, "reauth_required"))
        before = self.workspaces()
        self.assertEqual(self.workspaces(), before)

    def test_one_import_at_a_time(self):
        blob = self.export()
        data_import._running.add(self.bob.id)
        try:
            r = self.check(blob)
            self.assertEqual((r.status_code, r.json()["code"]), (429, "import_running"))
        finally:
            data_import._running.discard(self.bob.id)
        self.assertTrue(data_import._slots.acquire(blocking=False))
        try:
            r = self.check(blob)
            self.assertEqual((r.status_code, r.json()["code"]), (503, "import_busy"))
        finally:
            data_import._slots.release()
        self.assertEqual(self.check(blob).status_code, 200)
        self.assertNotIn(self.bob.id, data_import._running)

    def test_an_upload_that_stalls_gives_the_place_back(self):
        with helpers.client(self.bob.email) as c:
            with mock.patch.object(data_import, "READ_SECONDS", 0.001), \
                    mock.patch.object(data_import.Request, "stream", side_effect=lambda *a, **k: self._slow()):
                r = c.post("/api/account/import/check", content=b"PK" * 10)
        self.assertEqual((r.status_code, r.json()["code"]), (408, "import_slow"))
        self.assertNotIn(self.bob.id, data_import._running)
        self.assertTrue(data_import._slots.acquire(blocking=False))
        data_import._slots.release()

    @staticmethod
    async def _slow():
        import asyncio
        await asyncio.sleep(5)
        yield b"late"

    def test_a_failed_import_leaves_nothing_behind(self):
        blob = self.export()
        before = self.workspaces()
        uploads = set(os.listdir(docs.UPLOAD_DIR))
        with mock.patch.object(data_import.Importer, "chats", side_effect=RuntimeError("boom")):
            with self.assertRaises(RuntimeError):
                self.do_import(blob)
        self.assertEqual(self.workspaces(), before)
        self.assertEqual(set(os.listdir(docs.UPLOAD_DIR)), uploads, "the sealed originals were deleted again")
        self.assertNotIn(self.bob.id, data_import._running)

    def test_the_log_says_it_happened_with_counts_and_nothing_else(self):
        blob = self.export()
        with self.assertLogs("backend.data_import", level="INFO") as logs:
            self.assertEqual(self.do_import(blob).status_code, 200)
        logged = "\n".join(logs.output)
        self.assertRegex(logged, r"user \d+ imported 1 workspace\(s\): 7 jobs, 3 documents \(1 original files\), 1 chats")
        self.assertNotIn("MARK", logged)
        self.assertNotIn("My job search", logged)

    def test_the_import_source_never_unpacks_to_disk(self):
        for module in (data_import, import_file):
            with open(module.__file__, encoding="utf-8") as f:
                src = f.read()
            for forbidden in ("extractall", "extract(", "tempfile", "open(", "write_bytes", "mkdtemp"):
                self.assertNotIn(forbidden, src.replace("self.zip.open(", ""), (module.__name__, forbidden))


if __name__ == "__main__":
    unittest.main()


class MappingsTest(Base):
    """Phase I1: a source's page layout comes with its source, and the account's saved job-site layouts come with the
    file (backend/mappings.py); never a source on the local network, and never a layout of another shape."""

    def setUp(self):
        super().setUp()
        from backend import mapping_format, mappings

        self.careers = mapping_format.mapping("careers", {"list": {"selector": "li.job"}, "title": {"selector": "h3"},
                                                          "link": {"selector": ":scope", "attribute": "href"}})
        self.job = mapping_format.mapping("job", {"title": {"selector": "h1"}, "description": {"selector": "#ad"}})
        with helpers.as_account(self.bob):
            db = SessionLocal()
            try:
                src = db.query(CompanySource).filter(CompanySource.workspace_id == self.bob.workspace_id).first()
                src.mapping = self.careers
                mappings.save(db, self.bob.id, "jobs.example.org", self.job)
                mappings.save(db, self.bob.id, "board.example.net", self.job)
                db.commit()
            finally:
                db.close()

    def sources(self, ws, who):
        db = SessionLocal()
        try:
            with helpers.as_account(who):
                return [(s.careers_url, s.mapping) for s in
                        db.query(CompanySource).filter(CompanySource.workspace_id == ws).order_by(CompanySource.id)]
        finally:
            db.close()

    def sites(self, who):
        from backend import mappings

        db = SessionLocal()
        try:
            with helpers.as_account(who):
                return {r.site: r.mapping for r in mappings.for_account(db, who.id)}
        finally:
            db.close()

    def test_layouts_go_out_and_come_back_and_an_accounts_own_are_kept(self):
        from backend import mappings

        blob = self.export()
        data = self.import_data(blob)
        self.assertEqual(data["workspaces"][0]["sources"][0]["mapping"], self.careers)
        self.assertEqual([m["site"] for m in data["site_mappings"]], ["board.example.net", "jobs.example.org"])
        alice = helpers.make_user(f"MARK-name-alice-{os.urandom(4).hex()}")
        own = {**self.job, "fields": {**self.job["fields"], "title": {"selector": "h1.alice"}}}
        with helpers.as_account(alice):
            db = SessionLocal()
            mappings.save(db, alice.id, "jobs.example.org", own)
            db.commit()
            db.close()
        r = self.do_import(blob, who=alice)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["site_mappings"], 1)
        ws = r.json()["workspaces"][0]["id"]
        self.assertEqual([m for _, m in self.sources(ws, alice)], [self.careers])
        self.assertEqual(self.sites(alice), {"jobs.example.org": own, "board.example.net": self.job})
        self.assertEqual(len(self.sites(self.bob)), 2, "the exporting account's are untouched")

    def test_local_sources_and_layouts_of_another_shape_are_left_out(self):
        blob = self.export()
        data = self.import_data(blob)
        sources = data["workspaces"][0]["sources"]
        sources[0]["mapping"] = {**self.careers, "fields": {"list": {"selector": "li:hover"}}}
        sources.append({"ref": 90, "label": "Router", "url": "http://192.168.1.1/jobs", "enabled": True})
        sources.append({"ref": 91, "label": "Here", "url": "http://localhost:8000/careers", "enabled": True})
        sources.append({"ref": 92, "label": "Fine", "url": "https://careers.example.org/fine", "enabled": True,
                        "mapping": "javascript:alert(1)"})
        data["site_mappings"] = [{"site": "javascript:alert(1)", "mapping": self.job},
                                 {"site": "ok.example", "mapping": {**self.job, "kind": "careers"}},
                                 {"site": "fine.example", "mapping": self.job}, "junk"]
        r = self.do_import(self.replace_import_file(blob, data))
        self.assertEqual(r.status_code, 200, r.text)
        ws = r.json()["workspaces"][0]["id"]
        self.assertEqual(self.sources(ws, self.bob), [("https://careers.example/bob", None),
                                                      ("https://careers.example.org/fine", None)])
        self.assertEqual(r.json()["site_mappings"], 1)
        self.assertIn("fine.example", self.sites(self.bob))
