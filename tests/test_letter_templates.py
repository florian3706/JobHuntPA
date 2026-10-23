"""Tests for the cover letter template and the .docx/.odt/.pdf downloads (backend.letter_templates). No network, and
no LibreOffice: the PDF backend is patched (a fake ``soffice`` stands in for the conversion).

Run:  python3 -m unittest tests.test_letter_templates -v
"""
from __future__ import annotations

import io
import os
import sys
import tempfile
import textwrap
import unittest
from pathlib import Path
from unittest import mock

_TMP = tempfile.mkdtemp()
os.environ.setdefault("JOBHUNT_DB_PATH", os.path.join(_TMP, "test.db"))  # never the real database
os.environ.setdefault("JOBHUNT_SCHEDULER", "off")
os.environ.setdefault("JOBHUNT_ENV_PATH", os.path.join(_TMP, "test.env"))

from tests import helpers  # noqa: E402  (signed-in clients; sets the test environment)

from backend import docs  # noqa: E402
from backend import letter_templates as lt  # noqa: E402
from backend.db import UPLOAD_DIR, Job, LetterTemplate, SessionLocal, init_db  # noqa: E402

init_db()

VALUES = {"company": "Acme", "role": "Program Manager", "location": "", "date": "3 October 2026"}
LETTER = "Dear Hiring Manager,\n\nFirst paragraph,\nwith a line break.\n\n\nSecond   paragraph.\n\nKind regards"
ODT_STYLE = "Letter_20_Body"  # (the style "Letter Body" as ODF names it)
LETTER_PARAGRAPHS = ["Dear Hiring Manager,", "First paragraph,\nwith a line break.", "Second   paragraph.",
                     "Kind regards"]


def setUpModule():
    helpers.enter_admin_scope()


def docx_template(marker: bool = True) -> bytes:
    """A letter template: {company} in the header, {date} in the footer, {role} in a table cell, and the {letter}
    paragraph (style "Body Text") split across two runs, the first one bold."""
    import docx

    d = docx.Document()
    section = d.sections[0]
    section.header.paragraphs[0].text = "Application to {company}"
    section.footer.paragraphs[0].text = "Sent on {date}"
    d.add_paragraph("Dear Hiring Manager,")
    p = d.add_paragraph(style="Body Text")
    if marker:
        p.add_run("{let").bold = True
        p.add_run("ter}")
    else:
        p.add_run("no marker here")
    in_one_run = d.add_paragraph()
    in_one_run.add_run("Kept ")
    in_one_run.add_run("{company}").italic = True
    in_one_run.add_run(" formatting, {location} stays")
    d.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run("Role: {role}")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def odt_template(marker: bool = True) -> bytes:
    """The same as an .odt: a master page with a header and a footer, the {letter} paragraph (style "Letter Body")
    with the marker split between a bold span and the text after it, and a table cell with {role}."""
    from odf.opendocument import OpenDocumentText
    from odf.style import Footer, Header, MasterPage, PageLayout, Style, TextProperties
    from odf.table import Table, TableCell, TableColumn, TableRow
    from odf.text import P, Span

    doc = OpenDocumentText()
    body = Style(name="Letter Body", family="paragraph")
    doc.styles.addElement(body)
    bold = Style(name="Bold", family="text")
    bold.addElement(TextProperties(fontweight="bold"))
    doc.automaticstyles.addElement(bold)
    layout = PageLayout(name="Page")
    doc.automaticstyles.addElement(layout)
    master = MasterPage(name="Standard", pagelayoutname=layout)
    header, footer = Header(), Footer()
    header.addElement(P(text="Application to {company}"))
    footer.addElement(P(text="Sent on {date}"))
    master.addElement(header)
    master.addElement(footer)
    doc.masterstyles.addElement(master)
    doc.text.addElement(P(text="Dear Hiring Manager,"))
    p = P(stylename=body)
    if marker:
        p.addElement(Span(stylename=bold, text="{let"))
        p.addText("ter}")
    else:
        p.addText("no marker here")
    doc.text.addElement(p)
    table = Table(name="T")
    table.addElement(TableColumn())
    row, cell = TableRow(), TableCell()
    cell.addElement(P(text="Role: {role}"))
    row.addElement(cell)
    table.addElement(row)
    doc.text.addElement(table)
    out = io.BytesIO()
    doc.write(out)
    return out.getvalue()


def docx_texts(data: bytes):
    """(body paragraphs as (style, text), header text, footer text, first table cell text)."""
    import docx

    d = docx.Document(io.BytesIO(data))
    section = d.sections[0]
    return ([(p.style.name, p.text) for p in d.paragraphs], section.header.paragraphs[0].text,
            section.footer.paragraphs[0].text, d.tables[0].cell(0, 0).text)


def odt_paragraphs(data: bytes) -> list[tuple[str, str]]:
    """Every paragraph of the .odt (header and footer first, in the master page), as (style, text)."""
    from odf import teletype
    from odf.opendocument import load
    from odf.text import P

    doc = load(io.BytesIO(data))
    out = []
    for root in (doc.masterstyles, doc.text):
        for p in root.getElementsByType(P):
            out.append((p.getAttribute("stylename") or "", teletype.extractText(p)))
    return out


class FillTest(unittest.TestCase):
    def test_docx(self):
        body, header, footer, cell = docx_texts(lt.fill_docx(docx_template(), LETTER, VALUES))
        self.assertEqual(header, "Application to Acme")
        self.assertEqual(footer, "Sent on 3 October 2026")
        self.assertEqual(cell, "Role: Program Manager")
        self.assertEqual(body[0], ("Normal", "Dear Hiring Manager,"))
        self.assertEqual(body[1:5], [("Body Text", t) for t in LETTER_PARAGRAPHS])  # in order, each the marker's style
        self.assertEqual(body[5], ("Normal", "Kept Acme formatting, {location} stays"))  # no value: left as written
        self.assertFalse(any("{letter}" in t for _, t in body))

    def test_docx_formatting(self):
        import docx

        d = docx.Document(io.BytesIO(lt.fill_docx(docx_template(), LETTER, VALUES)))
        letter = d.paragraphs[1:5]
        for p in letter:  # the marker's first run was bold: every letter paragraph is
            self.assertTrue([r for r in p.runs if r.text][0].bold, p.text)
        kept = d.paragraphs[5]
        self.assertEqual([(r.text, r.italic) for r in kept.runs if r.text][1], ("Acme", True),
                         "a placeholder in one run is filled in place")
        self.assertEqual(len(letter[1]._p.xpath(".//w:br")), 1, "a line break in a paragraph stays one")

    def test_docx_text_around_the_marker(self):
        import docx

        d = docx.Document()
        d.add_paragraph("Before {letter} after {company}")
        out = io.BytesIO()
        d.save(out)
        body = docx_texts_simple(lt.fill_docx(out.getvalue(), "One\n\nTwo", VALUES))
        self.assertEqual(body, ["Before One", "Two after Acme"])
        body = docx_texts_simple(lt.fill_docx(out.getvalue(), "Only", VALUES))
        self.assertEqual(body, ["Before Only after Acme"])

    def test_odt(self):
        paras = odt_paragraphs(lt.fill_odt(odt_template(), LETTER, VALUES))
        texts = [t for _, t in paras]
        self.assertEqual(texts[:2], ["Application to Acme", "Sent on 3 October 2026"])  # header and footer
        self.assertEqual(texts[2], "Dear Hiring Manager,")
        self.assertEqual(paras[3:7], [(ODT_STYLE, t) for t in LETTER_PARAGRAPHS])
        self.assertEqual(texts[7], "Role: Program Manager")
        self.assertFalse(any("{letter}" in t for t in texts))

    def test_odt_keeps_the_first_span_and_spaces(self):
        from odf.opendocument import load
        from odf.text import P, Span

        doc = load(io.BytesIO(lt.fill_odt(odt_template(), LETTER, VALUES)))
        letter = [p for p in doc.text.getElementsByType(P) if p.getAttribute("stylename") == ODT_STYLE]
        self.assertEqual(len(letter), 4)
        for p in letter:
            spans = p.getElementsByType(Span)
            self.assertEqual([s.getAttribute("stylename") for s in spans], ["Bold"])

    def test_refused_templates(self):
        for fill, template in ((lt.fill_docx, docx_template), (lt.fill_odt, odt_template)):
            with self.assertRaises(lt.TemplateError) as caught:
                fill(template(marker=False), LETTER, VALUES)
            self.assertEqual(caught.exception.kind, "no_marker")
            with self.assertRaises(lt.TemplateError) as caught:
                fill(b"PK\x03\x04 not really a zip", LETTER, VALUES)
            self.assertEqual(caught.exception.kind, "unreadable")

    def test_characters_xml_cant_hold_are_dropped(self):
        body, *_ = docx_texts(lt.fill_docx(docx_template(), "A\x00B\x0bC", VALUES))
        self.assertEqual(body[1][1], "ABC")

    def test_plain_documents(self):
        import docx

        d = docx.Document(io.BytesIO(lt.plain_docx(LETTER)))
        self.assertEqual([p.text for p in d.paragraphs], LETTER_PARAGRAPHS)
        self.assertEqual([t for _, t in odt_paragraphs(lt.plain_odt(LETTER))], LETTER_PARAGRAPHS)

    def test_plain_pdf(self):
        from pypdf import PdfReader

        pdf = lt.to_pdf_plain("Dear Zoë, “quotes” {braces} € Łódź 日本​.\n\n" + "More words. " * 400)
        self.assertTrue(pdf.startswith(b"%PDF"))
        reader = PdfReader(io.BytesIO(pdf))
        self.assertGreater(len(reader.pages), 1, "long letters run onto a second page")
        text = reader.pages[0].extract_text()
        self.assertIn("Dear Zoë, “quotes” {braces} € Lódz ??.", text)

    def test_cp1252_text(self):
        self.assertEqual(lt.cp1252_text("Zoë “x” {y} — Ł ő 日​\ttab"), "Zoë “x” {y} — L o ? tab")


def docx_texts_simple(data: bytes) -> list[str]:
    import docx

    return [p.text for p in docx.Document(io.BytesIO(data)).paragraphs]


class _Workspace:
    """A fresh workspace of the test account with a job and a saved letter, deleted after."""

    def __init__(self, c, letter=LETTER):
        self.c = c
        self.id = c.post("/api/workspaces", json={"name": "Templates"}).json()["id"]
        self.h = {"X-Workspace": str(self.id)}
        db = SessionLocal()
        j = Job(workspace_id=self.id, url=f"https://templates.example/job/{self.id}", title="Program Manager",
                company="Acme", location_text="", status="to_review", detail_status="full")
        db.add(j)
        db.commit()
        self.job_id = j.id
        db.close()
        r = c.put(f"/api/jobs/{self.job_id}/cover-letter", json={"text": letter}, headers=self.h)
        assert r.status_code == 200, r.text

    def upload(self, name, data, headers=None):
        return self.c.post("/api/letter-template", files={"file": (name, data, "application/octet-stream")},
                           headers=self.h if headers is None else headers)

    def download(self, fmt):
        return self.c.get(f"/api/jobs/{self.job_id}/cover-letter.{fmt}?ws={self.id}")

    def stored_names(self) -> list[str]:
        db = SessionLocal()
        try:
            return [t.stored_name for t in db.query(LetterTemplate).filter(LetterTemplate.workspace_id == self.id)]
        finally:
            db.close()

    def close(self):
        self.c.delete(f"/api/workspaces/{self.id}")


class TemplateApiTest(unittest.TestCase):
    def setUp(self):
        self.c = helpers.client()
        self.c.__enter__()
        self.addCleanup(self.c.__exit__, None, None, None)
        self.ws = _Workspace(self.c)
        self.addCleanup(self.ws.close)

    def test_upload_get_replace_delete(self):
        ws = self.ws
        r = self.c.get("/api/letter-template", headers=ws.h)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIsNone(r.json()["template"])
        self.assertIn(r.json()["pdf"], ("libreoffice", "plain"))
        self.assertFalse(r.json()["pdf_uses_template"])

        r = ws.upload("My letterhead.docx", docx_template())
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["template"]["filename"], "My letterhead.docx")
        self.assertEqual(r.json()["template"]["filetype"], ".docx")
        self.assertTrue(r.json()["template"]["created_at"])
        first = ws.stored_names()
        self.assertEqual(len(first), 1)
        self.assertTrue((UPLOAD_DIR / first[0]).is_file())
        self.assertEqual(docs.read_upload(first[0]), docx_template(), "sealed, and opens to the upload")
        with mock.patch.object(lt, "pdf_backend", return_value="libreoffice"):
            self.assertTrue(self.c.get("/api/letter-template", headers=ws.h).json()["pdf_uses_template"])
        with mock.patch.object(lt, "pdf_backend", return_value="plain"):
            got = self.c.get("/api/letter-template", headers=ws.h).json()
        self.assertEqual((got["pdf"], got["pdf_uses_template"]), ("plain", False))
        self.assertEqual(got["template"]["filename"], "My letterhead.docx")

        r = ws.upload("Other.ODT", odt_template())
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["template"]["filetype"], ".odt")
        second = ws.stored_names()
        self.assertEqual(len(second), 1)
        self.assertNotEqual(second, first)
        self.assertFalse((UPLOAD_DIR / first[0]).exists(), "the old template's file is gone")
        self.assertEqual(self.c.get("/api/letter-template", headers=ws.h).json()["template"]["filename"], "Other.ODT")

        self.assertEqual(self.c.delete("/api/letter-template", headers=ws.h).status_code, 204)
        self.assertFalse((UPLOAD_DIR / second[0]).exists())
        self.assertIsNone(self.c.get("/api/letter-template", headers=ws.h).json()["template"])
        self.assertEqual(self.c.delete("/api/letter-template", headers=ws.h).status_code, 404)

    def test_refusals(self):
        ws = self.ws
        cases = [
            ("letter.txt", b"{letter}", 400, "must be a .docx or .odt"),
            ("letter.doc", b"{letter}", 400, "must be a .docx or .odt"),
            ("no marker.docx", docx_template(marker=False), 400, lt.NO_MARKER),
            ("no marker.odt", odt_template(marker=False), 400, lt.NO_MARKER),
            ("broken.docx", b"PK\x03\x04 this is not a zip", 400, lt.UNREADABLE),
            ("broken.odt", b"not an odt", 400, lt.UNREADABLE),
            ("odt named docx.docx", odt_template(), 400, lt.UNREADABLE),
        ]
        for name, data, status, detail in cases:
            with self.subTest(name=name):
                r = ws.upload(name, data)
                self.assertEqual(r.status_code, status, r.text)
                self.assertIn(detail, r.json()["detail"])
        with mock.patch.object(docs, "MAX_UPLOAD_BYTES", 100):
            r = ws.upload("big.docx", docx_template())
        self.assertEqual(r.status_code, 413, r.text)
        self.assertEqual(ws.stored_names(), [], "nothing refused was stored")
        self.assertIsNone(self.c.get("/api/letter-template", headers=ws.h).json()["template"])

    def test_a_refused_upload_keeps_the_old_template(self):
        ws = self.ws
        self.assertEqual(ws.upload("Good.docx", docx_template()).status_code, 200)
        before = ws.stored_names()
        self.assertEqual(ws.upload("Bad.docx", docx_template(marker=False)).status_code, 400)
        self.assertEqual(ws.stored_names(), before)
        self.assertTrue((UPLOAD_DIR / before[0]).is_file())

    def test_other_workspace(self):
        ws = self.ws
        self.assertEqual(ws.upload("Mine.docx", docx_template()).status_code, 200)
        other = _Workspace(self.c)
        self.addCleanup(other.close)
        self.assertIsNone(self.c.get("/api/letter-template", headers=other.h).json()["template"])
        self.assertEqual(self.c.delete("/api/letter-template", headers=other.h).status_code, 404)
        r = self.c.get(f"/api/jobs/{ws.job_id}/cover-letter.docx?ws={other.id}")
        self.assertEqual(r.status_code, 404, "another workspace's job")
        self.assertEqual(len(ws.stored_names()), 1, "still there")
        bob = helpers.make_user()
        with helpers.client(bob.email) as theirs:
            for method in ("get", "delete"):
                r = getattr(theirs, method)("/api/letter-template", headers=ws.h)
                self.assertEqual(r.status_code, 404, f"{method}: another account's workspace")
            r = theirs.post("/api/letter-template", headers=ws.h,
                            files={"file": ("x.docx", docx_template(), "application/octet-stream")})
            self.assertEqual(r.status_code, 404)
        self.assertEqual(len(ws.stored_names()), 1)

    def test_deleting_the_workspace_removes_the_file(self):
        other = _Workspace(self.c)
        self.assertEqual(other.upload("T.docx", docx_template()).status_code, 200)
        stored = other.stored_names()
        other.close()
        self.assertFalse((UPLOAD_DIR / stored[0]).exists())

    def test_a_copied_workspace_gets_its_own_sealed_copy(self):
        ws = self.ws
        self.assertEqual(ws.upload("T.docx", docx_template()).status_code, 200)
        r = self.c.post("/api/workspaces", json={"name": "Copy", "copy_from": ws.id, "copy_documents": True})
        self.assertEqual(r.status_code, 201, r.text)
        copy_id = r.json()["id"]
        self.addCleanup(self.c.delete, f"/api/workspaces/{copy_id}")
        db = SessionLocal()
        try:
            copied = db.query(LetterTemplate).filter(LetterTemplate.workspace_id == copy_id).one()
            self.assertEqual((copied.filename, copied.filetype), ("T.docx", ".docx"))
            self.assertNotEqual(copied.stored_name, ws.stored_names()[0])
            self.assertEqual(docs.read_upload(copied.stored_name), docx_template())
        finally:
            db.close()


class DownloadTest(unittest.TestCase):
    def setUp(self):
        self.c = helpers.client()
        self.c.__enter__()
        self.addCleanup(self.c.__exit__, None, None, None)
        self.ws = _Workspace(self.c)
        self.addCleanup(self.ws.close)

    def check_name(self, r, ext):
        self.assertEqual(r.headers["content-disposition"],
                         f'attachment; filename="Cover letter - Acme - Program Manager.{ext}"')

    def test_with_a_docx_template(self):
        ws = self.ws
        self.assertEqual(ws.upload("T.docx", docx_template()).status_code, 200)
        r = ws.download("docx")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.headers["content-type"], lt.MEDIA_TYPES["docx"])
        self.check_name(r, "docx")
        body, header, footer, cell = docx_texts(r.content)
        self.assertEqual(body[1:5], [("Body Text", t) for t in LETTER_PARAGRAPHS])
        self.assertEqual((header, cell), ("Application to Acme", "Role: Program Manager"))
        self.assertRegex(footer, r"^Sent on \d{1,2} [A-Z][a-z]+ \d{4}$")  # today, as "23 October 2026"
        self.assertFalse(any("{letter}" in t for _, t in body))
        r = ws.download("odt")  # the template is a .docx: a plain .odt
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.headers["content-type"], lt.MEDIA_TYPES["odt"])
        self.check_name(r, "odt")
        self.assertEqual([t for _, t in odt_paragraphs(r.content)], LETTER_PARAGRAPHS)

    def test_with_an_odt_template(self):
        ws = self.ws
        self.assertEqual(ws.upload("T.odt", odt_template()).status_code, 200)
        r = ws.download("odt")
        self.assertEqual(r.status_code, 200, r.text)
        paras = odt_paragraphs(r.content)
        self.assertEqual(paras[0][1], "Application to Acme")
        self.assertEqual(paras[3:7], [(ODT_STYLE, t) for t in LETTER_PARAGRAPHS])
        self.assertEqual(paras[7][1], "Role: Program Manager")
        self.assertFalse(any("{letter}" in t for _, t in paras))
        r = ws.download("docx")  # a plain .docx
        self.assertEqual(docx_texts_simple(r.content), LETTER_PARAGRAPHS)

    def test_pdf_without_libreoffice(self):
        from pypdf import PdfReader

        ws = self.ws
        self.assertEqual(ws.upload("T.docx", docx_template()).status_code, 200)
        with mock.patch.object(lt, "pdf_backend", return_value="plain"):
            r = ws.download("pdf")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.headers["content-type"], "application/pdf")
        self.check_name(r, "pdf")
        self.assertTrue(r.content.startswith(b"%PDF"))
        text = PdfReader(io.BytesIO(r.content)).pages[0].extract_text()
        self.assertTrue(text.startswith("Dear Hiring Manager,"), text[:80])
        self.assertIn("Kind regards", text)
        self.assertNotIn("Application to", text, "the plain PDF never uses the template")

    def test_without_a_template(self):
        ws = self.ws
        r = ws.download("docx")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(docx_texts_simple(r.content), LETTER_PARAGRAPHS)
        r = ws.download("odt")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual([t for _, t in odt_paragraphs(r.content)], LETTER_PARAGRAPHS)
        with mock.patch.object(lt, "pdf_backend", return_value="plain"):
            r = ws.download("pdf")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.content.startswith(b"%PDF"))

    def test_unknown_formats_and_no_letter(self):
        ws = self.ws
        for fmt in ("rtf", "DOCX", "doc", "pdf.exe"):
            r = ws.download(fmt)
            self.assertEqual(r.status_code, 404, fmt)
        db = SessionLocal()
        j = Job(workspace_id=ws.id, url=f"https://templates.example/job/{ws.id}/2", title="T", company="C",
                status="to_review", detail_status="full")
        db.add(j)
        db.commit()
        job_id = j.id
        db.close()
        r = self.c.get(f"/api/jobs/{job_id}/cover-letter.pdf?ws={ws.id}")
        self.assertEqual((r.status_code, r.json()["detail"]), (404, "No cover letter drafted yet"))


FAKE_SOFFICE = textwrap.dedent("""\
    #!{python}
    import json, pathlib, sys
    pathlib.Path({calls!r}).write_text(json.dumps(sys.argv[1:]))
    args = sys.argv[1:]
    if {fail!r}:
        sys.exit(1)
    outdir = pathlib.Path(args[args.index("--outdir") + 1])
    source = pathlib.Path(args[-1])
    (outdir / (source.stem + ".pdf")).write_bytes(b"%PDF-1.4 fake from " + source.suffix.encode())
""")


@unittest.skipIf(os.name == "nt", "the fake soffice is a script with a #! line")
class LibreOfficeTest(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.data = self.dir / "data"
        self.data.mkdir()
        self.calls = self.dir / "calls.json"
        patcher = mock.patch.object(lt, "_temp_parent", return_value=self.data)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.addCleanup(lt.reset_pdf_backend)

    def fake(self, fail=False) -> str:
        path = self.dir / ("fail-soffice" if fail else "soffice")
        path.write_text(FAKE_SOFFICE.format(python=sys.executable, calls=str(self.calls), fail=fail))
        path.chmod(0o755)
        patcher = mock.patch.object(lt, "_backend", ("libreoffice", str(path)))
        patcher.start()
        self.addCleanup(patcher.stop)
        return str(path)

    def test_the_conversion_and_its_folder(self):
        import json

        exe = self.fake()
        self.assertEqual(lt.pdf_backend(), "libreoffice")
        pdf = lt.to_pdf_libreoffice(b"the filled template", ".odt")
        self.assertEqual(pdf, b"%PDF-1.4 fake from .odt")
        self.assertEqual(list(self.data.iterdir()), [], "the conversion's folder is gone")
        args = json.loads(self.calls.read_text())
        self.assertEqual(args[:3], ["--headless", "--norestore", "--nologo"])
        self.assertTrue(args[3].startswith("-env:UserInstallation=file://"), args[3])
        self.assertEqual(args[4:6], ["--convert-to", "pdf"])
        self.assertTrue(args[-1].endswith("letter.odt"))
        self.assertTrue(Path(args[-1]).parent.name.startswith(lt.TEMP_PREFIX))
        self.assertTrue(exe)

    def test_a_failed_conversion(self):
        self.fake(fail=True)
        with self.assertRaises(lt.PdfError):
            lt.to_pdf_libreoffice(b"x", ".docx")
        self.assertEqual(list(self.data.iterdir()), [])

    def test_stale_folders_are_removed(self):
        self.fake()
        stale = self.data / f"{lt.TEMP_PREFIX}old"
        stale.mkdir()
        (stale / "letter.docx").write_bytes(b"left by a crash")
        os.utime(stale, (1, 1))
        fresh = self.data / f"{lt.TEMP_PREFIX}running"
        fresh.mkdir()
        lt.to_pdf_libreoffice(b"x", ".docx")
        self.assertFalse(stale.exists())
        self.assertTrue(fresh.exists(), "a conversion running now is left alone")

    def test_over_http(self):
        with helpers.client() as c:
            ws = _Workspace(c)
            try:
                self.fake()
                r = ws.download("pdf")  # no template: the plain .docx goes to LibreOffice
                self.assertEqual((r.status_code, r.content), (200, b"%PDF-1.4 fake from .docx"))
                self.assertEqual(ws.upload("T.odt", odt_template()).status_code, 200)
                self.assertTrue(c.get("/api/letter-template", headers=ws.h).json()["pdf_uses_template"])
                r = ws.download("pdf")  # the filled .odt
                self.assertEqual((r.status_code, r.content), (200, b"%PDF-1.4 fake from .odt"))
                self.fake(fail=True)
                with self.assertLogs("backend.letter_templates", "WARNING") as logged:
                    r = ws.download("pdf")
                self.assertEqual((r.status_code, r.json()["detail"]), (502, lt.PDF_FAILED))
                self.assertEqual(logged.output, ["WARNING:backend.letter_templates:LibreOffice could not turn a "
                                                 "cover letter into a PDF: exit status 1"])
                self.assertEqual(list(self.data.iterdir()), [])
            finally:
                ws.close()

    def test_finding_libreoffice(self):
        lt.reset_pdf_backend()
        with mock.patch.object(lt, "_candidates", return_value=[str(self.dir / "missing")]):
            self.assertEqual(lt.pdf_backend(), "plain")
        lt.reset_pdf_backend()
        exe = self.dir / "soffice-found"
        exe.write_text("#!/bin/sh\n")
        exe.chmod(0o755)
        with mock.patch.object(lt, "_candidates", return_value=[str(self.dir / "missing"), str(exe)]):
            self.assertEqual(lt._find(), ("libreoffice", str(exe)))
            with mock.patch.object(lt, "_candidates", return_value=[]):
                self.assertEqual(lt.pdf_backend(), "libreoffice", "found once, for the process")
        lt.reset_pdf_backend()


if __name__ == "__main__":
    unittest.main()
