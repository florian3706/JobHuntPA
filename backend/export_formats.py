"""The export's text files as Word (DOCX) and OpenDocument (ODT) documents, made in memory from its Markdown.

The export (backend/export.py) writes every text file as Markdown first; ``convert`` turns that Markdown into the
other file types the person chose. It reads only the Markdown the exporter itself writes: headings (``#``), bullet
lists (``-``, nested by two spaces), tables (``| a | b |`` with a ``|---|`` line), ``---`` rules, ``**bold**``,
links ``[label](address)`` and paragraphs (a line break inside one is kept). Everything else is text.

Other people's text in that Markdown is escaped (``export.text``): ``\\[``, ``\\|``, ``&lt;`` and ``&amp;`` come
back here as the characters they stand for. In a DOCX or ODT file text is only ever text (the libraries escape it
for XML), so nothing in an ad or an AI answer can become markup; characters XML can't hold at all (control
characters) are dropped. Links stay links only for http(s) addresses and for the relative links between the
export's own files, whose ``.md`` ending becomes the file type's own, so they point at the file of the same type.

Nothing here touches the disk: both libraries write into a ``BytesIO``.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Optional

FORMATS = ("md", "docx", "odt")
EXTENSIONS = {"md": ".md", "docx": ".docx", "odt": ".odt"}
LONGEST_EXT = max(len(e) for e in EXTENSIONS.values())

_NOT_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ud800-\udfff￾￿]")
_PUNCT = set("!\"#$%&'()*+,-./:;<=>?@[\\]^_`{|}~")
_HEADING = re.compile(r"^(#{1,6}) +(.*)$")
_BULLET = re.compile(r"^( *)- (.*)$")
_SEPARATOR = re.compile(r"^\|(\s*:?-+:?\s*\|)+\s*$")
_ENTITY = re.compile(r"&(lt|amp);")


# ---------------------------------------------------------------------------
# Reading the exporter's Markdown
# ---------------------------------------------------------------------------

@dataclass
class Run:
    text: str
    bold: bool = False
    href: Optional[str] = None


@dataclass
class Block:
    kind: str                     # heading | bullet | table | rule | para
    level: int = 0                # a heading's 1-6, a bullet's depth from 0
    runs: list[Run] = field(default_factory=list)
    rows: list[list[list[Run]]] = field(default_factory=list)  # a table's rows of cells (the first is the header)


def clean(value: str) -> str:
    """Text with the characters XML 1.0 can't hold taken out."""
    return _NOT_XML.sub("", value)


def _plain(value: str) -> str:
    """Escaped Markdown text as the characters it stands for (backslash escapes and the two entities)."""
    out, i = [], 0
    while i < len(value):
        c = value[i]
        if c == "\\" and i + 1 < len(value) and value[i + 1] in _PUNCT:
            out.append(value[i + 1])
            i += 2
            continue
        out.append(c)
        i += 1
    return clean(_ENTITY.sub(lambda m: "<" if m.group(1) == "lt" else "&", "".join(out)))


def _link_end(value: str, start: int) -> Optional[tuple[int, int, int]]:
    """For a ``[`` at ``start``: (end of the label, start of the address, end of the address) when a link follows."""
    i = start + 1
    while i < len(value):
        c = value[i]
        if c == "\\":
            i += 2
            continue
        if c == "[":
            return None
        if c == "]":
            if value.startswith("(", i + 1):
                close = value.find(")", i + 2)
                if close > i + 2 and not re.search(r"\s", value[i + 2:close]):
                    return i, i + 2, close
            return None
        i += 1
    return None


def inline(value: str, ext: str) -> list[Run]:
    """The runs of one line or paragraph: plain text, ``**bold**`` and links (relative ``.md`` targets get ``ext``)."""
    runs: list[Run] = []
    buf: list[str] = []
    bold = False

    def flush() -> None:
        if buf:
            runs.append(Run(_plain("".join(buf)), bold))
            buf.clear()

    i = 0
    while i < len(value):
        c = value[i]
        if c == "\\" and i + 1 < len(value):
            buf.append(value[i:i + 2])
            i += 2
            continue
        if c == "`" and value.find("`", i + 1) > i + 1:  # a code span (file names in the README): the text inside
            close = value.find("`", i + 1)
            buf.append(value[i + 1:close])
            i = close + 1
            continue
        if value.startswith("**", i):
            flush()
            bold = not bold
            i += 2
            continue
        if c == "[":
            found = _link_end(value, i)
            if found is not None:
                label_end, target_start, target_end = found
                flush()
                runs.append(Run(_plain(value[i + 1:label_end]), bold, _target(value[target_start:target_end], ext)))
                i = target_end + 1
                continue
        buf.append(c)
        i += 1
    flush()
    return [r for r in runs if r.text]


def _target(address: str, ext: str) -> Optional[str]:
    """A link's address in the other file type: http(s) as it is, one of the export's own files with ``ext``,
    anything else no link."""
    if re.match(r"^https?://", address, re.I):
        return address
    if re.match(r"^[a-z][a-z0-9+.-]*:", address, re.I) or address.startswith(("/", "\\")):
        return None
    return address[:-3] + ext if address.endswith(".md") else address


def _cells(row: str) -> list[str]:
    """A table row's cells: split on the pipes that aren't escaped."""
    body = row.strip()
    body = body[1:] if body.startswith("|") else body
    cells, buf, i = [], [], 0
    while i < len(body):
        c = body[i]
        if c == "\\" and i + 1 < len(body):
            buf.append(body[i:i + 2])
            i += 2
            continue
        if c == "|":
            cells.append("".join(buf).strip())
            buf = []
        else:
            buf.append(c)
        i += 1
    if "".join(buf).strip():
        cells.append("".join(buf).strip())
    return cells


def parse(markdown: str, ext: str) -> list[Block]:
    """The exporter's Markdown as blocks."""
    lines = markdown.replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    para: list[str] = []

    def end_para() -> None:
        if para:
            blocks.append(Block("para", runs=inline("\n".join(para), ext)))
            para.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        heading, bullet = _HEADING.match(line), _BULLET.match(line)
        if not line.strip():
            end_para()
        elif heading:
            end_para()
            blocks.append(Block("heading", level=len(heading.group(1)), runs=inline(heading.group(2), ext)))
        elif line.strip() == "---":
            end_para()
            blocks.append(Block("rule"))
        elif bullet:
            end_para()
            blocks.append(Block("bullet", level=min(len(bullet.group(1)) // 2, 2), runs=inline(bullet.group(2), ext)))
        elif line.startswith("|") and i + 1 < len(lines) and _SEPARATOR.match(lines[i + 1]):
            end_para()
            rows = [_cells(line)]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(_cells(lines[i]))
                i += 1
            width = max(len(r) for r in rows)
            blocks.append(Block("table", rows=[[inline(c, ext) for c in r + [""] * (width - len(r))] for r in rows]))
            continue
        else:
            para.append(line)
        i += 1
    end_para()
    return blocks


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------

def _docx_link(paragraph, run: Run) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rel = paragraph.part.relate_to(run.href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    r = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    if run.bold:
        props.append(OxmlElement("w:b"))
    props.append(color)
    props.append(underline)
    r.append(props)
    t = OxmlElement("w:t")
    t.text = run.text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    link.append(r)
    paragraph._p.append(link)


def _docx_runs(paragraph, runs: list[Run]) -> None:
    for run in runs:
        if run.href:
            _docx_link(paragraph, run)
            continue
        for n, piece in enumerate(run.text.split("\n")):
            if n:
                paragraph.add_run().add_break()
            if piece:
                paragraph.add_run(piece).bold = run.bold or None


def to_docx(markdown: str) -> bytes:
    import docx

    document = docx.Document()
    for block in parse(markdown, ".docx"):
        if block.kind == "heading":
            _docx_runs(document.add_heading("", level=min(block.level, 4)), block.runs)
        elif block.kind == "bullet":
            style = "List Bullet" if block.level == 0 else f"List Bullet {block.level + 1}"
            _docx_runs(document.add_paragraph(style=style), block.runs)
        elif block.kind == "table":
            table = document.add_table(rows=len(block.rows), cols=len(block.rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(block.rows):
                for c, cell_runs in enumerate(row):
                    paragraph = table.cell(r, c).paragraphs[0]
                    _docx_runs(paragraph, [Run(x.text, x.bold or r == 0, x.href) for x in cell_runs])
        elif block.kind == "rule":
            document.add_paragraph("—" * 12)
        else:
            _docx_runs(document.add_paragraph(), block.runs)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# ODT (odfpy)
# ---------------------------------------------------------------------------

def _odt_styles(document):
    from odf.style import ListLevelProperties, ParagraphProperties, Style, TableCellProperties, TextProperties
    from odf.text import ListLevelStyleBullet, ListStyle

    styles = {}
    for level, size in ((1, "18pt"), (2, "15pt"), (3, "13pt"), (4, "12pt")):
        s = Style(name=f"Heading {level}", family="paragraph", defaultoutlinelevel=str(level))
        s.addElement(ParagraphProperties(margintop="0.3cm", marginbottom="0.15cm", keepwithnext="always"))
        s.addElement(TextProperties(fontsize=size, fontweight="bold"))
        document.styles.addElement(s)
        styles[f"h{level}"] = s
    bold = Style(name="Strong", family="text")
    bold.addElement(TextProperties(fontweight="bold"))
    document.styles.addElement(bold)
    styles["bold"] = bold
    bullets = ListStyle(name="Bullets")
    for level in range(1, 4):
        item = ListLevelStyleBullet(level=str(level), bulletchar="•")
        item.addElement(ListLevelProperties(spacebefore=f"{0.6 * (level - 1):.1f}cm", minlabelwidth="0.6cm"))
        bullets.addElement(item)
    document.automaticstyles.addElement(bullets)
    styles["list"] = bullets
    cell = Style(name="Cell", family="table-cell")
    cell.addElement(TableCellProperties(border="0.5pt solid #808080", padding="0.08cm"))
    document.automaticstyles.addElement(cell)
    styles["cell"] = cell
    return styles


def _odt_runs(element, runs: list[Run], styles) -> None:
    from odf import teletype
    from odf.text import A, Span

    for run in runs:
        target = element
        if run.href:
            target = A(href=run.href, type="simple")
            element.addElement(target)
        if run.bold:
            span = Span(stylename=styles["bold"])
            target.addElement(span)
            target = span
        teletype.addTextToElement(target, run.text)


def to_odt(markdown: str) -> bytes:
    from odf.opendocument import OpenDocumentText
    from odf.table import Table, TableCell, TableColumn, TableRow
    from odf.text import H, List, ListItem, P

    document = OpenDocumentText()
    styles = _odt_styles(document)
    lists: list = []  # the open lists, one per depth
    tables = 0
    for block in parse(markdown, ".odt"):
        if block.kind != "bullet":
            lists = []
        if block.kind == "heading":
            level = min(block.level, 4)
            h = H(outlinelevel=level, stylename=styles[f"h{level}"])
            _odt_runs(h, block.runs, styles)
            document.text.addElement(h)
        elif block.kind == "bullet":
            depth = min(block.level, len(lists))
            del lists[depth + 1:]
            if depth == len(lists):
                new = List(stylename=styles["list"]) if depth == 0 else List()
                if depth == 0:
                    document.text.addElement(new)
                else:
                    lists[-1][1].addElement(new)  # nested inside the last item of the list above
                lists.append([new, None])
            item = ListItem()
            p = P()
            _odt_runs(p, block.runs, styles)
            item.addElement(p)
            lists[depth][0].addElement(item)
            lists[depth][1] = item
            del lists[depth + 1:]
        elif block.kind == "table":
            tables += 1
            table = Table(name=f"Table{tables}")
            table.addElement(TableColumn(numbercolumnsrepeated=len(block.rows[0])))
            for r, row in enumerate(block.rows):
                tr = TableRow()
                for cell_runs in row:
                    tc = TableCell(valuetype="string", stylename=styles["cell"])
                    p = P()
                    _odt_runs(p, [Run(x.text, x.bold or r == 0, x.href) for x in cell_runs], styles)
                    tc.addElement(p)
                    tr.addElement(tc)
                table.addElement(tr)
            document.text.addElement(table)
        elif block.kind == "rule":
            document.text.addElement(P(text="—" * 12))
        else:
            p = P()
            _odt_runs(p, block.runs, styles)
            document.text.addElement(p)
    out = io.BytesIO()
    document.write(out)
    return out.getvalue()


def convert(markdown: str, fmt: str) -> bytes:
    """The Markdown text file as ``fmt`` (one of FORMATS)."""
    if fmt == "md":
        return markdown.encode("utf-8")
    if fmt == "docx":
        return to_docx(markdown)
    if fmt == "odt":
        return to_odt(markdown)
    raise ValueError("unknown export file type")
