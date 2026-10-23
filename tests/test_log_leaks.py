"""No log line may hold anything of a person's (accounts design §2.5: "Logs must never contain plaintext"; Phase L).

One account goes through the whole app, end to end, with the network and the AI service mocked: sign-up, the search
profile, documents, a search (mocked sources returning planted ads), scoring, chat, a cover letter, company research,
the browser extension, a data export (in every file type) and its import, merged and unmerged ads, a copied
workspace, a password change, new recovery codes shown again, a scheduled search, and a password reset
with a recovery code (and, in a process of its own, the first run's folder step). Every personal field holds a planted marker (``ZQX``), and every secret the run makes
(passwords, recovery codes, session and CSRF tokens, extension keys) is collected. Logging is captured at DEBUG for every logger, with the
text of each record formatted the way a log file would show it, tracebacks included, and then none of the markers or
secrets may be in it. The test also checks that the markers really are in the person's data (so a pass isn't just an
empty run) and that the log isn't empty.

Run:  python3 -m unittest tests.test_log_leaks -v
"""
from __future__ import annotations

import contextlib
import io
import json
import logging
import os
import time
import unittest
import zipfile
from dataclasses import replace
from datetime import datetime, timezone
from unittest import mock

from tests import helpers  # noqa: E402  (sets the test environment first)

from fastapi.testclient import TestClient  # noqa: E402

from backend import app as app_module, ratelimit, users  # noqa: E402
from backend.adapters.base import Adapter, Posting  # noqa: E402
from backend.db import SessionLocal, User, Workspace  # noqa: E402

MARK = "ZQX"
PASSWORD = "zqx-first-password-1"
NEW_PASSWORD = "zqx-second-password-2"
LLM_ENV = {"LLM_API_KEY": "zqx-llm-key-never-logged", "LLM_BASE_URL": "https://api.openai.com/v1", "LLM_MODEL": "m",
           "LLM_REASONING_EFFORT": ""}
NOW = datetime.now(timezone.utc)

SCORE_REPLY = json.dumps({
    "score": 81, "summary": "ZQX fit summary: strong delivery background.", "gaps": ["ZQX gap: no cloud"],
    "requirements": [{"point": "ZQX requirement", "matched": True, "must_have": True,
                      "evidence": [{"bullet": "ZQX evidence from the CV", "sub_bullets": []}]}]})
RESEARCH_REPLY = json.dumps({
    "official_name": "ZQX Corp Pty Ltd", "is_company": True, "is_recruiter": False,
    "business_model": "ZQX sells delivery software.", "ownership": "Private.", "headquarters": "Zqxtown, Australia",
    "controversies": [], "controversy_note": "ZQX note.",
    "sources": [{"title": "About ZQX", "url": "https://zqx-corp.example/about"}]})


def _tiny_pdf(text: str) -> bytes:
    """A one-page PDF with ``text`` (no library needed), enough for pypdf to read it back."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
               b"/Resources << /Font << /F1 5 0 R >> >> >>",
               b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    start = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    return out + b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, start)


def _docx(text: str, template: bool = False) -> bytes:
    """A one-paragraph .docx; as a ``template``, a letterhead (``text`` in the header and footer) with the
    {{letter}} marker in a font this computer may not have."""
    import docx
    from docx.shared import Pt

    document = docx.Document()
    if template:
        document.sections[0].header.paragraphs[0].text = f"{text} {{{{date}}}}"
        document.sections[0].footer.paragraphs[0].text = text
        run = document.add_paragraph().add_run("{{letter}}")
        run.font.name, run.font.size = "ZQX Garamond", Pt(11)
    else:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class _Response:
    """What ``httpx.post`` gives back, as far as the AI service code reads it."""

    def __init__(self, status: int, data: dict):
        self.status_code, self._data, self.headers = status, data, {}
        self.text = json.dumps(data)

    def json(self):
        return self._data


class _Source(Adapter):
    """A source with planted ads, listed and detailed from memory."""

    source, method = "greenhouse:zqx-corp", "test"

    def __init__(self, postings: list[Posting]):
        self.postings = postings

    failure: Exception | None = None

    def list_postings(self):
        if self.failure:
            raise self.failure
        return [replace(p) for p in self.postings]

    def fetch_detail(self, posting):
        return replace(posting, description=posting.description + " " + MARK + " the full advertisement.",
                       detail_status="full")


def _ad(n: int, **fields) -> Posting:
    base = dict(url=f"https://jobs.zqx.example/{n}", title="ZQX Delivery Lead", company="ZQX Corp", external_id=str(n),
                location_text="Zqxtown NSW", country="AU", description="ZQX include: lead the delivery program.")
    return Posting(**{**base, **fields})


class Everything(logging.Handler):
    """Every record, as a log file would show it (traceback included)."""

    def __init__(self):
        super().__init__(logging.NOTSET)
        self.lines: list[str] = []
        self.setFormatter(logging.Formatter("%(name)s %(levelname)s %(message)s"))

    def emit(self, record):
        try:
            line = self.format(record)
        except Exception as exc:  # noqa: BLE001  (a bad format string is a bug of its own)
            line = f"{record.name}: could not format ({exc!r}) {record.msg!r}"
        if record.name.startswith(("httpx", "httpcore")) and "testserver" in line:
            return  # the test client's own request lines, not the app's
        self.lines.append(line)


@contextlib.contextmanager
def capture_all():
    """DEBUG for the root and every logger that exists (SQLAlchemy, multipart and the rest included), one handler on
    all of them, and what is printed or written to stderr. Yields (the handler, the printed text). The loggers the
    app deliberately holds quiet (app.QUIET_LOGGERS: httpx, the Claude SDK, pypdf) keep that level: it is what
    protects a person whatever the root logger is set to, and ``test_the_quiet_loggers_stay_quiet`` checks it."""
    handler, printed = Everything(), io.StringIO()
    root = logging.getLogger()
    quiet = tuple(app_module.QUIET_LOGGERS)
    named = [lg for name, lg in list(logging.Logger.manager.loggerDict.items()) if isinstance(lg, logging.Logger)
             and not any(name == q or name.startswith(q + ".") for q in quiet)]
    saved = [(lg, lg.level, lg.propagate) for lg in [root, *named]]
    for lg in [root, *named]:
        lg.setLevel(logging.DEBUG)
    root.addHandler(handler)
    for lg in named:
        if not lg.propagate:
            lg.addHandler(handler)
    try:
        with contextlib.redirect_stdout(printed), contextlib.redirect_stderr(printed):
            yield handler, printed
    finally:
        root.removeHandler(handler)
        for lg in named:
            lg.removeHandler(handler)
        for lg, level, propagate in saved:
            lg.setLevel(level)


class FullRun(unittest.TestCase):
    def setUp(self):
        helpers.started()
        ratelimit.reset_all()
        self.addCleanup(ratelimit.reset_all)
        patches = [
            mock.patch.dict(os.environ, LLM_ENV),
            mock.patch("backend.pipeline.source_adapter", side_effect=lambda *a, **k: self.source),
            mock.patch("backend.llm_providers.httpx.post", side_effect=self.llm),
            mock.patch("backend.geo.httpx.get", side_effect=self.nominatim),
            mock.patch("backend.geo.time.sleep"),
            mock.patch("backend.research.run_response", side_effect=self.web_search),
        ]
        for patch in patches:
            patch.start()
            self.addCleanup(patch.stop)
        self.source = _Source([])
        self.secrets: dict[str, str] = {}
        self.llm_calls = 0
        self.llm_status = 200  # an error status: the provider answers it, echoing the request
        self.mapping_reply: dict | None = None  # the AI's answer to a page outline (None: a score, refused)

    # -- clients --------------------------------------------------------------------------------------------------

    def anonymous(self, **kwargs) -> TestClient:
        c = TestClient(app_module.app, **kwargs)
        c.__enter__()
        self.addCleanup(c.__exit__, None, None, None)
        return c

    def signed_in(self) -> helpers.SignedInClient:
        c = helpers.client()
        c.__enter__()
        self.addCleanup(c.__exit__, None, None, None)
        return c

    # -- the outside world, mocked ------------------------------------------------------------------------------

    def llm(self, url, json=None, headers=None, timeout=None, **_):
        self.llm_calls += 1
        prompt = " ".join(str(m.get("content")) for m in (json or {}).get("messages", []))
        if self.llm_status != 200:
            return _Response(self.llm_status, {"error": {"message": f"rejected, you sent: {prompt[:200]}"}})
        if "ZQXBOOM" in prompt and "response_format" in (json or {}):  # a provider that echoes the request in its error
            return _Response(400, {"error": {"message": f"bad request, you sent: {prompt[:200]}"}})
        text = SCORE_REPLY if "response_format" in (json or {}) else (
            "Dear hiring team,\n\nZQX answer: I led delivery programs.\n\nZqx Quillon")
        if "PAGE OUTLINE" in prompt and self.mapping_reply is not None:  # an AI mapping (backend/mappings.py)
            text = self.mapping_reply["careers" if '"list" picks every job entry' in prompt else "job"]
        return _Response(200, {"choices": [{"message": {"content": text}, "finish_reason": "stop"}],
                               "usage": {"prompt_tokens": 10, "completion_tokens": 5}})

    def nominatim(self, url, params=None, **_):
        import httpx

        if "Zqxfail" in (params or {}).get("q", ""):
            raise httpx.ConnectError(f"cannot reach the geocoder for {params['q']}")
        hits = [{"lat": "-33.8", "lon": "151.2", "importance": 0.5, "addresstype": "suburb", "place_rank": 20,
                 "display_name": "Zqxtown, New South Wales", "address": {"country_code": "au"}}]
        return mock.Mock(json=lambda: hits, raise_for_status=lambda: None)

    def web_search(self, payload, cfg, what="web search"):
        return {"id": "resp_1", "status": "completed", "output": [
            {"type": "web_search_call", "status": "completed", "results": [
                {"type": "text_result", "title": "About ZQX", "url": "https://zqx-corp.example/about"}]},
            {"type": "message", "content": [{"type": "output_text", "text": RESEARCH_REPLY, "annotations": []}]}]}

    # -- the person ---------------------------------------------------------------------------------------------

    def signed_up(self):
        """An account whose sign-in name is personal text (the app makes no second account itself: tests/helpers.py
        writes it as an older install's would be), signed in through the app."""
        email = "Zqx Quillon zqxowner.private@example.invalid"
        account = helpers.extra_account(email, PASSWORD)
        c = self.anonymous()
        r = c.post("/api/auth/login", json={"name": email, "password": PASSWORD})
        self.assertEqual(r.status_code, 200, r.text)
        c.headers["X-CSRF-Token"] = r.json()["csrf"]
        self.secrets.update(password=PASSWORD, csrf=r.json()["csrf"], session=c.cookies.get("jh_session") or "",
                            device=c.cookies.get("jh_device") or "",
                            **{f"recovery_code_{i}": code for i, code in enumerate(account.recovery_codes)})
        self.email = email
        return c

    def the_install_settings_and_a_new_name(self, c):
        """The name changed in Settings, and the AI key and the public address saved there: none of
        these values may reach the log (only the names of the settings that changed)."""
        from backend import config
        env_file = config.ENV_PATH
        saved = env_file.read_text() if env_file.exists() else None
        self.addCleanup(lambda: env_file.write_text(saved) if saved is not None else env_file.unlink(missing_ok=True))
        new_name = "Zqx Quillon Renamed zqxowner"
        r = c.put("/api/auth/name", json={"name": new_name})
        self.assertEqual(r.status_code, 200, r.text)
        self.email = new_name
        helpers._passwords[new_name] = PASSWORD
        llm = {k: os.environ.get(k) for k in ("LLM_API_KEY", "LLM_BASE_URL", "LLM_MODEL")}
        key = "sk-zqx-secret-ai-key-000111"
        r = c.put("/api/settings/llm", json={"base_url": "https://zqx-llm.example/v1", "model": "zqx-model",
                                              "api_key": key})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(c.put("/api/settings/address", json={"public_base_url": "https://zqx-host.example",
                                                               "allowed_hosts": ["zqx-alias.example"]}).status_code, 200)
        self.assertEqual(c.put("/api/settings/address", json={"public_base_url": "", "allowed_hosts": []}).status_code,
                         200)
        r = c.put("/api/settings/llm", json={"base_url": llm["LLM_BASE_URL"] or "", "model": llm["LLM_MODEL"] or "",
                                              "api_key": llm["LLM_API_KEY"] or ""})
        self.assertEqual(r.status_code, 200, r.text)
        self.secrets.update(ai_key=key, host="zqx-host", alias="zqx-alias")

    def run_to_end(self, c, run: dict) -> dict:
        for _ in range(600):
            got = c.get(f"/api/runs/{run['id']}").json()
            if got["state"] not in ("queued", "running"):
                return got
            time.sleep(0.05)
        self.fail("the run didn't finish")

    def test_nothing_personal_reaches_a_log(self):
        notice = []
        with capture_all() as (log, printed):
            logging.getLogger("test").info("the capture works")
            c = self.signed_up()
            self.the_install_settings_and_a_new_name(c)
            self.sign_up_and_set_up(c)
            ads = self.search_and_score(c)
            self.talk_to_the_ai(c, ads)
            self.documents_and_saved_pages(c)
            self.commute_links(c)
            self.browser_extension(c)
            self.when_things_fail(c)
            self.with_claude(c)
            self.export(c)
            self.merges_a_copy_and_the_password(c)
            self.scheduled_search_and_recovery(c, notice)
        everything = "\n".join(log.lines + [printed.getvalue()])
        self.assertGreater(len(log.lines), 100, "the capture saw almost nothing")
        self.assertIn("the capture works", everything)
        planted = {"marker": MARK.lower(), "email": "zqxowner", "name": "quillon",
                   **{k: v.lower() for k, v in self.secrets.items() if v}}
        for what, value in planted.items():
            found = [line[:300] for line in log.lines if value in line.lower()]
            self.assertEqual(found, [], f"{what} ({value!r}) is in the log")
            self.assertNotIn(value, printed.getvalue().lower(), f"{what} was printed")
        self.assertTrue(notice, "the steps didn't all run")

    # -- the steps ----------------------------------------------------------------------------------------------

    def sign_up_and_set_up(self, c):
        # The search profile, a pin, a source and documents.
        self.assertEqual(c.put("/api/profile", json={
            "titles": ["ZQX Delivery Lead"], "keywords_include": ["zqx include"], "keywords_exclude": ["zqx exclude"],
            "seek_enabled": False, "seek_locations": ["Zqxtown NSW"], "salary_floor": 100000}).status_code, 200)
        self.assertEqual(c.put("/api/dealbreakers", json={"industries": ["zqx industry"],
                                                           "keywords": ["zqx dealbreaker"]}).status_code, 200)
        self.assertEqual(c.post("/api/pins", json={"label": "ZQX home", "kind": "home", "lat": -33.8, "lng": 151.2,
                                                   "radius_km": 20}).status_code, 201)
        self.assertEqual(c.post("/api/sources", json={"url": "https://careers.zqx.example/jobs",
                                                      "label": "ZQX dream employer"}).status_code, 201)
        # Careers pages in bulk: the preview, the add and the export (then the source goes, so the search skips it).
        text = '"ZQX bulk, employer",https://careers.zqx-bulk.example/jobs\nZQX local,http://192.168.0.9/zqx\n'
        rows = c.post("/api/sources/bulk/preview", json={"text": text}).json()["rows"]
        self.assertEqual([x["status"] for x in rows], ["new", "refused"])
        r = c.post("/api/sources/bulk", json={"rows": [{"company": x["company"], "url": x["url"]} for x in rows]})
        self.assertEqual([x["result"] for x in r.json()["results"]], ["added", "refused"], r.text)
        self.assertIn('"ZQX bulk, employer"', c.get("/api/sources/export").text)
        self.assertEqual(c.delete(f"/api/sources/{r.json()['results'][0]['id']}").status_code, 200)
        r = c.post("/api/documents/notes", json={"title": "ZQX project notes", "text": "ZQX note: I ran the ZQX launch.",
                                                 "kind": "other"})
        self.assertEqual(r.status_code, 201, r.text)
        for name, data, kind in (("ZQX-cv.docx", _docx("Zqx Quillon. ZQX resume: delivery lead at ZQX Corp."), "resume"),
                                 ("ZQX-letter.pdf", _tiny_pdf("ZQX earlier cover letter"), "cover_letter"),
                                 ("ZQX-extra.txt", b"ZQX extra document text", "other")):
            r = c.post("/api/documents", files={"file": (name, data, "application/octet-stream")}, data={"kind": kind})
            self.assertEqual(r.status_code, 200, r.text)
        texts = " ".join(d["text"] if "text" in d else "" for d in c.get("/api/documents").json())
        self.assertEqual(len(c.get("/api/documents").json()), 4, texts)

    def search_and_score(self, c):
        self.source.postings = [
            _ad(1),
            _ad(2, title="ZQX Delivery Lead (boom)", description="ZQX include. ZQXBOOM please lead."),
            _ad(3, description="ZQX include. ZQX exclude: not for you."),
            _ad(4, location_text="Zqxfail NSW"),
        ]
        run = c.post("/api/search/run")
        self.assertEqual(run.status_code, 200, run.text)
        done = self.run_to_end(c, run.json())
        self.assertEqual(done["state"], "done", done)
        jobs = c.get("/api/jobs?include=excluded,hidden").json()
        self.assertTrue([j for j in jobs if MARK in j["title"]], "the planted ads should be in the job list")
        self.assertTrue([j for j in jobs if MARK in (j.get("excluded_reason") or "")] or
                        [j for j in jobs if j.get("excluded_reason")], "one ad should have been excluded")
        scored = [j for j in jobs if (j.get("fit") or {}).get("status") == "ok"]
        self.assertTrue(scored, "something should have been scored")
        job_id = scored[0]["id"]
        r = c.post(f"/api/score/{job_id}")
        self.assertEqual(r.status_code, 200, r.text)
        r = c.post("/api/score/run", json={"mode": "all"})
        self.assertEqual(self.run_to_end(c, r.json())["state"], "done")
        self.assertEqual(c.patch(f"/api/jobs/{job_id}/status", json={"status": "shortlisted"}).status_code, 200)
        self.assertEqual(c.get(f"/api/jobs/{job_id}").status_code, 200)
        self.assertEqual(c.post("/api/jobs/refilter").status_code, 200)
        self.job_id = job_id
        return jobs

    def talk_to_the_ai(self, c, jobs):
        job_id = self.job_id
        r = c.post("/api/chats", json={"job_ids": [job_id], "message": "ZQX chat question: how do I fit?"})
        self.assertEqual(r.status_code, 200, r.text)
        chat = r.json()
        r = c.post(f"/api/chats/{chat['id']}/messages", json={"message": "ZQX follow-up with the web", "web_search": True})
        self.assertEqual(r.status_code, 200, r.text)
        r = c.post(f"/api/jobs/{job_id}/cover-letter", json={"instructions": "ZQX instruction: mention my launch"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(c.put(f"/api/jobs/{job_id}/cover-letter", json={"text": "ZQX edited letter"}).status_code, 200)
        self.assertEqual(c.get(f"/api/jobs/{job_id}/cover-letter.docx").status_code, 200)
        # The letter as PDF, .odt and .docx from the plain layout, then from a letter template (a letterhead with
        # the marker in a font this computer doesn't have: the substitution is answered, never logged).
        for fmt in ("pdf", "odt"):
            self.assertEqual(c.get(f"/api/jobs/{job_id}/cover-letter.{fmt}").status_code, 200)
        r = c.post("/api/documents", files={"file": ("ZQX-letterhead.docx", _docx("ZQX letterhead", template=True),
                                                     "application/octet-stream")}, data={"kind": "letter_template"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()["current_template"])
        info = c.get("/api/documents/letter-template").json()
        self.assertEqual(info["template"]["filename"], "ZQX-letterhead.docx")
        self.assertIn("ZQX Garamond", [f["asked"] for f in info["fonts"]])
        r = c.get(f"/api/jobs/{job_id}/cover-letter.pdf")
        self.assertEqual(r.status_code, 200)
        self.assertIn("ZQX Garamond", r.headers["x-font-substitutions"])
        self.assertEqual(c.get(f"/api/jobs/{job_id}/cover-letter.docx").status_code, 200)
        self.assertEqual(c.get(f"/api/jobs/{job_id}/cover-letter.odt").status_code, 409)  # a .docx template: no .odt
        r = c.post("/api/documents", files={"file": ("ZQX-no-marker.docx", _docx("ZQX no marker here"),
                                                     "application/octet-stream")}, data={"kind": "letter_template"})
        self.assertEqual(r.status_code, 400, r.text)
        self.assertEqual(c.put("/api/documents/letter-template", json={"id": None}).json()["template"], None)
        r = c.post("/api/companies/research", json={"mode": "all"})
        self.assertEqual(r.status_code, 200, r.text)
        listed = c.get("/api/tasks")  # the tasks tray, while the research may still run: kinds and ids only
        self.assertEqual(listed.status_code, 200, listed.text)
        self.assertNotIn(MARK, listed.text)
        self.assertEqual(self.run_to_end(c, r.json())["state"], "done")
        r = c.post("/api/companies/research", json={"name": "ZQX Corp"})
        self.assertEqual(self.run_to_end(c, r.json())["state"], "done")
        self.assertIn("zqx corp", c.get("/api/companies").json())

    def documents_and_saved_pages(self, c):
        """The AI's work on documents (a concise version, a summary of the ticked ones, title suggestions), a rename and
        the text, a place looked up for the walkthrough (in the query string), and saved pages: one imported as jobs
        and one uploaded as a job's full ad."""
        docs = c.get("/api/documents").json()
        cv = next(d for d in docs if d["kind"] == "resume")
        self.assertEqual(c.patch(f"/api/documents/{cv['id']}", json={"title": "ZQX renamed resume"}).status_code, 200)
        r = c.post(f"/api/documents/{cv['id']}/concise")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn(MARK, c.get(f"/api/documents/{cv['id']}/text").text)
        self.assertEqual(c.post("/api/documents/summary").status_code, 201)
        r = c.post("/api/documents/suggest-titles")
        self.assertLess(r.status_code, 500, r.text)
        self.assertTrue(c.get("/api/geo/lookup", params={"q": "ZQX Street, Zqxtown NSW"}).json()["found"])
        self.assertFalse(c.get("/api/geo/lookup", params={"q": "Zqxfail Road NSW"}).json()["found"])
        posting = {"@context": "https://schema.org", "@type": "JobPosting", "title": "ZQX Imported Lead",
                   "description": "ZQX saved page description.", "url": "https://jobs.zqx-saved.example/1",
                   "hiringOrganization": {"@type": "Organization", "name": "ZQX Saved Co"},
                   "jobLocation": {"@type": "Place", "address": {"addressLocality": "Zqxtown"}},
                   "datePosted": "2026-10-01"}
        page = (f'<html><head><title>ZQX saved</title><script type="application/ld+json">{json.dumps(posting)}'
                '</script></head><body>ZQX page body</body></html>').encode()
        r = c.post("/api/import/page", files={"file": ("ZQX saved page.html", page, "text/html")})
        self.assertEqual(r.status_code, 200, r.text)
        r = c.post(f"/api/jobs/{self.job_id}/import-page", files={"file": ("ZQX full ad.html", page, "text/html")})
        self.assertLess(r.status_code, 500, r.text)

    def commute_links(self, c):
        """Starting points (Search setup > Commute, labels and typed places), a refused one, and the job page's places
        for its Google Maps links, after an office is picked from the employer's offices."""
        profile = c.get("/api/profile").json()
        pin = c.get("/api/pins").json()[0]["id"]
        profile["starting_points"] = [{"label": "ZQX station", "place": "ZQX Station, Zqxtown NSW", "use": "transit"},
                                      {"label": "ZQX parents", "pin_id": pin, "use": "driving"}]
        self.assertEqual(c.put("/api/profile", json=profile).status_code, 200)
        stale = {**profile, "starting_points": [*profile["starting_points"],
                                                {"label": "ZQX nowhere", "pin_id": 987654, "use": "both"}]}
        r = c.put("/api/profile", json=stale)  # a pin that isn't in this search: dropped, the rest saved
        self.assertEqual((r.status_code, len(r.json()["starting_points"])), (200, 2))
        bad = {**profile, "starting_points": [{"label": "ZQX both", "pin_id": pin, "place": "ZQX Place", "use": "both"}]}
        self.assertEqual(c.put("/api/profile", json=bad).status_code, 422)
        self.assertEqual(c.put(f"/api/jobs/{self.job_id}/office", json={"office": "1 Zqx Street, Zqxtown NSW"}).status_code,
                         200)
        r = c.get(f"/api/jobs/{self.job_id}/commute")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn("ZQX Station, Zqxtown NSW", [s["origin"] for s in r.json()["starts"]])

    def browser_extension(self, c):
        """The browser extension: a key named by the person, then jobs it read from pages in their browser (the
        values, the description and the addresses), previewed and saved, a page of search results (each card's
        address and summary), and a source added and tested through the key."""
        from backend import ratelimit
        from backend.adapters.base import SourceError

        r = c.post("/api/extension/keys", json={"label": "ZQX Firefox on Zqx's laptop"})
        self.assertEqual(r.status_code, 201, r.text)
        self.secrets["extension_key"] = r.json()["key"]
        self.assertIn("ZQX", c.get("/api/extension/keys").json()[0]["label"])
        k = self.anonymous()
        k.headers["Authorization"] = f"Bearer {r.json()['key']}"
        self.assertIn("ZQX", k.get("/api/extension/status").json()["label"])
        ad = "ZQX include: lead the ZQX delivery program for Zqx Quillon's team. " * 6
        page = {"url": "https://careers.zqx.example/jobs/zqx-77?zqx=1", "canonical_url": "https://careers.zqx.example/c/zqx",
                "job": {"title": "ZQX Delivery Lead", "company": "ZQX Corp", "location_text": "Zqxtown NSW",
                        "salary_text": "ZQX $150k", "employment_type": "ZQX full time", "work_mode": "hybrid",
                        "posted_at": "2026-10-01", "url": "https://careers.zqx.example/jobs/zqx-77",
                        "description": ad, "description_from": "site", "method": "site"}}
        self.assertEqual(k.post("/api/extension/preview", json=page).status_code, 200)
        self.assertEqual(k.post("/api/extension/jobs", json=page).status_code, 201)
        edited = {**page, "job": {**page["job"], "description": "ZQX selected ad text", "description_from": "selection"}}
        self.assertEqual(k.post("/api/extension/jobs", json=edited).status_code, 200)
        self.assertEqual(k.post("/api/extension/preview", json=page).json()["existing"]["title"], "ZQX Delivery Lead")
        for bad in ({**page, "job": {**page["job"], "title": ""}}, {**page, "url": "ftp://zqx.example/zqx"},
                    {**page, "job": {**page["job"], "work_mode": "ZQX moon"}}, {**page, "html": "<p>ZQX</p>"},
                    {**page, "job": {**page["job"], "posted_at": "ZQX yesterday"}}):
            self.assertEqual(k.post("/api/extension/jobs", json=bad).status_code, 422)
        too_long = {**page, "job": {**page["job"], "title": "ZQX " * 200}}
        self.assertEqual(k.post("/api/extension/jobs", json=too_long).status_code, 413)
        listed = {"url": "https://jobs.zqx.example/search?q=zqx+delivery&zqx=new",
                  "jobs": [{"url": f"/roles/zqx-{n}?zqx=card", "title": f"ZQX Program Lead {n}", "company": "ZQX Corp",
                            "location_text": "Zqxtown NSW", "salary_text": "ZQX $160k", "employment_type": "ZQX full time",
                            "work_mode": "remote", "posted_at": "2026-10-09",
                            "description": "ZQX include: run the ZQX program for Zqx Quillon.\n\n- ZQX point"}
                           for n in (1, 2)]}
        r = k.post("/api/extension/jobs/batch", json=listed)
        self.assertEqual((r.status_code, r.json()["added"]), (200, 2), r.text)
        self.assertEqual(k.post("/api/extension/jobs/batch", json=listed).json()["existing"], 2)
        for bad in ({**listed, "jobs": [{**listed["jobs"][0], "title": ""}]},
                    {**listed, "jobs": [{**listed["jobs"][0], "url": "https://other.zqx.example/roles/zqx-9"}]},
                    {**listed, "jobs": [{**listed["jobs"][0], "posted_at": "ZQX 3d ago"}]},
                    {**listed, "jobs": [{**listed["jobs"][0], "html": "<article>ZQX</article>"}]}):
            self.assertEqual(k.post("/api/extension/jobs/batch", json=bad).status_code, 422)
        too_long = {**listed, "jobs": [{**listed["jobs"][0], "description": "ZQX " * 2000}]}
        self.assertEqual(k.post("/api/extension/jobs/batch", json=too_long).status_code, 413)
        r = k.post("/api/sources", json={"url": "https://careers.zqx-two.example/jobs", "label": "ZQX second employer"})
        self.assertEqual(r.status_code, 201, r.text)
        self.source.postings = [_ad(20)]
        self.source.notes = []  # (what a source test shows besides the sample)
        with mock.patch("backend.sources.source_adapter", side_effect=lambda *a, **kw: self.source):
            self.assertEqual(k.post(f"/api/sources/{r.json()['id']}/test").json()["count"], 1)
            self.source.failure = SourceError("ZQX source down: https://careers.zqx-two.example/jobs")
            self.assertFalse(k.post(f"/api/sources/{r.json()['id']}/test").json()["ok"])
            self.source.failure = None
        self.assertEqual(k.get("/api/jobs").status_code, 403)
        self.ai_mappings(c, k)
        wrong = self.anonymous()
        wrong.headers["Authorization"] = "Bearer " + self.secrets["extension_key"][:-4] + "zqxq"
        self.secrets["wrong_extension_key"] = wrong.headers["Authorization"][7:]
        self.assertEqual(wrong.get("/api/extension/status").status_code, 401)
        ratelimit.BAD_KEY.reset()

    @staticmethod
    def outline(n: int = 5) -> dict:
        """A page outline as the extension sends it: the person's page, marked everywhere."""
        return {"title": "ZQX Delivery Lead at ZQX Corp", "root": {"tag": "body", "attrs": {"class": "zqx-page"},
                "children": [{"tag": "h1", "attrs": {"class": "zqx-title", "id": "zqx-h1"}, "text": "ZQX Delivery Lead"},
                             {"tag": "a", "attrs": {"href": "/zqx/apply?zqx=token", "aria-label": "ZQX apply"},
                              "text": "ZQX apply"},
                             {"tag": "div", "attrs": {"id": "zqx-ad", "data-testid": "zqx-ad"}, "children": [
                                 {"tag": "p", "text": f"ZQX paragraph {i} for Zqx Quillon"} for i in range(n)]}]}}

    def ai_mappings(self, c, k):
        """AI mappings (Phase I1): a job page and a careers page worked out from the person's page outline, answers
        that aren't selectors refused, an oversized outline refused, the layouts saved, read and forgotten. The
        outline, the sites and the selectors never reach the log."""
        from tests.test_mappings import Pages

        site = "jobs.zqx-site.example"
        job = {"kind": "job", "url": f"https://{site}/zqx-role/7?zqx=private", "outline": self.outline()}
        r = k.post("/api/extension/work-out", json=job)  # the AI answers with a score: not a mapping
        self.assertEqual((r.status_code, r.json()["code"]), (422, "not_worked_out"), r.text)
        self.mapping_reply = {
            "job": json.dumps({"title": {"selector": "h1.zqx-title"}, "description": {"selector": "#zqx-ad"},
                               "company": {"selector": "[data-testid='zqx-company']"}}),
            "careers": json.dumps({"list": {"selector": "li.zqx-job"}, "title": {"selector": "h3"},
                                   "link": {"selector": "a", "attribute": "href"}})}
        r = k.post("/api/extension/work-out", json=job)
        self.assertEqual(r.status_code, 200, r.text)
        mapping = r.json()["mapping"]
        self.assertEqual(mapping["fields"]["title"], {"selector": "h1.zqx-title"})
        for refused in ({"title": {"selector": "a[href^='javascript:zqx()']"}, "description": {"selector": "#zqx"}},
                        {"title": {"selector": "h1"}, "description": {"selector": "#zqx"}, "zqx": {"selector": "p"}}):
            self.mapping_reply["job"] = json.dumps(refused)
            self.assertEqual(k.post("/api/extension/work-out", json=job).status_code, 422)
        big = {**job, "outline": self.outline(2100)}
        self.assertEqual(k.post("/api/extension/work-out", json=big).status_code, 413)
        bad = {**job, "outline": {"root": {"tag": "script", "text": "ZQX token = 1"}}}
        self.assertEqual(k.post("/api/extension/work-out", json=bad).status_code, 422)
        self.assertEqual(k.put(f"/api/extension/mappings/{site}", json={"mapping": mapping}).status_code, 201)
        self.assertEqual(c.get(f"/api/extension/mappings/{site}").json()["site"], site)
        self.assertEqual([m["site"] for m in k.get("/api/extension/mappings").json()["mappings"]], [site])
        bad_mapping = {**mapping, "fields": {**mapping["fields"], "title": {"selector": "zqx:hover"}}}
        self.assertEqual(k.put(f"/api/extension/mappings/{site}", json={"mapping": bad_mapping}).status_code, 422)
        careers = "https://careers.zqx-three.example/zqx-jobs"
        page = ('<html><body><ul>' + "".join(f'<li class="zqx-job"><h3>ZQX Analyst {n}</h3>'
                                             f'<a href="/zqx/o/{n}">ZQX more</a></li>' for n in range(3))
                + "</ul></body></html>")
        with Pages({careers: page}).served():
            r = k.post("/api/extension/work-out", json={"kind": "careers", "url": careers, "outline": self.outline()})
        self.assertEqual((r.status_code, r.json()["test"]["count"]), (200, 3), r.text)
        r = k.post("/api/sources", json={"url": careers, "label": "ZQX careers", "mapping": r.json()["mapping"]})
        self.assertEqual(r.status_code, 201, r.text)
        self.assertEqual(c.delete(f"/api/extension/mappings/{site}").status_code, 200)
        self.assertEqual(k.get(f"/api/extension/mappings/{site}").status_code, 404)

    def when_things_fail(self, c):
        """Errors carry the text that caused them (a provider echoes the prompt, a library names the address): none
        of it may be logged, only what kind of failure it was."""
        from backend.adapters.base import SourceError
        from backend.scorer import ScorerError

        r = c.post("/api/documents", files={"file": ("ZQX-broken.pdf", b"ZQX is not a pdf", "application/pdf")},
                   data={"kind": "resume"})
        self.assertEqual(r.status_code, 400, r.text)
        # The AI service refuses (a bad key, a bad request): scoring, chat, a cover letter and research all fail.
        self.llm_status = 401
        for path, body in (("/api/chats", {"job_ids": [self.job_id], "message": "ZQX failing question"}),
                           (f"/api/jobs/{self.job_id}/cover-letter", {"instructions": "ZQX failing instruction"}),
                           (f"/api/score/{self.job_id}", None),
                           ("/api/extension/work-out", {"kind": "job", "url": "https://jobs.zqx-fail.example/zqx",
                                                        "outline": self.outline()})):
            self.assertGreaterEqual(c.post(path, json=body).status_code, 400, path)
        self.assertEqual(self.run_to_end(c, c.post("/api/score/run", json={"mode": "all"}).json())["state"], "done")
        self.llm_status = 200
        with mock.patch("backend.research.run_response", side_effect=ScorerError("ZQX research failed for ZQX Corp")):
            r = c.post("/api/companies/research", json={"name": "ZQX Corp"})
            self.run_to_end(c, r.json())
        # A source breaks, with its address in the message; a whole run breaks, with the person's address in it.
        for failure in (SourceError("ZQX source down: https://careers.zqx.example/jobs"),
                        RuntimeError("ZQX unexpected, for zqxowner.private@example.invalid")):
            self.source.failure = failure
            self.run_to_end(c, c.post("/api/search/run").json())
        self.source.failure = None
        with mock.patch("backend.pipeline.run_search", side_effect=RuntimeError("ZQX run failed for zqxowner.private")):
            self.assertEqual(self.run_to_end(c, c.post("/api/search/run").json())["state"], "failed")
        self.assertEqual(self.anonymous().post("/api/auth/login", json={"email": "zqxnobody@example.invalid",
                                                                        "password": "zqx nobody"}).status_code, 401)

    def with_claude(self, c):
        """The same through the Claude SDK (its own logger, which logs whole requests at DEBUG), over a mock transport."""
        import anthropic
        import httpx2

        def handler(request: httpx2.Request) -> httpx2.Response:
            self.assertIn(b"ZQX", request.content)  # the prompt really carries the person's data
            return httpx2.Response(200, json={
                "id": "msg_1", "type": "message", "role": "assistant", "model": "claude-test", "stop_reason": "end_turn",
                "content": [{"type": "text", "text": SCORE_REPLY}], "usage": {"input_tokens": 10, "output_tokens": 5}})

        client = anthropic.Anthropic(api_key="zqx-claude-key", max_retries=0,
                                     http_client=httpx2.Client(transport=httpx2.MockTransport(handler)))
        with mock.patch.dict(os.environ, {"LLM_BASE_URL": "https://api.anthropic.com", "LLM_MODEL": "claude-sonnet-4-5",
                                          "LLM_API_KEY": "zqx-claude-key"}), \
                mock.patch("backend.llm_anthropic._client", return_value=client):
            self.assertEqual(c.post(f"/api/score/{self.job_id}").status_code, 200)
            r = c.post("/api/chats", json={"job_ids": [self.job_id], "message": "ZQX Claude question"})
            self.assertEqual(r.status_code, 200, r.text)
            r = c.post(f"/api/jobs/{self.job_id}/cover-letter", json={"instructions": "ZQX Claude instruction"})
            self.assertEqual(r.status_code, 200, r.text)

    def export(self, c):
        ws = c.get("/api/workspaces").json()[0]["id"]
        self.assertEqual(c.patch(f"/api/workspaces/{ws}", json={"name": "ZQX dream search"}).status_code, 200)
        r = c.post("/api/account/export", json={"formats": ["md", "docx", "odt"]})
        self.assertEqual(r.status_code, 200, r.text)
        names = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
        text = b"".join(zipfile.ZipFile(io.BytesIO(r.content)).read(n) for n in names)
        self.assertIn(b"ZQX", text, "the export should hold the planted data")
        self.assertIn(b"zqxowner", text.lower())
        self.import_it_back(c, r.content)

    def import_it_back(self, c, blob):
        """The export brought back as a new workspace (the names, documents, ads and chats in it are the person's), and
        an import file that is refused: its refusal names only where the problem is."""
        headers = {"Content-Type": "application/zip"}
        r = c.post("/api/account/import/check", content=blob, headers=headers)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn("ZQX", json.dumps(r.json()))
        r = c.post("/api/account/import", content=blob, headers=headers, params={"pick": "0"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertGreater(r.json()["workspaces"][0]["jobs"], 0)
        z = zipfile.ZipFile(io.BytesIO(blob))
        files = {n: z.read(n) for n in z.namelist()}
        name = next(n for n in files if n.endswith("import file.json"))
        data = json.loads(files[name])
        data["workspaces"][0]["jobs"][0]["status"] = "ZQX not a status"
        files[name] = json.dumps(data).encode()
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as bad:
            for n, b in files.items():
                bad.writestr(n, b)
        r = c.post("/api/account/import", content=out.getvalue(), headers=headers, params={"pick": "0"})
        self.assertEqual(r.status_code, 422, r.text)
        self.assertNotIn("ZQX", r.text)

    def merges_a_copy_and_the_password(self, c):
        """Two planted ads merged and unmerged (by hand and as a suggested pair), a workspace copied from this one with
        its documents, the password changed (planted old and new values) and changed back, and new recovery codes
        made and shown again with the password."""
        jobs = [j for j in c.get("/api/jobs?include=excluded,hidden").json() if not j.get("duplicate_of")]
        self.assertGreaterEqual(len(jobs), 2, "two planted ads to merge")
        a, b = jobs[0]["id"], jobs[1]["id"]
        r = c.post("/api/jobs/merge", json={"ids": [a, b]})
        self.assertEqual(r.status_code, 200, r.text)
        other = b if r.json()["id"] == a else a
        self.assertEqual(c.post(f"/api/jobs/{other}/unmerge").status_code, 200)
        for decision in ("merge", "distinct"):
            r = c.post("/api/duplicates/decide", json={"a": a, "b": b, "decision": decision})
            self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(c.get("/api/duplicates").status_code, 200)

        ws = c.get("/api/workspaces").json()[0]["id"]
        r = c.post("/api/workspaces", json={"name": "ZQX copied search", "copy_from": ws, "copy_documents": True})
        self.assertEqual(r.status_code, 201, r.text)

        changed = "zqx-changed-password-3"
        r = c.post("/api/auth/password", json={"current_password": "zqx wrong current", "new_password": changed,
                                                "sign_out_others": False})
        self.assertEqual(r.status_code, 422, r.text)
        r = c.post("/api/auth/password", json={"current_password": PASSWORD, "new_password": changed,
                                                "sign_out_others": False})
        self.assertEqual(r.status_code, 200, r.text)
        r = c.post("/api/auth/password", json={"current_password": changed, "new_password": PASSWORD,
                                                "sign_out_others": False})
        self.assertEqual(r.status_code, 200, r.text)
        self.secrets.update(changed_password=changed, wrong_current="zqx wrong current")

        r = c.post("/api/auth/recovery-codes", json={"password": PASSWORD})
        self.assertEqual(r.status_code, 200, r.text)
        codes = r.json()["recovery_codes"]
        self.assertEqual(c.post("/api/auth/recovery-codes/show", json={"password": "zqx wrong password"}).status_code,
                         422)
        r = c.post("/api/auth/recovery-codes/show", json={"password": PASSWORD})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["codes"], codes)
        self.secrets.update({f"recovery_code_{i}": code for i, code in enumerate(codes)})
        ratelimit.reset_all()  # (the two wrong passwords above would otherwise slow the sign-ins of the last step)

    def scheduled_search_and_recovery(self, c, notice):
        # Background access (so scheduled runs can open the data), then a scheduled run.
        self.assertEqual(c.post("/api/account/background", json={"enabled": True}).status_code, 200)
        db = SessionLocal()
        try:
            uid = users.find_by_name(db, self.email).id
            ws = db.query(Workspace.id).filter(Workspace.owner_id == uid).order_by(Workspace.id).first()[0]
        finally:
            db.close()
        acct = helpers.Account(uid, self.email, PASSWORD, ws)
        self.source.postings = [_ad(10, title="ZQX Delivery Lead scheduled one"),
                                _ad(11, title="ZQX Delivery Lead scheduled two")]
        with helpers.as_account(acct):
            run = app_module.start_search_run(ws, scheduled=True)
        self.assertEqual(self.run_to_end(c, run)["state"], "done")

        # A forgotten password, reset with a recovery code (a wrong one first): the data is kept.
        anon = self.anonymous()
        r = anon.post("/api/auth/recover", json={"name": self.email, "recovery_code": "ZQXQ-ZQXQ-ZQXQ-ZQXQ",
                                                 "new_password": NEW_PASSWORD})
        self.assertEqual(r.status_code, 401, r.text)
        r = anon.post("/api/auth/recover", json={"name": self.email, "recovery_code": self.secrets["recovery_code_0"],
                                                 "new_password": NEW_PASSWORD})
        self.assertEqual(r.status_code, 200, r.text)
        self.secrets.update(new_password=NEW_PASSWORD, recovered_csrf=r.json()["csrf"])
        self.assertEqual(anon.post("/api/auth/login", json={"name": self.email, "password": "zqx wrong password"}
                                   ).status_code, 401)
        notice.append(True)


class StripQueryTest(unittest.TestCase):
    """uvicorn's access log: the query string (which can carry a token) is dropped."""

    def test_the_access_log_has_no_query_string(self):
        record = logging.LogRecord("uvicorn.access", logging.INFO, "", 0, '%s - "%s %s HTTP/%s" %d',
                                   ("127.0.0.1:5", "GET", "/verify?token=ZQXsecret&x=1", "1.1", 200), None)
        logging.getLogger("uvicorn.access").filters[0].filter(record)
        self.assertNotIn("ZQXsecret", record.getMessage())
        self.assertIn("/verify ", record.getMessage())

    def test_the_quiet_loggers_stay_quiet(self):
        """What the app holds back is held back even with the root logger at DEBUG."""
        self.assertTrue({"httpx", "httpcore", "httpx2", "httpcore2", "anthropic", "pypdf"} <= set(app_module.QUIET_LOGGERS))
        root = logging.getLogger()
        handler, old = Everything(), root.level
        root.addHandler(handler)
        root.setLevel(logging.DEBUG)
        try:
            for name in app_module.QUIET_LOGGERS:
                self.assertGreaterEqual(logging.getLogger(name).getEffectiveLevel(), logging.WARNING, name)
            logging.getLogger("pypdf._reader").warning("invalid pdf header: b'ZQX private'")  # (made on demand, like pypdf's)
            logging.getLogger("anthropic._base_client").debug("Request options: ZQX prompt")
            logging.getLogger("httpx2").info("HTTP Request: GET https://api.example/?q=ZQX")
        finally:
            root.removeHandler(handler)
            root.setLevel(old)
        self.assertEqual(handler.lines, [])

    def test_sqlalchemy_never_logs_statements_or_rows(self):
        """At DEBUG it would log every row it reads, opened (decrypted) by the app's column types."""
        helpers.started()
        sql = logging.getLogger("sqlalchemy.engine.Engine")
        handler, old = Everything(), sql.level
        sql.addHandler(handler)
        sql.setLevel(logging.DEBUG)
        try:
            db = SessionLocal()
            try:
                self.assertTrue(db.query(User.id, User.status).all())
            finally:
                db.close()
            sql.warning("a warning still comes through")
        finally:
            sql.removeHandler(handler)
            sql.setLevel(old)
        self.assertEqual(handler.lines, ["sqlalchemy.engine.Engine WARNING a warning still comes through"])


class FolderStepLogTest(unittest.TestCase):
    """The first run's step 1 takes a folder path, which can hold the person's name (``/home/<name>/...``): no log line
    may hold it, whether the folder is checked, refused or chosen (backend/data_folder.py logs the kind only). It runs
    in a process of its own, before any data folder exists (tests/test_data_folder.py has the sandbox)."""

    def test_the_folder_path_never_reaches_a_log(self):
        from tests.test_data_folder import TOKEN, Sandbox

        box = Sandbox()
        self.addCleanup(box.remove)
        folder = box.home / "ZQX Person Name" / "JobHuntPA"
        busy = box.top / "ZQX busy"
        busy.mkdir()
        (busy / "ZQX-private.txt").write_text("mine")
        got = box.run(f'''
            import logging
            logging.getLogger().setLevel(logging.DEBUG)
            for handler in logging.getLogger().handlers:
                handler.setLevel(logging.DEBUG)
            check = "/api/auth/data-folder/check"
            c.post(check, json={{"token": {TOKEN!r}, "path": {str(folder)!r}}})
            c.post(check, json={{"token": {TOKEN!r}, "path": "relative ZQX folder"}})
            c.post(check, json={{"token": {TOKEN!r}, "path": {str(busy)!r}}})
            c.post("/api/auth/data-folder", json={{"token": {TOKEN!r}, "path": {str(busy)!r}}})
            r = c.post("/api/auth/data-folder", json={{"token": {TOKEN!r}, "path": {str(folder)!r}}})
            out["status"] = r.status_code
        ''')
        self.assertEqual(got["status"], 200)
        self.assertIn("data folder is chosen", got["_log"])  # (the log isn't empty)
        self.assertNotIn(MARK, got["_log"])


if __name__ == "__main__":
    unittest.main()
